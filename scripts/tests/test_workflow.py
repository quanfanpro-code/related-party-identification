# 关联方识别与核查 — 原创编排层
# Copyright (C) 2026 CPA-Q (quanfanpro-code)
#
# 本文件是 related-party-identification 的原创编排层,采用 GNU Affero General
# Public License v3.0 (AGPL-3.0) 发布。
# 完整协议见项目根目录 LICENSE 文件(AGPL-3.0)。
#
# 本项目的 scripts/cicpa/ 目录包含改编自 nigo/nigo-skills(MIT) 和
# jackwener/OpenCLI(Apache-2.0) 的代码,分别保留原始许可证。
# 详见 NOTICE 和 references/SOURCES.json。
import json
from datetime import datetime
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

import openpyxl

from scripts.cicpa.exporter import ExportState, save_export_state
from scripts.discovery import Candidate, DiscoveryResult
from scripts.related_party_workflow import (
    TaskState,
    auth_status,
    ensure_login,
    extract_seed_export_candidates,
    load_task_state,
    pick_path,
    preflight,
    read_name_list,
    run_workflow,
    save_task_state,
    setup_opencli,
)
from scripts.validate_bundle import validate_bundle


class FakeRoot:
    def __init__(self):
        self.withdrawn = False
        self.destroyed = False

    def withdraw(self):
        self.withdrawn = True

    def attributes(self, *_args):
        return None

    def destroy(self):
        self.destroyed = True


class FakeFileDialog:
    def __init__(self, file_result="", directory_result=""):
        self.file_result = file_result
        self.directory_result = directory_result
        self.file_calls = []
        self.directory_calls = []

    def askopenfilename(self, **kwargs):
        self.file_calls.append(kwargs)
        return self.file_result

    def askdirectory(self, **kwargs):
        self.directory_calls.append(kwargs)
        return self.directory_result


class FakeStore:
    def __init__(self, cookies=None, error=None):
        self.cookies = cookies or {}
        self.error = error

    def load_cookies(self):
        if self.error:
            raise self.error
        return self.cookies


class PreflightTests(unittest.TestCase):
    def test_缺少依赖只给独立环境建议而不执行安装(self):
        command_calls = []

        def find_spec(name):
            return None if name == "requests" else object()

        result = preflight(
            platform_name="nt",
            version_info=(3, 11, 0),
            import_finder=find_spec,
            browser_detector=lambda: {
                "name": "Microsoft Edge",
                "executable_path": r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            },
            auth_checker=lambda: {"status": "idle", "message_zh": "尚未登录"},
            local_app_data=r"C:\Users\测试\AppData\Local",
            command_runner=lambda command: command_calls.append(command),
        )

        self.assertEqual(result["status"], "needs_user_approval")
        self.assertEqual(result["missing_dependencies"], ["requests"])
        self.assertEqual(
            result["recommended_venv"],
            r"C:\Users\测试\AppData\Local\related-party-identification\venv",
        )
        self.assertEqual(command_calls, [])

    def test_python_低于三点九会停止(self):
        result = preflight(
            platform_name="nt",
            version_info=(3, 8, 10),
            import_finder=lambda _name: object(),
            browser_detector=lambda: {"name": "Microsoft Edge", "executable_path": "edge.exe"},
            auth_checker=lambda: {"status": "idle", "message_zh": "尚未登录"},
            local_app_data=r"C:\Users\测试\AppData\Local",
        )

        self.assertEqual(result["status"], "unsupported_python")
        self.assertIn("3.9", result["message_zh"])

    def test_预检结果不包含认证秘密(self):
        result = preflight(
            platform_name="nt",
            version_info=(3, 11, 0),
            import_finder=lambda _name: object(),
            browser_detector=lambda: {"name": "Microsoft Edge", "executable_path": "edge.exe"},
            auth_checker=lambda: {"status": "authenticated", "message_zh": "登录有效"},
            opencli_checker=lambda _path: True,
            local_app_data=r"C:\Users\测试\AppData\Local",
        )
        serialized = json.dumps(result, ensure_ascii=False).lower()

        self.assertNotIn("cookie", serialized)
        self.assertNotIn("password", serialized)
        self.assertNotIn("authorization", serialized)

    def test_edge_未安装_opencli_时先推荐_firefox_再给安装路线(self):
        result = preflight(
            platform_name="nt",
            version_info=(3, 11, 0),
            import_finder=lambda _name: object(),
            browser_detector=lambda: {
                "name": "Microsoft Edge",
                "executable_path": r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            },
            opencli_checker=lambda _path: False,
            auth_checker=lambda: {"status": "idle", "message_zh": "尚未登录"},
            local_app_data=r"C:\Users\测试\AppData\Local",
        )

        self.assertEqual(result["status"], "needs_opencli_extension")
        self.assertLess(
            result["message_zh"].index("Firefox"),
            result["message_zh"].index("OpenCLI"),
        )
        self.assertIn("Gitee 国内镜像", result["message_zh"])

    def test_opencli_安装准备只让用户先做一个动作(self):
        extension_dir = Path(tempfile.mkdtemp(prefix="rpi_opencli_extension_"))
        opened = []
        browser = {
            "name": "Microsoft Edge",
            "executable_path": r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        }

        result = setup_opencli(
            browser_detector=lambda: browser,
            downloader=lambda: extension_dir,
            guide_opener=lambda selected, path: opened.append((selected, path)),
        )

        self.assertEqual(result["status"], "waiting_user")
        self.assertEqual(opened, [(browser, extension_dir)])
        self.assertIn("开发人员模式", result["message_zh"])
        self.assertIn("打开了", result["message_zh"])


class FilePickerTests(unittest.TestCase):
    def test_输入文件路径来自_windows_选择窗口(self):
        root = FakeRoot()
        dialog = FakeFileDialog(file_result=r"C:\审计资料\客户名单.xlsx")

        result = pick_path(
            "input-file",
            root_factory=lambda: root,
            filedialog_module=dialog,
        )

        self.assertEqual(result["status"], "selected")
        self.assertEqual(result["path"], r"C:\审计资料\客户名单.xlsx")
        self.assertTrue(root.withdrawn)
        self.assertTrue(root.destroyed)
        self.assertEqual(len(dialog.file_calls), 1)

    def test_取消选择返回可恢复状态(self):
        root = FakeRoot()
        dialog = FakeFileDialog(directory_result="")

        result = pick_path(
            "output-folder",
            root_factory=lambda: root,
            filedialog_module=dialog,
        )

        self.assertEqual(result["status"], "cancelled")
        self.assertEqual(result["path"], "")
        self.assertTrue(root.destroyed)


class TaskStateTests(unittest.TestCase):
    def test_任务状态可恢复且_json_不带_bom(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_workflow_state_"))
        state_path = artifact_dir / "task.json"
        state = TaskState(
            task_id="task-001",
            mode="discovery",
            audited_entity="甲公司",
            stage="preflight",
            status="waiting_user",
            output_dir=r"C:\审计输出",
        )

        save_task_state(state_path, state)
        loaded = load_task_state(state_path)
        raw = state_path.read_bytes()

        self.assertEqual(loaded, state)
        self.assertFalse(raw.startswith(b"\xef\xbb\xbf"))
        serialized = raw.decode("utf-8").lower()
        self.assertNotIn("cookie", serialized)
        self.assertNotIn("password", serialized)


class AuthenticationStatusTests(unittest.TestCase):
    def test_已有加密状态仍以官方接口验证结果为准(self):
        result = auth_status(
            store=FakeStore({"cicpa_token": "secret"}),
            verifier=lambda _cookies: True,
        )

        self.assertEqual(result["status"], "authenticated")
        self.assertEqual(set(result), {"status", "message_zh"})

    def test_失效状态不回显_cookie(self):
        result = auth_status(
            store=FakeStore({"cicpa_token": "secret"}),
            verifier=lambda _cookies: False,
        )
        serialized = json.dumps(result, ensure_ascii=False).lower()

        self.assertEqual(result["status"], "expired")
        self.assertNotIn("cookie", serialized)
        self.assertNotIn("secret", serialized)

    def test_每次调用先检查_已有登录可用就直接继续(self):
        starts = []

        result = ensure_login(
            auth_checker=lambda: {
                "status": "authenticated",
                "message_zh": "现有登录状态有效",
            },
            login_starter=lambda: starts.append(True),
        )

        self.assertEqual(result["status"], "authenticated")
        self.assertEqual(starts, [])

    def test_每次调用发现登录失效就自动打开浏览器(self):
        starts = []

        result = ensure_login(
            auth_checker=lambda: {
                "status": "expired",
                "message_zh": "现有登录状态已失效",
            },
            login_starter=lambda: starts.append(True) or {
                "status": "waiting_user",
                "message_zh": "已打开系统浏览器",
            },
        )

        self.assertEqual(result["status"], "waiting_user")
        self.assertEqual(starts, [True])


class ScriptEntryTests(unittest.TestCase):
    def test_直接运行脚本时_help_可用(self):
        skill_root = Path(__file__).resolve().parents[2]
        script = skill_root / "scripts" / "related_party_workflow.py"

        completed = subprocess.run(
            [sys.executable, "-X", "utf8", str(script), "--help"],
            cwd=skill_root,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn("关联方识别 AI 对话式工作流", completed.stdout)


class NameListTests(unittest.TestCase):
    def test_xlsx_csv_txt_只提取公司名称并去重(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_name_lists_"))
        xlsx_path = artifact_dir / "客户名单.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["客户名称", "金额"])
        sheet.append(["甲客户有限公司", 100])
        sheet.append(["甲客户有限公司", 200])
        sheet.append(["123456", 300])
        workbook.save(xlsx_path)
        csv_path = artifact_dir / "供应商名单.csv"
        csv_path.write_text(
            "供应商名称,备注\n乙供应商有限公司,有效\n",
            encoding="utf-8-sig",
        )
        txt_path = artifact_dir / "其他名单.txt"
        txt_path.write_text("丙公司\n\n丙公司\n", encoding="utf-8-sig")

        xlsx_result = read_name_list(xlsx_path)
        csv_result = read_name_list(csv_path)
        txt_result = read_name_list(txt_path)

        self.assertEqual(xlsx_result.names, ["甲客户有限公司"])
        self.assertEqual(xlsx_result.rejected, ["123456"])
        self.assertEqual(csv_result.names, ["乙供应商有限公司"])
        self.assertEqual(txt_result.names, ["丙公司"])

    def test_多列且无明确名称列时暂停让用户确认(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_name_columns_"))
        path = artifact_dir / "待确认.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["字段甲", "字段乙"])
        sheet.append(["甲公司", "乙公司"])
        workbook.save(path)

        result = read_name_list(path)

        self.assertEqual(result.status, "needs_column_confirmation")
        self.assertEqual(result.candidate_columns, ["字段甲", "字段乙"])

    def test_总账和明细账不会被当成名单解析(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_ledger_reject_"))
        path = artifact_dir / "销售明细账.xlsx"
        path.write_bytes(b"not-used")

        result = read_name_list(path)

        self.assertEqual(result.status, "unsupported_input")
        self.assertIn("名单", result.message_zh)


def make_seed_counterparty_export(
    path,
    owner_name,
    relation_label,
    counterparty_name,
):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "序号",
        "公司名称",
        "公告时间",
        "金额",
        "占比（%）",
        "与本公司关系",
        "货币代码",
        "关联方名称",
        "关联方ID",
        "关联理由",
    ])
    sheet.append([
        1,
        owner_name,
        "2026-07-01",
        100,
        10,
        relation_label,
        "CNY",
        counterparty_name,
        "org-1",
        "公开公告",
    ])
    workbook.save(path)


def copy_with_wrong_dimension(source, destination):
    with zipfile.ZipFile(source, "r") as input_zip:
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as output_zip:
            for item in input_zip.infolist():
                content = input_zip.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    text = content.decode("utf-8")
                    text = re.sub(
                        r'<dimension ref="[^"]+"',
                        '<dimension ref="A1:A2"',
                        text,
                        count=1,
                    )
                    content = text.encode("utf-8")
                output_zip.writestr(item, content)


class SeedExportCandidateTests(unittest.TestCase):
    def test_客户供应商第八列是公开交易对手而不是关联方认定(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_counterparty_extract_"))
        make_seed_counterparty_export(
            artifact_dir / "客户.xlsx",
            "甲公司",
            "客户",
            "客户公司",
        )
        make_seed_counterparty_export(
            artifact_dir / "供应商.xlsx",
            "甲公司",
            "供应商",
            "供应商公司",
        )

        candidates = extract_seed_export_candidates(artifact_dir)

        self.assertEqual(
            [(item.name, item.relation_type) for item in candidates],
            [("客户公司", "公开客户关系"), ("供应商公司", "公开供应商关系")],
        )
        self.assertNotIn("甲公司", [item.name for item in candidates])

    def test_错误工作表范围信息不影响交易对手提取(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_wrong_dimension_"))
        source = artifact_dir / "原始客户.xlsx"
        customer_path = artifact_dir / "客户.xlsx"
        make_seed_counterparty_export(
            source,
            "甲公司",
            "客户",
            "客户公司",
        )
        copy_with_wrong_dimension(source, customer_path)

        candidates = extract_seed_export_candidates(artifact_dir)

        self.assertEqual(
            [(item.name, item.relation_type) for item in candidates],
            [("客户公司", "公开客户关系")],
        )


class FakeChecker:
    def __init__(self):
        self.calls = []

    def __call__(self, **kwargs):
        self.calls.append(kwargs)
        output = Path(kwargs["output_path"])
        workbook = openpyxl.Workbook()
        workbook.save(output)
        return SimpleNamespace(output_path=output, summary=[], errors=[], limitations=[])


class FakeExporter:
    def __init__(self, extract_dir):
        self.extract_dir = Path(extract_dir)
        self.start_calls = []
        self.wait_calls = []
        self.download_calls = []

    def start_export(self, company_names, state_path):
        self.start_calls.append(list(company_names))
        state = ExportState(
            batch_no="batch-1",
            company_names=list(company_names),
            created_at="2026-07-26 16:00",
            triggered_at="2026-07-26 16:00",
            staging_dir=str(Path(state_path).parent),
            status="waiting",
        )
        save_export_state(state_path, state)
        return state

    def wait_for_task(self, state, state_path, max_polls=18):
        self.wait_calls.append(state.batch_no)
        return {
            "task_id": "task-1",
            "batch_id": state.batch_no,
            "type": "批量查询",
            "status": 1,
            "date": "2026-07-26 16:01",
            "url": "/task-1.zip",
        }

    def download_and_validate(self, state, task, state_path):
        self.download_calls.append(task["task_id"])
        state.status = "completed"
        state.task_id = task["task_id"]
        state.extract_dir = str(self.extract_dir)
        save_export_state(state_path, state)
        return self.extract_dir


class WorkflowModeTests(unittest.TestCase):
    def test_已有数据模式不触发登录或注协查询(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_existing_mode_"))
        data_dir = artifact_dir / "existing_files"
        data_dir.mkdir()
        checker = FakeChecker()
        state_path = artifact_dir / "task.json"
        state = TaskState(
            task_id="existing-1",
            mode="existing_export",
            audited_entity="甲公司",
            input_files=[str(data_dir)],
            output_dir=str(artifact_dir),
        )

        result = run_workflow(
            state,
            state_path,
            checker=checker,
            client_factory=lambda: self.fail("已有数据模式不应创建注协客户端"),
        )

        self.assertEqual(result.status, "completed")
        self.assertEqual(len(checker.calls), 1)
        self.assertEqual(checker.calls[0]["data_dir"], data_dir)
        self.assertEqual(
            result.report_paths,
            [artifact_dir / "甲公司_关联方核查报告.xlsx"],
        )

    def test_名单核查模式只导出用户明确提供的名称(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_list_mode_"))
        data_dir = artifact_dir / "甲公司_注协原始导出"
        data_dir.mkdir()
        list_path = artifact_dir / "客户名单.txt"
        list_path.write_text("乙公司\n丙公司\n", encoding="utf-8-sig")
        exporter = FakeExporter(data_dir)
        checker = FakeChecker()
        state_path = artifact_dir / "task.json"
        state = TaskState(
            task_id="list-1",
            mode="list_check",
            audited_entity="甲公司",
            input_files=[str(list_path)],
            output_dir=str(artifact_dir),
        )

        result = run_workflow(
            state,
            state_path,
            client_factory=lambda: object(),
            exporter_factory=lambda _client: exporter,
            checker=checker,
        )

        self.assertEqual(result.status, "completed")
        self.assertEqual(exporter.start_calls, [["甲公司", "乙公司", "丙公司"]])
        self.assertEqual(checker.calls[0]["task_mode"], "list_check")
        self.assertEqual(
            set(result.report_paths),
            {
                data_dir,
                artifact_dir / "甲公司_关联方核查报告.xlsx",
            },
        )

    def test_主动发现模式只给被审计单位做完整维度导出(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_discovery_mode_"))
        data_dir = artifact_dir / "甲公司_注协原始导出"
        data_dir.mkdir()
        exporter = FakeExporter(data_dir)
        checker = FakeChecker()
        state_path = artifact_dir / "task.json"
        state = TaskState(
            task_id="discovery-1",
            mode="discovery",
            audited_entity="甲公司",
            output_dir=str(artifact_dir),
        )
        discovery_result = DiscoveryResult(
            status="completed",
            seed_name="甲公司",
            candidates={
                "乙公司": Candidate(
                    name="乙公司",
                    company_id="org-b",
                    depth=1,
                    parent_name="甲公司",
                    relation_type="对外投资",
                    ratio=60.0,
                    reasons=["对外投资持股比例 60%"],
                    paths=["甲公司 --对外投资 60%--> 乙公司"],
                )
            },
        )

        result = run_workflow(
            state,
            state_path,
            client_factory=lambda: object(),
            exporter_factory=lambda _client: exporter,
            discoverer=lambda *_args, **_kwargs: discovery_result,
            checker=checker,
        )

        self.assertEqual(result.status, "completed")
        self.assertEqual(exporter.start_calls, [["甲公司"]])
        self.assertEqual(result.candidate_count, 1)
        self.assertTrue(
            (artifact_dir / "甲公司_主动发现候选清单.xlsx").exists()
        )
        self.assertEqual(
            set(result.report_paths),
            {
                data_dir,
                artifact_dir / "甲公司_主动发现候选清单.xlsx",
                artifact_dir / "甲公司_关联方核查报告.xlsx",
            },
        )

    def test_默认在线导出直接使用带公司名的成果目录(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_default_export_dir_"))
        checker = FakeChecker()
        state_path = artifact_dir / "task.json"
        state = TaskState(
            task_id="default-export-1",
            mode="discovery",
            audited_entity="甲公司",
            output_dir=str(artifact_dir),
        )
        discovery_result = DiscoveryResult(
            status="completed",
            seed_name="甲公司",
            candidates={},
        )

        def build_exporter(_client, *, artifact_dir):
            artifact_dir.mkdir(parents=True)
            return FakeExporter(artifact_dir)

        try:
            with patch(
                "scripts.related_party_workflow.CicpaExporter",
                side_effect=build_exporter,
            ):
                result = run_workflow(
                    state,
                    state_path,
                    client_factory=lambda: object(),
                    exporter_factory=None,
                    discoverer=lambda *_args, **_kwargs: discovery_result,
                    checker=checker,
                )
        except TypeError as exc:
            self.fail("默认在线导出尚未接收成果目录：{}".format(exc))

        raw_export_dir = artifact_dir / "甲公司_注协原始导出"
        self.assertTrue(raw_export_dir.is_dir())
        self.assertIn(raw_export_dir, result.report_paths)

    def test_公司名称清理后用于报告文件名(self):
        long_name = "甲" * 79 + ".乙"
        cases = (
            (" 甲<乙>:公司. ", "甲_乙__公司_关联方核查报告.xlsx"),
            ("CON", "_CON_关联方核查报告.xlsx"),
            ("...", "未命名公司_关联方核查报告.xlsx"),
            (long_name, "甲" * 79 + "_关联方核查报告.xlsx"),
        )
        for index, (audited_entity, expected_name) in enumerate(cases):
            with self.subTest(audited_entity=audited_entity):
                artifact_dir = Path(
                    tempfile.mkdtemp(prefix="rpi_safe_company_name_")
                )
                data_dir = artifact_dir / "existing_files"
                data_dir.mkdir()
                state = TaskState(
                    task_id="safe-name-{}".format(index),
                    mode="existing_export",
                    audited_entity=audited_entity,
                    input_files=[str(data_dir)],
                    output_dir=str(artifact_dir),
                )

                result = run_workflow(
                    state,
                    artifact_dir / "task.json",
                    checker=FakeChecker(),
                )

                self.assertEqual(result.report_paths[0].name, expected_name)

    def test_同名报告存在时增加时间戳且不覆盖旧文件(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_report_collision_"))
        data_dir = artifact_dir / "existing_files"
        data_dir.mkdir()
        existing_report = artifact_dir / "甲公司_关联方核查报告.xlsx"
        existing_report.write_bytes(b"existing-report")
        state = TaskState(
            task_id="collision-1",
            mode="existing_export",
            audited_entity="甲公司",
            input_files=[str(data_dir)],
            output_dir=str(artifact_dir),
        )

        result = run_workflow(
            state,
            artifact_dir / "task.json",
            checker=FakeChecker(),
        )

        self.assertEqual(existing_report.read_bytes(), b"existing-report")
        self.assertNotEqual(result.report_paths[0], existing_report)
        self.assertTrue(
            result.report_paths[0].name.startswith(
                "甲公司_关联方核查报告_"
            )
        )

    def test_同一秒已有时间戳报告时继续编号且不覆盖(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_report_collision_"))
        data_dir = artifact_dir / "existing_files"
        data_dir.mkdir()
        original = artifact_dir / "甲公司_关联方核查报告.xlsx"
        timestamped = (
            artifact_dir / "甲公司_关联方核查报告_20260726_120000.xlsx"
        )
        original.write_bytes(b"original")
        timestamped.write_bytes(b"timestamped")
        state = TaskState(
            task_id="collision-2",
            mode="existing_export",
            audited_entity="甲公司",
            input_files=[str(data_dir)],
            output_dir=str(artifact_dir),
        )

        with patch("scripts.related_party_workflow.datetime") as mocked_datetime:
            mocked_datetime.now.return_value = datetime(2026, 7, 26, 12, 0, 0)
            result = run_workflow(
                state,
                artifact_dir / "task.json",
                checker=FakeChecker(),
            )

        self.assertEqual(original.read_bytes(), b"original")
        self.assertEqual(timestamped.read_bytes(), b"timestamped")
        self.assertEqual(
            result.report_paths[0].name,
            "甲公司_关联方核查报告_20260726_120000_2.xlsx",
        )

    def test_恢复等待中的导出不会再次上传名单(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_resume_export_"))
        data_dir = artifact_dir / "exported_files"
        data_dir.mkdir()
        list_path = artifact_dir / "客户名单.txt"
        list_path.write_text("乙公司\n", encoding="utf-8-sig")
        export_state_path = artifact_dir / "export-state.json"
        save_export_state(
            export_state_path,
            ExportState(
                batch_no="batch-existing",
                company_names=["甲公司", "乙公司"],
                created_at="2026-07-26 16:00",
                triggered_at="2026-07-26 16:00",
                staging_dir=str(artifact_dir),
                status="waiting",
            ),
        )
        exporter = FakeExporter(data_dir)
        checker = FakeChecker()
        state_path = artifact_dir / "task.json"
        state = TaskState(
            task_id="resume-1",
            mode="list_check",
            audited_entity="甲公司",
            input_files=[str(list_path)],
            output_dir=str(artifact_dir),
            export_state_path=str(export_state_path),
        )

        result = run_workflow(
            state,
            state_path,
            client_factory=lambda: object(),
            exporter_factory=lambda _client: exporter,
            checker=checker,
        )

        self.assertEqual(result.status, "completed")
        self.assertEqual(exporter.start_calls, [])
        self.assertEqual(exporter.wait_calls, ["batch-existing"])


class SkillDocumentationContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.skill_root = Path(__file__).resolve().parents[2]

    def _read_if_exists(self, relative_path):
        path = self.skill_root / relative_path
        return path.read_text(encoding="utf-8-sig") if path.exists() else ""

    def test_对话手册覆盖三个入口和非技术用户关键节点(self):
        skill_text = self._read_if_exists("SKILL.md")
        user_flow = self._read_if_exists("references/user-flow.md")
        combined = skill_text + "\n" + user_flow

        required_phrases = [
            "主动发现",
            "名单核查",
            "已有数据",
            "Windows 文件选择窗口",
            "安装前征得用户同意",
            "登录好了",
            "每分钟最多 15 次",
            "每小时最多 300 次",
            "候选超过 100 家",
            "只保存到当前 Windows 用户",
            "不等于最终关联方结论",
        ]
        for phrase in required_phrases:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, combined)

    def test_用户文案不暴露技术操作或要求手填绝对路径(self):
        skill_text = self._read_if_exists("SKILL.md")
        user_flow = self._read_if_exists("references/user-flow.md")
        combined = skill_text + "\n" + user_flow

        forbidden_phrases = [
            "F12",
            "开发者工具",
            "复制 Cookie",
            "复制Cookie",
            "手填绝对路径",
            "python3 ",
            "python ",
            "~/.claude/skills/",
            "~/.agents/skills/",
        ]
        for phrase in forbidden_phrases:
            with self.subTest(phrase=phrase):
                self.assertNotIn(phrase, combined)

    def test_目标运行时不引用外部_skill_路径(self):
        runtime_files = [
            path
            for path in (self.skill_root / "scripts").rglob("*.py")
            if "tests" not in path.parts
        ]
        runtime_text = "\n".join(
            path.read_text(encoding="utf-8-sig") for path in runtime_files
        )

        self.assertNotIn("cicpa-company-query/scripts", runtime_text)
        self.assertNotIn("workbuddy skills/cicpa-company-query", runtime_text)
        self.assertNotIn("workbuddy skills\\cicpa-company-query", runtime_text)

    def test_来源说明完整且不虚构独立许可证文件(self):
        sources_path = self.skill_root / "references" / "SOURCES.json"
        notice_path = self.skill_root / "NOTICE"

        self.assertTrue(sources_path.exists())
        self.assertTrue(notice_path.exists())
        sources = json.loads(sources_path.read_text(encoding="utf-8"))
        notice = notice_path.read_text(encoding="utf-8-sig")
        serialized = json.dumps(sources, ensure_ascii=False)
        combined = serialized + "\n" + notice

        for phrase in [
            "nigo",
            "https://github.com/nigo81/nigo-skills",
            "cicpa-company-query",
            "0657de949b42b5177a7e2a9a1f9c19b5deac926a",
            "MIT",
            "本地改造",
        ]:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, combined)
        self.assertFalse(sources["sources"][0]["license_evidence"]["license_file_verified"])

    def test_分发元数据和依赖文件存在且彼此一致(self):
        requirements_path = self.skill_root / "requirements.txt"
        metadata_path = self.skill_root / "agents" / "openai.yaml"

        self.assertTrue(requirements_path.exists())
        self.assertTrue(metadata_path.exists())
        requirements = requirements_path.read_text(encoding="utf-8")
        metadata = metadata_path.read_text(encoding="utf-8")

        self.assertIn("requests", requirements)
        self.assertIn("openpyxl", requirements)
        self.assertNotIn("playwright", requirements.lower())
        self.assertIn('display_name: "关联方识别与核查"', metadata)
        self.assertIn("$related-party-identification", metadata)

    def test_规则说明区分公开候选和有效核查证据(self):
        skill = self._read_if_exists("SKILL.md")
        readme = self._read_if_exists("README.md")
        rules = self._read_if_exists("references/rules.md")
        dimensions = self._read_if_exists("references/dimensions.md")
        cases = self._read_if_exists("references/cases.md")
        user_flow = self._read_if_exists("references/user-flow.md")
        combined = "\n".join(
            [skill, readme, rules, dimensions, cases, user_flow]
        )

        for phrase in [
            "候选不等于关联方",
            "红旗不等于硬关联",
            "公开客户关系",
            "公开供应商关系",
            "只作为候选",
            "不能单独形成风险证据",
            "初始 2 层",
            "最多 5 层",
            "只能从完整维度导出",
        ]:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, combined)

        for phrase in [
            "注协标记关联关系",
            "导出表内关联标注",
        ]:
            with self.subTest(phrase=phrase):
                self.assertNotIn(phrase, combined)

        for phrase in [
            "<公司名称>_关联方核查报告.xlsx",
            "<公司名称>_主动发现候选清单.xlsx",
            "<公司名称>_注协原始导出",
        ]:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, combined)

    def test_gitignore_不含搜索工具无法解析的末尾反斜杠(self):
        lines = (self.skill_root / ".gitignore").read_text(
            encoding="utf-8-sig"
        ).splitlines()
        invalid = [
            line
            for line in lines
            if line.strip()
            and not line.lstrip().startswith("#")
            and line.endswith("\\")
        ]

        self.assertEqual(invalid, [])


class BundleValidationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.skill_root = Path(__file__).resolve().parents[2]

    @staticmethod
    def _passing_tests(_root):
        return True, 0, "测试由当前用例代替"

    def test_校验器检查当前技能的分发契约(self):
        report = validate_bundle(
            self.skill_root,
            test_runner=self._passing_tests,
        )

        self.assertTrue(report.is_valid, report.errors)
        for check in [
            "required_files",
            "python_syntax",
            "unit_tests",
            "runtime_paths",
            "secret_leaks",
            "encoding",
            "metadata",
            "provenance",
            "dialog_entries",
        ]:
            with self.subTest(check=check):
                self.assertTrue(report.checks[check])

    def test_外部路径和模拟明文凭据会指出具体文件(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_bundle_bad_"))
        bundle = artifact_dir / "bundle"
        shutil.copytree(self.skill_root, bundle)
        bad_path = bundle / "scripts" / "bad_runtime.py"
        external_path = "~/" + ".claude/skills/" + "other/SKILL.md"
        simulated_secret = "cicpa_" + "token=" + "A" * 32
        bad_path.write_text(
            f'EXTERNAL = "{external_path}"\nSECRET = "{simulated_secret}"\n',
            encoding="utf-8-sig",
        )

        report = validate_bundle(bundle, test_runner=self._passing_tests)
        combined = "\n".join(report.errors)

        self.assertFalse(report.is_valid)
        self.assertFalse(report.checks["runtime_paths"])
        self.assertFalse(report.checks["secret_leaks"])
        self.assertIn("scripts/bad_runtime.py", combined.replace("\\", "/"))

    def test_未复制_notice_的临时副本会失败但不删除源文件(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_bundle_notice_"))
        bundle = artifact_dir / "bundle"

        def ignore_notice(directory, names):
            if Path(directory).resolve() == self.skill_root.resolve() and "NOTICE" in names:
                return {"NOTICE"}
            return set()

        shutil.copytree(self.skill_root, bundle, ignore=ignore_notice)
        report = validate_bundle(bundle, test_runner=self._passing_tests)

        self.assertFalse(report.checks["required_files"])
        self.assertIn("NOTICE", "\n".join(report.errors))
        self.assertTrue((self.skill_root / "NOTICE").exists())

    def test_秘密扫描不会读取技能目录外的旧_cookie_文件(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_bundle_scope_"))
        outside_secret = artifact_dir / ".cicpa_cookies.json"
        outside_secret.write_text(
            '{"cicpa_token":"' + "B" * 32 + '"}',
            encoding="utf-8",
        )
        bundle = artifact_dir / "bundle"
        shutil.copytree(self.skill_root, bundle)

        report = validate_bundle(bundle, test_runner=self._passing_tests)

        self.assertTrue(report.checks["secret_leaks"], report.errors)


if __name__ == "__main__":
    unittest.main()
