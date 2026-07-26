# 关联方识别与核查 — 原创编排层
# Copyright (C) 2026 CPA-Q (quanfanpro-code)
#
# 本文件是 related-party-identification 的原创编排层,采用 GNU Affero General
# Public License v3.0 (AGPL-3.0) 发布。
# 完整协议见项目根目录 LICENSE 文件(AGPL-3.0)。
#
# 本项目的 scripts/cicpa/ 目录包含改编自 nigo/nigo-skills(MIT) 和
# jackwener/OpenCLI(Apache-2.0) 的代码,分别保留原始许可证。
# 详见 NOTICE 和 references/SOURCES.json。
#from datetime import date
from pathlib import Path
import tempfile
import unittest

import openpyxl

from scripts.related_party_check import (
    Company,
    build_company,
    rule3_counterparty_profile,
    rule4_disclosed_gap,
    run_check,
)


def blank_basic_row(name="甲公司"):
    row = [None] * 23
    row[0] = name
    row[2] = "存续"
    row[3] = "张三"
    row[4] = "1000 万元"
    row[5] = "2020-01-01"
    row[15] = 10
    row[21] = "成都市测试路 1 号"
    row[22] = "软件开发与技术服务"
    return tuple(row)


def make_basic_export(directory):
    path = Path(directory) / "基础工商信息.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["字段{}".format(index) for index in range(1, 24)])
    target = list(blank_basic_row("甲公司"))
    counterparty = list(blank_basic_row("乙公司"))
    counterparty[3] = "李四"
    counterparty[5] = "2025-01-01"
    counterparty[15] = 0
    sheet.append(target)
    sheet.append(counterparty)
    workbook.save(path)
    return path


class RowBoundaryTests(unittest.TestCase):
    def test_最终受益人最短合法行不会越界(self):
        dimensions = {
            "ultimate_beneficiary": {
                "甲公司": [
                    ("序号", "甲公司", "字段3", "字段4", "字段5", "张三")
                ]
            }
        }

        company = build_company("甲公司", dimensions, blank_basic_row())

        self.assertEqual(company.beneficiaries, [("张三", "")])

    def test_变更记录缺少变更后字段时保留已取得内容(self):
        dimensions = {
            "change": {
                "甲公司": [
                    ("序号", "甲公司", "2026-01-01", "地址", "旧地址")
                ]
            }
        }

        company = build_company("甲公司", dimensions, blank_basic_row())

        self.assertEqual(
            company.changes,
            [("2026-01-01", "地址", "旧地址", "")],
        )


class CounterpartyProfileTests(unittest.TestCase):
    def test_成立时间按核查基准日动态计算二十四个月(self):
        target = Company(name="甲公司", business_scope="软件开发与技术服务")
        party = Company(
            name="乙公司",
            found_date="2025-01",
            insured=10,
            capital=10_000_000,
            business_scope="软件开发与技术服务",
        )

        hits = rule3_counterparty_profile(
            target,
            party,
            {"甲公司"},
            as_of_date=date(2026, 7, 1),
        )

        self.assertEqual(len(hits), 1)
        self.assertIn("成立不足24个月", hits[0][2])
        self.assertNotIn("2018", hits[0][2])

    def test_零参保使用专门红旗(self):
        target = Company(name="甲公司", business_scope="软件开发与技术服务")
        party = Company(
            name="乙公司",
            found_date="2010-01",
            insured=0,
            capital=10_000_000,
            business_scope="软件开发与技术服务",
        )

        hits = rule3_counterparty_profile(
            target,
            party,
            {"甲公司"},
            as_of_date=date(2026, 7, 1),
        )

        self.assertIn("参保人数0(空壳特征)", hits[0][2])
        self.assertNotIn("参保人数0(疑似空壳)", hits[0][2])


class DisclosedPartySemanticsTests(unittest.TestCase):
    def setUp(self):
        self.dimensions = {
            "customer": {
                "甲公司": [
                    (
                        "序号",
                        "甲公司",
                        "客户公司",
                        "",
                        "",
                        "",
                        "",
                        "乙公司",
                    )
                ]
            },
            "supplier": {},
        }

    def test_没有自报名单时只能称为注协标记关系(self):
        hits = rule4_disclosed_gap(
            self.dimensions,
            "甲公司",
            ["甲公司", "乙公司"],
            disclosed_parties=None,
        )

        self.assertEqual(hits[0][4], "注协标记关联关系")
        self.assertNotIn("未披露", hits[0][2])
        self.assertNotIn("未在用户自报名单", hits[0][2])

    def test_有自报名单时才计算披露差异(self):
        hits = rule4_disclosed_gap(
            self.dimensions,
            "甲公司",
            ["甲公司", "乙公司"],
            disclosed_parties={"丙公司"},
        )

        self.assertEqual(hits[0][4], "用户披露差异")
        self.assertIn("未在用户自报名单", hits[0][2])

    def test_标注公司不在基础工商表时仍保留为候选证据(self):
        hits = rule4_disclosed_gap(
            self.dimensions,
            "甲公司",
            ["甲公司"],
            disclosed_parties=None,
        )

        self.assertEqual(len(hits), 1)
        self.assertIn("乙公司", hits[0][2])


class RunCheckIntegrationTests(unittest.TestCase):
    def test_读取警告进入结果和数据质量工作表(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_check_report_"))
        data_dir = artifact_dir / "complete-dimensions_files"
        data_dir.mkdir()
        make_basic_export(data_dir)
        (data_dir / "主要人员（高管）.xlsx").write_bytes(b"not-an-xlsx")
        output_path = artifact_dir / "关联方核查报告.xlsx"

        result = run_check(
            data_dir=data_dir,
            target_names=["甲公司"],
            output_path=output_path,
            as_of_date=date(2026, 7, 1),
            task_mode="existing_export",
            scope_metadata={"初始穿透层数": 2, "候选上限": 100},
        )

        self.assertEqual(result.output_path, output_path)
        self.assertTrue(output_path.exists())
        self.assertTrue(
            any("主要人员（高管）.xlsx" in item["source"] for item in result.errors)
        )
        workbook = openpyxl.load_workbook(output_path, read_only=True)
        try:
            self.assertIn("数据质量与限制", workbook.sheetnames)
            self.assertIn("任务范围", workbook.sheetnames)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
