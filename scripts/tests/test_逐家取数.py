# -*- coding: utf-8 -*-
# 批次三：保底通道与编排修复的行为测试。
# 保底写补采子目录、模糊匹配留痕、部分覆盖留痕、基准日传递、候选上限可继续、降级提示归类、对手方列不猜位置。
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path
from unittest import mock

import openpyxl

from scripts import related_party_workflow as workflow
from scripts.cicpa.exporter import ExportError
from scripts.related_party_check import run_check
from scripts.related_party_workflow import TaskState, extract_seed_export_candidates, run_workflow
from scripts.逐家取数 import collect_by_company, fetch_company


class FakeMatch:
    def __init__(self, name, org_id):
        self.name = name
        self.org_id = org_id


class FakeClient:
    """逐家取数用的假注协客户端：默认搜索精确命中，接口返回最小有效载荷。"""

    def __init__(self, search_results=None, holders=None):
        self.search_results = search_results or {}
        self.holders = holders or {}

    def search_companies(self, name):
        if name in self.search_results:
            return self.search_results[name]
        return [FakeMatch(name, "id-" + name)]

    def request_json(self, method, url, params=None, cache_key=None):
        orgid = (params or {}).get("orgid", "")
        name = orgid[3:] if orgid.startswith("id-") else orgid
        if "find_company_basic_info" in url:
            return {"status_code": 0, "data": {
                "name": name, "legal_representative": "张三",
                "reg_capital": "1000万", "established_date": "2010-01-01",
                "corp_address": "某市高新区科技园1号",
                "operating_scope": "一般项目：信息技术咨询服务；软件开发；企业管理咨询。",
            }}
        if "stock_holder_newest" in url:
            return {"status_code": 0, "data": {"list": list(self.holders.get(name, [
                {"holder_name": "股东" + name, "held_ratio": "60%"},
            ]))}}
        if "enterprise_equity" in url:
            return {"status_code": 0, "data": {
                "controllers": [{"name": "张三"}],
                "final_beneficiaries": [{"name": "张三"}],
                "invests": [],
            }}
        return {"status_code": 0, "data": {}}


def write_name_list(path, names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["公司名称"])
    for name in names:
        ws.append([name])
    wb.save(path)
    wb.close()


class Test保底子目录(unittest.TestCase):
    def test_保底写入补采子目录_不覆盖主目录证据(self):
        with tempfile.TemporaryDirectory() as d:
            raw = Path(d) / "注协原始导出"
            raw.mkdir()
            sentinel = raw / "基础工商信息.xlsx"
            sentinel.write_bytes(b"existing-evidence")
            state = TaskState(task_id="t", mode="list_check", audited_entity="甲公司")
            result_dir = workflow._fallback_to_single_company(
                state, Path(d) / "task.json", FakeClient(), raw, ["甲公司"], "批量接口持续报错")
            self.assertEqual(sentinel.read_bytes(), b"existing-evidence", "主目录已有证据被保底覆盖")
            self.assertTrue((raw / "逐家取数补采" / "基础工商信息.xlsx").is_file(), "保底未写入逐家取数补采子目录")
            self.assertEqual(result_dir, raw)

    def test_run_check自动纳入补采子目录(self):
        with tempfile.TemporaryDirectory() as d:
            raw = Path(d) / "注协原始导出"
            collect_by_company(FakeClient(), ["甲公司", "乙公司"], raw / "逐家取数补采", fallback_reason="测试")
            with redirect_stdout(StringIO()):
                result = run_check(data_dir=raw, target_names=["甲公司"],
                                   output_path=Path(d) / "报告.xlsx", as_of_date="2026-10-07")
            self.assertIn("甲公司", result.companies)
            self.assertIn("乙公司", result.companies)


class Test模糊匹配留痕(unittest.TestCase):
    def test_单条模糊匹配照常采用并留痕(self):
        client = FakeClient(search_results={
            "甲测试有限公司": [FakeMatch("甲测试（集团）有限公司", "id-fuzzy")],
        })
        profile = fetch_company(client, "甲测试有限公司")
        self.assertEqual(profile.company_id, "id-fuzzy")
        self.assertTrue(
            any("模糊匹配" in w and "甲测试（集团）有限公司" in w for w in profile.warnings),
            "模糊匹配未留痕",
        )

    def test_精确匹配不留模糊警告(self):
        profile = fetch_company(FakeClient(), "甲测试有限公司")
        self.assertFalse(any("模糊匹配" in w for w in profile.warnings))


class Test部分覆盖留痕(unittest.TestCase):
    def test_维度未覆盖全部公司时记入warnings(self):
        client = FakeClient(holders={"乙公司": []})
        with tempfile.TemporaryDirectory() as d:
            _results, warnings = collect_by_company(
                client, ["甲公司", "乙公司"], Path(d), fallback_reason="测试")
            self.assertTrue(
                any("最新公示股东" in w and "乙公司" in w and "未覆盖" in w for w in warnings),
                "维度部分覆盖未留痕: " + repr(warnings),
            )
            note = open(Path(d) / "取数说明.json", "r", encoding="utf-8-sig").read()
            self.assertIn("未覆盖", note)


class Test基准日传递(unittest.TestCase):
    def _run_existing_export(self, state, captured):
        def fake_checker(**kwargs):
            captured.update(kwargs)
        with tempfile.TemporaryDirectory() as d:
            data_dir = Path(d) / "导出"
            data_dir.mkdir()
            state.input_files = [str(data_dir)]
            state.output_dir = d
            result = run_workflow(state, Path(d) / "task.json", checker=fake_checker)
            return result

    def test_run接受基准日并传到核查(self):
        captured = {}
        state = TaskState(task_id="t", mode="existing_export", audited_entity="甲公司",
                          as_of_date="2026-06-30")
        result = self._run_existing_export(state, captured)
        self.assertEqual(result.status, "completed")
        self.assertEqual(captured.get("as_of_date"), date(2026, 6, 30))

    def test_未给基准日时仍为运行当天(self):
        captured = {}
        state = TaskState(task_id="t", mode="existing_export", audited_entity="甲公司")
        result = self._run_existing_export(state, captured)
        self.assertEqual(result.status, "completed")
        self.assertEqual(captured.get("as_of_date"), date.today())


class Test候选上限可继续(unittest.TestCase):
    def _run_cli(self, extra_args):
        captured = {}

        def fake_run(state, state_path, **kwargs):
            captured["state"] = state
            return workflow.WorkflowResult(status="completed", message_zh="ok")

        with tempfile.TemporaryDirectory() as d:
            argv = ["run", "--mode", "list_check", "--audited-entity", "甲公司",
                    "--state-path", str(Path(d) / "task.json"), "--output-dir", d] + extra_args
            with mock.patch.object(workflow, "ensure_login", return_value={"status": "authenticated"}), \
                 mock.patch.object(workflow, "run_workflow", fake_run), \
                 redirect_stdout(StringIO()):
                workflow.main(argv)
        return captured.get("state")

    def test_run命令接受candidate_cap参数(self):
        state = self._run_cli(["--candidate-cap", "500"])
        self.assertIsNotNone(state)
        self.assertEqual(state.candidate_cap, 500)

    def test_run命令candidate_cap默认100(self):
        state = self._run_cli([])
        self.assertIsNotNone(state)
        self.assertEqual(state.candidate_cap, 100)

    def test_run命令接受as_of_date参数(self):
        state = self._run_cli(["--as-of-date", "2026-06-30"])
        self.assertIsNotNone(state)
        self.assertEqual(state.as_of_date, "2026-06-30")


class Test降级提示归类(unittest.TestCase):
    def test_保底降级提示不进候选发现警告(self):
        class FailingExporter:
            def start_export(self, names, path):
                raise ExportError("批量接口不可用")

        captured = {}

        def fake_checker(**kwargs):
            captured.update(kwargs)

        with tempfile.TemporaryDirectory() as d:
            names_file = Path(d) / "名单.xlsx"
            write_name_list(names_file, ["乙公司"])
            state = TaskState(task_id="t", mode="list_check", audited_entity="甲公司",
                              input_files=[str(names_file)], output_dir=d)
            with redirect_stdout(StringIO()):
                result = run_workflow(
                    state, Path(d) / "task.json",
                    client_factory=lambda: FakeClient(),
                    exporter_factory=lambda client: FailingExporter(),
                    checker=fake_checker,
                )
        self.assertEqual(result.status, "completed")
        channel = captured.get("channel_warnings") or []
        self.assertTrue(any("批量导出通道不可用" in w for w in channel),
                        "保底降级提示未走取数通道降级类别: " + repr(captured))
        self.assertFalse(any("批量导出通道不可用" in w for w in state.discovery_warnings),
                         "保底降级提示仍混在候选发现警告里")

    def test_run_check把通道降级写入limitations(self):
        with tempfile.TemporaryDirectory() as d:
            collect_by_company(FakeClient(), ["甲公司"], Path(d), fallback_reason="测试")
            with redirect_stdout(StringIO()):
                result = run_check(data_dir=d, target_names=["甲公司"],
                                   output_path=Path(d) / "报告.xlsx", as_of_date="2026-10-07",
                                   channel_warnings=["批量导出通道不可用，已改用逐家取数"])
            self.assertTrue(any(item["category"] == "取数通道降级" for item in result.limitations))


class Test缺失维度提示不误报(unittest.TestCase):
    BASIC_HEADER = ["企业名称", "公司ID", "登记状态", "法定代表人", "注册资本", "成立日期",
                    "所在省份", "所在城市", "所在区县", "电话", "网址", "邮箱",
                    "统一社会信用代码", "注册号", "组织机构代码", "参保人数", "企业类型",
                    "行业门类", "行业大类", "行业中类", "曾用名", "企业地址", "经营范围"]
    SHAREHOLDER_HEADER = ["序号", "公司名称", "股东名称", "股东类型", "持股比例", "认缴出资额", "是否机构"]

    def _basic(self, name):
        return [name, "id-" + name, "存续", "法人甲", "1000万", "2010-01-01", "省", "市", "区",
                "01012345678", "", "mail@corp.com", "credit" + name, "", "", "50", "有限责任公司",
                "门类", "大类", "中类", "", "某市高新区科技园1号", "一般项目：信息技术咨询服务。"]

    def _write(self, folder, filename, header, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(header)
        for row in rows:
            ws.append(row)
        wb.save(Path(folder) / filename)
        wb.close()

    def test_补采子目录缺少的维度不报为缺失(self):
        with tempfile.TemporaryDirectory() as d:
            self._write(d, "基础工商信息.xlsx", self.BASIC_HEADER, [self._basic("甲公司")])
            self._write(d, "股东信息.xlsx", self.SHAREHOLDER_HEADER,
                        [[1, "甲公司", "股东甲", "自然人", "60%", "10万", "否"]])
            sub = Path(d) / "逐家取数补采"
            sub.mkdir()
            self._write(sub, "基础工商信息.xlsx", self.BASIC_HEADER, [self._basic("甲公司")])
            with redirect_stdout(StringIO()):
                result = run_check(data_dir=d, target_names=["甲公司"],
                                   output_path=Path(d) / "报告.xlsx", as_of_date="2026-10-07")
            sources = [item["source"] for item in result.limitations if item["category"] == "缺失维度"]
            self.assertNotIn("股东信息.xlsx", sources, "主目录已提供的维度被误报为缺失：{}".format(sources))
            self.assertEqual(sources.count("商标.xlsx"), 1, "同一维度不应重复报缺失：{}".format(sources))

    def test_主目录与补采目录都没有的维度报一次缺失(self):
        with tempfile.TemporaryDirectory() as d:
            self._write(d, "基础工商信息.xlsx", self.BASIC_HEADER, [self._basic("甲公司")])
            sub = Path(d) / "逐家取数补采"
            sub.mkdir()
            self._write(sub, "基础工商信息.xlsx", self.BASIC_HEADER, [self._basic("甲公司")])
            with redirect_stdout(StringIO()):
                result = run_check(data_dir=d, target_names=["甲公司"],
                                   output_path=Path(d) / "报告.xlsx", as_of_date="2026-10-07")
            sources = [item["source"] for item in result.limitations if item["category"] == "缺失维度"]
            self.assertEqual(sources.count("股东信息.xlsx"), 1, "两个目录都没有的维度应报一次")


class Test最终受益人比例(unittest.TestCase):
    class RatioClient(FakeClient):
        def request_json(self, method, url, params=None, cache_key=None):
            payload = super().request_json(method, url, params=params, cache_key=cache_key)
            if "enterprise_equity" in url:
                payload["data"]["final_beneficiaries"] = [{"name": "王五", "czbl": "35.5"}]
            return payload

    def test_受益人比例写入最终受益人表(self):
        from scripts.逐家取数 import fetch_company, write_dimension_files
        with tempfile.TemporaryDirectory() as d:
            profile = fetch_company(self.RatioClient(), "甲公司")
            write_dimension_files([profile], Path(d))
            wb = openpyxl.load_workbook(Path(d) / "最终受益人.xlsx")
            try:
                rows = list(wb.active.iter_rows(min_row=2, values_only=True))
            finally:
                wb.close()
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0][5], "王五")
            self.assertEqual(rows[0][6], "35.5%", "最终受益人比例被丢掉了")


class Test对手方列不猜位置(unittest.TestCase):
    def _write_customer(self, folder, headers, row):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        ws.append(row)
        wb.save(Path(folder) / "客户.xlsx")
        wb.close()

    def test_表头无对手方列名时跳过并告警(self):
        with tempfile.TemporaryDirectory() as d:
            self._write_customer(d, ["序号", "公司名称", "c", "d", "e", "f", "g", "h"],
                                 [1, "甲公司", "", "", "", "", "", "乙公司"])
            warnings = []
            candidates = extract_seed_export_candidates(d, warnings=warnings)
            self.assertEqual(candidates, [], "表头未命中时不应猜第8列")
            self.assertTrue(any("客户.xlsx" in w for w in warnings))

    def test_表头命中别名时仍提取(self):
        with tempfile.TemporaryDirectory() as d:
            self._write_customer(d, ["序号", "公司名称", "客户名称"],
                                 [1, "甲公司", "乙公司"])
            warnings = []
            candidates = extract_seed_export_candidates(d, warnings=warnings)
            self.assertEqual([item.name for item in candidates], ["乙公司"])
            self.assertEqual(warnings, [])


if __name__ == "__main__":
    unittest.main()
