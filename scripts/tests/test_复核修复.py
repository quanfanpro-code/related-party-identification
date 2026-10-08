# -*- coding: utf-8 -*-
# 2026-10-08 复核发现修复的行为测试：历史股东痕迹、既客又供、指纹稀缺性、跨类印证、曾用名股权。
# 用真实引擎 + 真实 Excel 输入验证行为，不打桩。
import tempfile
import unittest
from pathlib import Path

import openpyxl

from scripts.related_party_check import HARD, LOW, MEDIUM, run_check


BASIC_HEADER = ["企业名称", "公司ID", "登记状态", "法定代表人", "注册资本", "成立日期",
                "所在省份", "所在城市", "所在区县", "电话", "网址", "邮箱",
                "统一社会信用代码", "注册号", "组织机构代码", "参保人数", "企业类型",
                "行业门类", "行业大类", "行业中类", "曾用名", "企业地址", "经营范围"]

_BASIC_COL = {"legal_person": 3, "capital": 4, "found": 5, "phone": 9, "web": 10,
              "email": 11, "insured": 15, "former": 20, "address": 21, "scope": 22}


def basic_row(name, idx, **kw):
    row = [name, f"id{idx}", "存续", f"法人{idx}", "1000万", "2010-01-01",
           "省", "市", "区", f"139000{idx:05d}", "", f"mail{idx}@corp{idx}.com",
           f"credit{idx}", "", "", "50", "有限责任公司",
           "门类", "大类", "中类", "", f"某市高新区科技园{idx}号",
           "一般项目：信息技术咨询服务；软件开发；企业管理咨询。"]
    for key, value in kw.items():
        row[_BASIC_COL[key]] = value
    return row


def write_xlsx(folder, filename, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(Path(folder) / filename)
    wb.close()


CHANGE_HEADER = ["序号", "公司名称", "变更日期", "变更项目", "变更前", "变更后"]
CUST_HEADER = ["序号", "公司名称", "c", "d", "e", "f", "g", "关联方名称"]
BENEF_HEADER = ["序号", "公司名称", "c", "d", "e", "最终受益人名", "受益比例"]
TRADEMARK_HEADER = ["序号", "公司名称", "商标名"]
INVEST_HEADER = ["序号", "公司名称", "被投资企业", "d", "出资比例"]


def run(folder, target="甲测试有限公司"):
    return run_check(data_dir=folder, target_names=target,
                     output_path=Path(folder) / "报告.xlsx", as_of_date="2026-10-07")


def hits_between(result, a, b, field=None):
    pair = {a, b}
    found = [h for h in result.hits if {h["company_a"], h["company_b"]} == pair]
    if field:
        found = [h for h in found if h["field"] == field]
    return found


class Test历史股东痕迹(unittest.TestCase):
    def test_变更记录投资人项出现对方公司_命中(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "变更记录.xlsx", CHANGE_HEADER,
                       [[1, "乙测试有限公司", "2024-01-01", "投资人变更",
                         "甲测试有限公司持股80%", "丙测试有限公司持股80%"]])
            result = run(d)
            hits = hits_between(result, "甲测试有限公司", "乙测试有限公司", "past_investor")
            self.assertTrue(hits, "历史股东变更未命中")
            self.assertEqual(hits[0]["level"], MEDIUM)
            self.assertTrue(hits[0]["sources"], "命中缺少原始位置")

    def test_变更记录投资人项无对方公司_不命中(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "变更记录.xlsx", CHANGE_HEADER,
                       [[1, "乙测试有限公司", "2024-01-01", "投资人变更",
                         "丁测试有限公司持股80%", "戊测试有限公司持股80%"]])
            result = run(d)
            self.assertEqual(hits_between(result, "甲测试有限公司", "乙测试有限公司", "past_investor"), [])


class Test既客又供(unittest.TestCase):
    def test_同时出现在客户与供应商表_命中(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "客户.xlsx", CUST_HEADER, [[1, "甲测试有限公司", "", "", "", "", "", "乙测试有限公司"]])
            write_xlsx(d, "供应商.xlsx", CUST_HEADER, [[1, "甲测试有限公司", "", "", "", "", "", "乙测试有限公司"]])
            result = run(d)
            hits = hits_between(result, "甲测试有限公司", "乙测试有限公司", "dual_role")
            self.assertTrue(hits, "既客又供未命中")
            self.assertEqual(hits[0]["level"], MEDIUM)

    def test_只出现在客户表_不命中(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "客户.xlsx", CUST_HEADER, [[1, "甲测试有限公司", "", "", "", "", "", "乙测试有限公司"]])
            result = run(d)
            self.assertEqual(hits_between(result, "甲测试有限公司", "乙测试有限公司", "dual_role"), [])


class Test指纹稀缺性(unittest.TestCase):
    def test_地址六家共用_降为中级并注明(self):
        with tempfile.TemporaryDirectory() as d:
            rows = [basic_row(f"共用地址测试{i}有限公司", i, address="北京市海淀区中关村大街27号")
                    for i in range(1, 7)]
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER, rows)
            result = run(d, target="共用地址测试1有限公司")
            hits = hits_between(result, "共用地址测试1有限公司", "共用地址测试2有限公司", "address")
            self.assertTrue(hits, "地址命中丢失")
            self.assertEqual(hits[0]["level"], MEDIUM, "六家共用地址不应判高")
            self.assertIn("共用", hits[0]["evidence"])

    def test_地址两家独享_仍为高级(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1, address="北京市海淀区中关村大街27号"),
                        basic_row("乙测试有限公司", 2, address="北京市海淀区中关村大街27号")])
            result = run(d)
            hits = hits_between(result, "甲测试有限公司", "乙测试有限公司", "address")
            self.assertTrue(hits)
            self.assertEqual(hits[0]["level"], HARD)

    def test_电话五家共用_降为中级并注明(self):
        with tempfile.TemporaryDirectory() as d:
            rows = [basic_row(f"共用电话测试{i}有限公司", i, phone="01012348888")
                    for i in range(1, 6)]
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER, rows)
            result = run(d, target="共用电话测试1有限公司")
            hits = hits_between(result, "共用电话测试1有限公司", "共用电话测试2有限公司", "phone")
            self.assertTrue(hits, "电话命中丢失")
            self.assertEqual(hits[0]["level"], MEDIUM)
            self.assertIn("共用", hits[0]["evidence"])

    def test_邮箱完全相同六家共用_降为中级并注明(self):
        with tempfile.TemporaryDirectory() as d:
            rows = [basic_row(f"共用邮箱测试{i}有限公司", i, email="kuaiji@tongyi.com")
                    for i in range(1, 7)]
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER, rows)
            result = run(d, target="共用邮箱测试1有限公司")
            hits = hits_between(result, "共用邮箱测试1有限公司", "共用邮箱测试2有限公司", "email")
            self.assertTrue(hits, "邮箱命中丢失")
            self.assertEqual(hits[0]["level"], MEDIUM, "六家共用同一邮箱不应判高")
            self.assertIn("共用", hits[0]["evidence"])

    def test_邮箱完全相同两家独享_仍为高级(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1, email="one@tongyi.com"),
                        basic_row("乙测试有限公司", 2, email="one@tongyi.com")])
            result = run(d)
            hits = hits_between(result, "甲测试有限公司", "乙测试有限公司", "email")
            self.assertTrue(hits)
            self.assertEqual(hits[0]["level"], HARD)


class Test跨类印证(unittest.TestCase):
    def test_三条中风险三个类别_建议程序升级(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1, phone="01012340001"),
                        basic_row("乙测试有限公司", 2, phone="01012340002")])
            write_xlsx(d, "最终受益人.xlsx", BENEF_HEADER,
                       [[1, "甲测试有限公司", "", "", "", "王五", "30%"],
                        [2, "乙测试有限公司", "", "", "", "王五", "40%"]])
            write_xlsx(d, "商标.xlsx", TRADEMARK_HEADER,
                       [[1, "甲测试有限公司", "同名牌"], [2, "乙测试有限公司", "同名牌"]])
            result = run(d)
            entry = next(s for s in result.summary
                         if {s["company_a"], s["company_b"]} == {"甲测试有限公司", "乙测试有限公司"})
            self.assertTrue(entry["is_related"].startswith("中"), "印证情形不应改判高风险")
            self.assertIn("函证", entry["suggestion"], "印证后建议程序应升级为高风险档")
            self.assertIn("印证", entry["evidence"])

    def test_两条中风险_不升级(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1, phone="01012340001"),
                        basic_row("乙测试有限公司", 2, phone="01012340002")])
            write_xlsx(d, "商标.xlsx", TRADEMARK_HEADER,
                       [[1, "甲测试有限公司", "同名牌"], [2, "乙测试有限公司", "同名牌"]])
            result = run(d)
            entry = next(s for s in result.summary
                         if {s["company_a"], s["company_b"]} == {"甲测试有限公司", "乙测试有限公司"})
            self.assertNotIn("函证", entry["suggestion"])


class Test曾用名股权(unittest.TestCase):
    def test_对外投资使用对方曾用名_命中(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, former="乙旧称有限公司")])
            write_xlsx(d, "对外投资（新）.xlsx", INVEST_HEADER,
                       [[1, "甲测试有限公司", "乙旧称有限公司", "", "60%"]])
            result = run(d)
            hits = hits_between(result, "甲测试有限公司", "乙测试有限公司", "invest")
            self.assertTrue(hits, "曾用名未接通股权线")


class Test复核优化20261008(unittest.TestCase):
    """2026-10-08 独立复核（逻辑与报告可读性）发现修复的行为测试：
    名单按风险排序、集中注册类线索收敛、取数范围外不计缺口。"""

    def overview_rows(self, path):
        wb = openpyxl.load_workbook(path)
        try:
            return [row for row in wb["核查结论"].iter_rows(min_row=4, values_only=True) if row[0]]
        finally:
            wb.close()

    def test_名单按风险从高到低排列(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER, [
                basic_row("甲测试有限公司", 1),
                basic_row("乙测试有限公司", 2, phone="13900000001"),
                basic_row("丙测试有限公司", 3),
            ])
            write_xlsx(d, "商标.xlsx", TRADEMARK_HEADER,
                       [[1, "甲测试有限公司", "共同牌"], [2, "丙测试有限公司", "共同牌"]])
            result = run(d)
            clues = [row for row in self.overview_rows(result.output_path)
                     if str(row[1] or "") in {"高", "中", "低"}]
            # 乙与被审计单位撞手机号是高风险，丙与被审计单位撞商标是中风险
            self.assertEqual([(row[0], row[1]) for row in clues],
                             [("乙测试有限公司", "高"), ("丙测试有限公司", "中")],
                             "名单应按风险从高到低排")
            wb = openpyxl.load_workbook(result.output_path)
            try:
                rows = [row for row in wb["疑似关联方复核底稿"].iter_rows(min_row=4, values_only=True) if row[0]]
            finally:
                wb.close()
            self.assertEqual([row[1] for row in rows], ["乙测试有限公司", "丙测试有限公司"],
                             "复核底稿同样按风险排")

    def test_全部命中来自高共用度指纹的公司收敛(self):
        with tempfile.TemporaryDirectory() as d:
            rows = [basic_row(f"共用地址测试{i}有限公司", i, address="北京市海淀区中关村大街27号")
                    for i in range(1, 7)]
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER, rows)
            result = run(d, target="共用地址测试1有限公司")
            self.assertEqual(len(result.summary), 5, "只与被审计单位比对：被审计单位对另外 5 家各一条")
            self.assertTrue(all(item["converged"] for item in result.summary))
            self.assertTrue(all(hit["level"] == MEDIUM for hit in result.hits), "收敛不改命中本身的等级")
            overview = self.overview_rows(result.output_path)
            clues = [row for row in overview if str(row[1] or "") in {"高", "中", "低"}]
            self.assertEqual(clues, [], "收敛公司不列入疑似关联方名单")
            folded = [row for row in overview if str(row[0]).startswith("另有")]
            self.assertEqual(len(folded), 1)
            self.assertIn("5 家", folded[0][0])
            wb = openpyxl.load_workbook(result.output_path)
            try:
                metrics = {row[0]: row[1] for row in wb["核查范围与数量说明"].iter_rows(min_row=4, values_only=True) if row[0]}
                summary_rows = [row for row in wb["疑似关联方复核底稿"].iter_rows(min_row=4, values_only=True) if row[0]]
            finally:
                wb.close()
            self.assertEqual(metrics.get("其中只有集中注册类线索的公司数"), 5)
            self.assertEqual(len(summary_rows), 5)
            self.assertTrue(all(str(row[4]).startswith("【收敛") for row in summary_rows),
                            "复核底稿摘要应带收敛标记")

    def test_取数范围外不计入缺口_范围内未取得计入(self):
        from scripts.related_party_check import Company
        from scripts.报告输出 import DIMENSIONS, OUT_OF_SCOPE_STATUS, coverage_rows
        filled = {"legal_person": "某人", "phones": {"13900000000"}, "emails": ["a@corp.com"],
                  "addresses": ["某市某路1号"], "capital": 1000000.0, "found_date": "2010-01-01",
                  "insured": 10, "business_scope": "软件开发以及信息技术咨询服务"}
        companies = {"甲": Company("甲", **filled), "乙": Company("乙", **filled)}
        # 除商标维度外，其余维度两家公司均有记录，把缺口来源隔离到商标一处。
        file_info = [
            {"key": key, "file": f"{key}.xlsx", "exists": True, "readable": True,
             "companies": ["甲", "乙"], "scope": ["甲"], "terminal": "completed",
             "created_at": "", "malformed": False}
            for key in DIMENSIONS
        ]
        next(item for item in file_info if item["key"] == "trademark").update(
            {"companies": [], "terminal": ""})
        rows, incomplete, _severe = coverage_rows(["甲", "乙"], companies, file_info, [], {"甲"})
        yi_tm = next(row for row in rows if row[0] == "乙" and row[1] == "商标")
        self.assertEqual(yi_tm[2], OUT_OF_SCOPE_STATUS, "乙不在取数范围内，该维度应记为范围外")
        self.assertNotIn("乙", incomplete, "范围外不计入缺口")
        jia_tm = next(row for row in rows if row[0] == "甲" and row[1] == "商标")
        self.assertEqual(jia_tm[2], "未见该企业记录，范围待核实", "被审计单位在范围内，未取得应记为真缺口")
        self.assertIn("甲", incomplete)


if __name__ == "__main__":
    unittest.main()
