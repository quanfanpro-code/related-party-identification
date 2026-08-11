# 关联方识别与核查 — 原创编排层
# Copyright (C) 2026 CPA-Q (quanfanpro-code)
#
# 本文件是 related-party-identification 的原创编排层,采用 GNU Affero General
# Public License v3.0 (AGPL-3.0) 发布。
# 完整协议见项目根目录 LICENSE 文件(AGPL-3.0)。
#
# 本项目的 scripts/cicpa/ 目录包含改编自 nigo/nigo-skills(MIT) 和
# jackwener/OpenCLI(Apache-2.0) 的代码,分别保留原始许可证。
# 详见 NOTICE 和 references/SOURCES.json。
from io import BytesIO
import json
from pathlib import Path
import re
import tempfile
import unittest
import zipfile

import openpyxl

import scripts.cicpa.exporter as exporter_module

from scripts.cicpa.exporter import (
    CicpaExporter,
    ExportState,
    ExportTaskNotCorrelated,
    ExportValidationError,
    load_export_state,
    save_export_state,
)


class FakeClient:
    def __init__(self, json_results=None, binary_result=b""):
        self.json_results = list(json_results or [])
        self.binary_result = binary_result
        self.json_calls = []
        self.binary_calls = []

    def request_json(self, method, url, **kwargs):
        self.json_calls.append((method, url, kwargs))
        return self.json_results.pop(0)

    def request_bytes(self, method, url, **kwargs):
        self.binary_calls.append((method, url, kwargs))
        return self.binary_result


class AdaptiveClient(FakeClient):
    def __init__(self):
        super().__init__()
        self.current_task_id = ""
        self.task_archives = {}

    def request_json(self, method, url, **kwargs):
        self.json_calls.append((method, url, kwargs))
        if url.endswith("/upload_local"):
            return {"status_code": 0, "data": {"batch_id": "batch-adaptive"}}
        if url.endswith("/enter_search"):
            return {"status_code": 0, "data": {"list": [{"name": "甲公司"}]}}
        if url.endswith("/get_dimension_class"):
            return required_dimension_payload()
        if url.endswith("/task_list"):
            tasks = []
            if kwargs.get("kind") == "poll" and self.current_task_id:
                tasks.append(
                    {
                        "task_id": self.current_task_id,
                        "batch_id": "batch-adaptive",
                        "type": "批量查询",
                        "status": 1,
                        "date": "2026-08-06 22:01",
                        "url": "/{}.zip".format(self.current_task_id),
                    }
                )
            return {"status_code": 0, "data": {"list": tasks}}
        if url.endswith("/enter_search_out_type"):
            dimensions = kwargs["json"]["dimensions"]
            self.current_task_id = "task-{}".format(len(self.task_archives) + 1)
            self.task_archives[self.current_task_id] = make_dimension_zip(
                ["甲公司"],
                [(code, True) for code in dimensions],
            )
            return {"status_code": 0, "data": {"task_id": self.current_task_id}}
        raise AssertionError("未处理的测试接口：{}".format(url))

    def request_bytes(self, method, url, **kwargs):
        self.binary_calls.append((method, url, kwargs))
        task_id = Path(url).stem
        return self.task_archives[task_id]


class ThreeFailuresThenSingleClient(FakeClient):
    def __init__(self):
        super().__init__()
        self.export_dimensions = []

    def request_json(self, method, url, **kwargs):
        self.json_calls.append((method, url, kwargs))
        if url.endswith("/task_list"):
            return {"status_code": 0, "data": {"list": []}}
        if url.endswith("/enter_search_out_type"):
            dimensions = list(kwargs["json"]["dimensions"])
            self.export_dimensions.append(dimensions)
            if len(self.export_dimensions) <= 3:
                raise exporter_module.ExportError("模拟批量接口失败")
            return {"status_code": 0, "data": {"task_id": "task-single"}}
        raise AssertionError("未处理的测试接口：{}".format(url))


def make_export_zip(company_names, artifact_dir):
    workbook_path = Path(artifact_dir) / "基础工商信息.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["企业名称", "统一社会信用代码"])
    for index, name in enumerate(company_names, 1):
        sheet.append([name, "91510000TEST{:06d}".format(index)])
    workbook.save(workbook_path)
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(workbook_path, arcname="基础工商信息.xlsx")
    return output.getvalue()


def make_dimension_zip(company_names, dimension_rows):
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for code, has_rows in dimension_rows:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["企业名称", "核查字段"])
            if has_rows:
                for name in company_names:
                    sheet.append([name, "有效数据"])
            workbook_bytes = BytesIO()
            workbook.save(workbook_bytes)
            archive.writestr(
                exporter_module.DIMENSION_NAMES[code] + ".xlsx",
                workbook_bytes.getvalue(),
            )
    return output.getvalue()


def make_wrong_dimension_zip(company_name, code):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["序号", "公司名称", "股东"])
    sheet.append([1, company_name, "股东甲"])
    original = BytesIO()
    workbook.save(original)
    corrected = BytesIO()
    with zipfile.ZipFile(BytesIO(original.getvalue()), "r") as input_zip:
        with zipfile.ZipFile(corrected, "w", zipfile.ZIP_DEFLATED) as output_zip:
            for item in input_zip.infolist():
                content = input_zip.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(
                        rb'<dimension ref="[^"]+"',
                        b'<dimension ref="A1:A2"',
                        content,
                        count=1,
                    )
                output_zip.writestr(item, content)
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            exporter_module.DIMENSION_NAMES[code] + ".xlsx",
            corrected.getvalue(),
        )
    return result.getvalue()


def make_statistics_only_zip(company_name, code, count):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["导出维度记录数统计"])
    sheet.append(["公司名称", "现用名", exporter_module.DIMENSION_NAMES[code]])
    sheet.append([company_name, company_name, count])
    workbook_bytes = BytesIO()
    workbook.save(workbook_bytes)
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("统计表.xlsx", workbook_bytes.getvalue())
    return result.getvalue()


def waiting_state(artifact_dir):
    return ExportState(
        batch_no="batch-new",
        company_names=["甲公司"],
        created_at="2026-07-26T15:00:00+08:00",
        triggered_at="2026-07-26 15:00",
        preexisting_task_ids=["task-old"],
        staging_dir=str(artifact_dir),
        status="waiting",
    )


def required_dimension_payload():
    return {
        "status_code": 0,
        "status_msg": "success",
        "data": [
            {
                "name": "关联方核查",
                "children": [
                    {
                        "second_dimension_code": code,
                        "second_dimension_name": name,
                    }
                    for code, name in exporter_module.REQUIRED_DIMENSIONS
                ],
            }
        ],
    }


class ExportFlowTests(unittest.TestCase):
    def test_固定二十维首轮按十加十分组(self):
        self.assertTrue(hasattr(exporter_module, "REQUIRED_DIMENSIONS"))
        dimensions = exporter_module.REQUIRED_DIMENSIONS
        codes = [code for code, _name in dimensions]

        self.assertEqual(len(dimensions), 20)
        self.assertEqual(len(set(codes)), 20)
        self.assertEqual(codes[:2], ["S0000002", "S0000006"])
        self.assertEqual(codes[-2:], ["S0000036", "S0000101"])
        self.assertEqual(
            [len(group) for group in exporter_module.initial_dimension_groups()],
            [10, 10],
        )

    def test_恢复时下一组跳过已完成和明确无数据维度(self):
        codes = ["S0000002", "S0000006", "S0000103"]
        state = ExportState(
            batch_no="batch-new",
            company_names=["甲公司"],
            created_at="2026-08-06 22:00",
            triggered_at="",
            required_dimensions=codes,
            dimension_results={
                "S0000002": "completed",
                "S0000006": "no_data",
            },
            pending_groups=[codes],
        )

        self.assertEqual(exporter_module.next_dimension_group(state), ["S0000103"])

    def test_启动时只发送第一批十维且只上传一次名单(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_first_batch_"))
        state_path = artifact_dir / "state.json"
        waits = []
        client = FakeClient(
            json_results=[
                {"status_code": 0, "data": {"batch_id": "batch-new"}},
                {"status_code": 0, "data": {"list": [{"name": "甲公司"}]}},
                required_dimension_payload(),
                {"status_code": 0, "data": {"list": []}},
                {"status_code": 0, "data": {"task_id": "task-first"}},
            ]
        )
        exporter = CicpaExporter(
            client,
            artifact_dir=artifact_dir,
            sleeper=waits.append,
            randint=lambda lower, _upper: lower,
            now=lambda: "2026-08-06 22:00",
        )

        state = exporter.start_export(["甲公司"], state_path)

        upload_calls = [
            call for call in client.json_calls if call[1].endswith("/upload_local")
        ]
        export_calls = [
            call
            for call in client.json_calls
            if call[1].endswith("/enter_search_out_type")
        ]
        expected_first = [
            code for code, _name in exporter_module.REQUIRED_DIMENSIONS[:10]
        ]
        self.assertEqual(len(upload_calls), 1)
        self.assertEqual(len(export_calls), 1)
        self.assertEqual(export_calls[0][2]["json"]["dimensions"], expected_first)
        self.assertEqual(state.active_group, expected_first)
        self.assertEqual([len(group) for group in state.pending_groups], [10])
        self.assertEqual(waits, [])

    def test_活动任务存在时不触发第二个下载任务(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_single_task_"))
        state = waiting_state(artifact_dir)
        state.active_group = ["S0000002"]
        state.task_id = "task-active"
        exporter = CicpaExporter(FakeClient(), staging_root=artifact_dir)

        triggered = exporter.trigger_next_group(state, artifact_dir / "state.json")

        self.assertFalse(triggered)
        self.assertEqual(exporter.client.json_calls, [])

    def test_批量接口连续失败三次后才改为逐个下载(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_windows_"))
        state_path = artifact_dir / "state.json"
        codes = [code for code, _name in exporter_module.REQUIRED_DIMENSIONS[:5]]
        state = waiting_state(artifact_dir)
        state.required_dimensions = codes
        state.active_group = codes
        state.task_id = "task-active"
        exporter = CicpaExporter(FakeClient(), staging_root=artifact_dir)

        for failure_number in (1, 2):
            exporter.handle_poll_window_failure(state, state_path)
            self.assertEqual(state.batch_failure_streak, failure_number)
            self.assertFalse(state.single_dimension_mode)
            self.assertEqual(state.pending_groups, [codes])
            state.active_group = state.pending_groups.pop(0)
            state.task_id = "task-active-{}".format(failure_number + 1)

        exporter.handle_poll_window_failure(state, state_path)

        self.assertEqual(state.batch_failure_streak, 3)
        self.assertTrue(state.single_dimension_mode)
        self.assertEqual(state.active_group, [])
        self.assertEqual(state.task_id, "")
        self.assertEqual(state.pending_groups, [[code] for code in codes])
        self.assertIn("连续失败 3 次", state.message_zh)

    def test_触发接口连续失败三次后下一次请求为单维度(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_trigger_fallback_"))
        state_path = artifact_dir / "state.json"
        codes = [code for code, _name in exporter_module.REQUIRED_DIMENSIONS[:2]]
        state = waiting_state(artifact_dir)
        state.required_dimensions = codes
        state.pending_groups = [codes]
        client = ThreeFailuresThenSingleClient()
        exporter = CicpaExporter(client, staging_root=artifact_dir)

        triggered = exporter.trigger_next_group(state, state_path)

        self.assertTrue(triggered)
        self.assertEqual([len(group) for group in client.export_dimensions], [2, 2, 2, 1])
        self.assertEqual(state.batch_failure_streak, 3)
        self.assertTrue(state.single_dimension_mode)
        self.assertEqual(state.active_group, [codes[0]])
        self.assertEqual(state.pending_groups, [[codes[1]]])

    def test_单维度第二个轮询窗口失败后暂停(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_single_pause_"))
        state_path = artifact_dir / "state.json"
        state = waiting_state(artifact_dir)
        state.required_dimensions = ["S0000002"]
        state.active_group = ["S0000002"]
        state.task_id = "task-active"
        exporter = CicpaExporter(FakeClient(), staging_root=artifact_dir)

        exporter.handle_poll_window_failure(state, state_path)
        exporter.handle_poll_window_failure(state, state_path)

        self.assertEqual(state.status, "paused")
        self.assertEqual(state.active_group, ["S0000002"])
        self.assertIn("基础工商信息", state.message_zh)

    def test_未完成轮询交给客户端按_poll_节奏等待(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_poll_wait_"))
        state = waiting_state(artifact_dir)
        state.active_group = ["S0000002"]
        state.task_id = "task-active"
        waits = []
        pending = {
            "task_id": "task-active",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 0,
            "date": "2026-07-26 15:01",
        }
        completed = dict(pending, status=1, url="/batch.zip")
        client = FakeClient(
            json_results=[
                {"status_code": 0, "data": {"list": []}},
                {"status_code": 0, "data": {"list": [pending]}},
                {"status_code": 0, "data": {"list": [completed]}},
            ]
        )
        exporter = CicpaExporter(
            client,
            staging_root=artifact_dir,
            sleeper=waits.append,
            randint=lambda lower, _upper: lower,
        )

        task = exporter.wait_for_task(
            state,
            artifact_dir / "state.json",
            max_polls=3,
        )

        self.assertEqual(task["status"], 1)
        self.assertEqual(waits, [])
        self.assertTrue(
            all(call[2]["kind"] == "poll" for call in client.json_calls)
        )

    def test_两批十维串行完成且成功批次之间人工等待(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_complete_"))
        state_path = artifact_dir / "state.json"
        waits = []
        client = AdaptiveClient()
        exporter = CicpaExporter(
            client,
            artifact_dir=artifact_dir,
            sleeper=waits.append,
            randint=lambda lower, _upper: lower,
            now=lambda: "2026-08-06 22:00",
        )

        state = exporter.start_export(["甲公司"], state_path)
        result_dir = exporter.run_to_completion(state, state_path, max_polls=1)

        upload_calls = [
            call for call in client.json_calls if call[1].endswith("/upload_local")
        ]
        export_calls = [
            call
            for call in client.json_calls
            if call[1].endswith("/enter_search_out_type")
        ]
        self.assertEqual(len(upload_calls), 1)
        self.assertEqual([len(call[2]["json"]["dimensions"]) for call in export_calls], [10, 10])
        self.assertEqual(waits, [15])
        self.assertEqual(result_dir, artifact_dir)
        self.assertEqual(state.status, "completed")
        self.assertEqual(set(state.dimension_results.values()), {"completed"})
        self.assertEqual(len(state.dimension_results), 20)
        self.assertEqual(len(list(artifact_dir.glob("*.xlsx"))), 21)

    def test_真实流程保存批次号并只下载对应任务(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_flow_"))
        state_path = artifact_dir / "export-state.json"
        zip_bytes = make_export_zip(["甲公司"], artifact_dir)
        delivery_dir = artifact_dir / "甲公司_注协原始导出"
        client = FakeClient(
            json_results=[
                {
                    "status_code": 0,
                    "status_msg": "success",
                    "data": {
                        "batch_id": "batch-new",
                        "all_num": 1,
                        "hit_num": 1,
                        "not_hit_num": 0,
                    },
                },
                {
                    "status_code": 0,
                    "status_msg": "success",
                    "data": {"total": 1, "list": [{"name": "甲公司"}]},
                },
                required_dimension_payload(),
                {
                    "status_code": 0,
                    "data": {
                        "list": [
                            {
                                "task_id": "task-old",
                                "type": "批量查询",
                                "status": 1,
                                "date": "2026-07-26 14:00",
                                "url": "/old.zip",
                            }
                        ]
                    },
                },
                {"status_code": 0, "status_msg": "success", "data": {}},
                {
                    "status_code": 0,
                    "data": {
                        "list": [
                            {
                                "task_id": "task-new",
                                "batch_id": "batch-new",
                                "type": "批量查询",
                                "name": "批量查询-甲公司",
                                "status": 1,
                                "date": "2026-07-26 15:01",
                                "url": "/download/task-new.zip",
                            },
                            {
                                "task_id": "task-old",
                                "type": "批量查询",
                                "status": 1,
                                "date": "2026-07-26 14:00",
                                "url": "/old.zip",
                            },
                        ]
                    },
                },
            ],
            binary_result=zip_bytes,
        )
        try:
            exporter = CicpaExporter(
                client,
                artifact_dir=delivery_dir,
                now=lambda: "2026-07-26 15:00",
                sleeper=lambda _seconds: None,
            )
        except TypeError as exc:
            self.fail("导出器尚未支持精确成果目录：{}".format(exc))

        state = exporter.start_export(["甲公司"], state_path)
        task = exporter.wait_for_task(state, state_path, max_polls=1)
        extract_dir = exporter.download_and_validate(state, task, state_path)

        reloaded = load_export_state(state_path)
        self.assertEqual(Path(state.staging_dir), delivery_dir)
        self.assertTrue(state.direct_delivery)
        self.assertEqual(reloaded.batch_no, "batch-new")
        self.assertEqual(reloaded.task_id, "task-new")
        self.assertEqual(reloaded.status, "completed")
        self.assertEqual(extract_dir, delivery_dir)
        self.assertTrue((extract_dir / "基础工商信息.xlsx").exists())
        self.assertTrue((delivery_dir / "complete-dimensions.zip").exists())
        self.assertTrue((delivery_dir / "upload-companies.xlsx").exists())
        self.assertFalse((delivery_dir / "complete-dimensions_files").exists())
        self.assertEqual(
            client.binary_calls[0][1],
            "https://zsk-cmis.cicpa.org.cn/download/task-new.zip",
        )
        poll_calls = [
            kwargs
            for _method, url, kwargs in client.json_calls
            if url.endswith("/download/task_list")
        ]
        self.assertEqual(poll_calls[-1]["kind"], "poll")

    def test_超时后绝不下载旧的最新任务(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_timeout_"))
        state = waiting_state(artifact_dir)
        client = FakeClient(
            json_results=[
                {
                    "status_code": 0,
                    "data": {
                        "list": [
                            {
                                "task_id": "task-old",
                                "type": "批量查询",
                                "status": 1,
                                "date": "2026-07-26 14:00",
                                "url": "/old.zip",
                            }
                        ]
                    },
                }
            ]
        )
        exporter = CicpaExporter(client, staging_root=artifact_dir)

        with self.assertRaises(ExportTaskNotCorrelated):
            exporter.wait_for_task(state, artifact_dir / "state.json", max_polls=1)

        self.assertEqual(client.binary_calls, [])

    def test_多个无法唯一对应的新任务会立即阻断(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_ambiguous_"))
        state = waiting_state(artifact_dir)
        tasks = [
            {
                "task_id": "task-a",
                "type": "批量查询",
                "status": 1,
                "date": "2026-07-26 15:01",
                "url": "/a.zip",
            },
            {
                "task_id": "task-b",
                "type": "批量查询",
                "status": 1,
                "date": "2026-07-26 15:02",
                "url": "/b.zip",
            },
        ]
        exporter = CicpaExporter(FakeClient(), staging_root=artifact_dir)

        with self.assertRaises(ExportTaskNotCorrelated):
            exporter.correlate_task(state, tasks)

    def test_闭环遇到多个新任务立即暂停而不拆组(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_ambiguous_pause_"))
        state_path = artifact_dir / "state.json"
        state = waiting_state(artifact_dir)
        state.required_dimensions = ["S0000002", "S0000006"]
        state.active_group = list(state.required_dimensions)
        state.task_id = ""
        tasks = [
            {
                "task_id": "task-a",
                "type": "批量查询",
                "status": 1,
                "date": "2026-07-26 15:01",
                "url": "/a.zip",
            },
            {
                "task_id": "task-b",
                "type": "批量查询",
                "status": 1,
                "date": "2026-07-26 15:02",
                "url": "/b.zip",
            },
        ]
        exporter = CicpaExporter(
            FakeClient(json_results=[{"status_code": 0, "data": {"list": tasks}}]),
            staging_root=artifact_dir,
        )

        with self.assertRaises(ExportTaskNotCorrelated):
            exporter.run_to_completion(state, state_path, max_polls=1)

        self.assertEqual(state.status, "paused")
        self.assertEqual(state.poll_windows, 0)
        self.assertEqual(state.pending_groups, [])


class ExportValidationTests(unittest.TestCase):
    def test_统计表明确为零时缺失工作簿记为无数据(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_explicit_zero_"))
        code = "S0000107"
        state = waiting_state(artifact_dir)
        state.required_dimensions = [code]
        state.active_group = [code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/zero.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=make_statistics_only_zip("甲公司", code, 0)),
            artifact_dir=artifact_dir,
        )

        completed = exporter.download_current_group(
            state,
            task,
            artifact_dir / "state.json",
        )

        self.assertEqual(completed, [code])
        self.assertEqual(state.dimension_results[code], "no_data")
        self.assertEqual(state.pending_groups, [])
        self.assertEqual(state.status, "completed")

    def test_统计表为正数但工作簿缺失时仍然重排(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_positive_missing_"))
        code = "S0000107"
        state = waiting_state(artifact_dir)
        state.required_dimensions = [code]
        state.active_group = [code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/positive.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=make_statistics_only_zip("甲公司", code, 1)),
            artifact_dir=artifact_dir,
        )

        completed = exporter.download_current_group(
            state,
            task,
            artifact_dir / "state.json",
        )

        self.assertEqual(completed, [])
        self.assertNotIn(code, state.dimension_results)
        self.assertEqual(state.pending_groups, [[code]])

    def test_恢复同一批次时跳过已完成维度(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_skip_completed_"))
        completed_code = "S0000002"
        pending_code = "S0000006"
        existing = artifact_dir / "基础工商信息.xlsx"
        existing.write_bytes(b"confirmed evidence")
        state = waiting_state(artifact_dir)
        state.required_dimensions = [completed_code, pending_code]
        state.dimension_results = {completed_code: "completed"}
        state.active_group = [completed_code, pending_code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/resume.zip",
        }
        exporter = CicpaExporter(
            FakeClient(
                binary_result=make_dimension_zip(
                    ["甲公司"],
                    [(completed_code, True), (pending_code, True)],
                )
            ),
            artifact_dir=artifact_dir,
        )

        newly_completed = exporter.download_current_group(
            state,
            task,
            artifact_dir / "state.json",
        )

        self.assertEqual(newly_completed, [pending_code])
        self.assertEqual(existing.read_bytes(), b"confirmed evidence")
        self.assertEqual(state.status, "completed")

    def test_工作簿范围错误时仍读取真实公司名称列(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_wrong_range_"))
        code = "S0000006"
        state = waiting_state(artifact_dir)
        state.required_dimensions = [code]
        state.active_group = [code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/wrong-range.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=make_wrong_dimension_zip("甲公司", code)),
            artifact_dir=artifact_dir,
        )

        completed = exporter.download_current_group(
            state,
            task,
            artifact_dir / "state.json",
        )

        self.assertEqual(completed, [code])
        self.assertEqual(state.dimension_results[code], "completed")

    def test_恢复时不覆盖没有写入历史的孤立批次证据(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_orphan_"))
        batch_root = artifact_dir / "dimension-batches"
        orphan_dir = batch_root / "batch-001_files"
        orphan_dir.mkdir(parents=True)
        orphan_archive = batch_root / "batch-001.zip"
        orphan_archive.write_bytes(b"orphan evidence")
        (orphan_dir / "旧证据.txt").write_text("保留", encoding="utf-8-sig")
        code = "S0000002"
        state = waiting_state(artifact_dir)
        state.required_dimensions = [code]
        state.active_group = [code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/new.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=make_dimension_zip(["甲公司"], [(code, True)])),
            artifact_dir=artifact_dir,
        )

        exporter.download_current_group(state, task, artifact_dir / "state.json")

        self.assertEqual(orphan_archive.read_bytes(), b"orphan evidence")
        self.assertEqual(
            (orphan_dir / "旧证据.txt").read_text(encoding="utf-8-sig"),
            "保留",
        )
        self.assertTrue((batch_root / "batch-002.zip").exists())

    def test_分批工作簿企业范围不符时状态同步暂停(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_batch_mismatch_"))
        code = "S0000002"
        state_path = artifact_dir / "state.json"
        state = waiting_state(artifact_dir)
        state.required_dimensions = [code]
        state.active_group = [code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/wrong-company.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=make_dimension_zip(["乙公司"], [(code, True)])),
            artifact_dir=artifact_dir,
        )

        with self.assertRaises(ExportValidationError):
            exporter.download_current_group(state, task, state_path)

        reloaded = load_export_state(state_path)
        self.assertEqual(reloaded.status, "paused")
        self.assertIn("企业集合", reloaded.message_zh)

    def test_部分_zip_保留成功和空数据并只重排缺失维度(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_partial_"))
        state_path = artifact_dir / "state.json"
        codes = [code for code, _name in exporter_module.REQUIRED_DIMENSIONS[:3]]
        state = waiting_state(artifact_dir)
        state.required_dimensions = codes
        state.active_group = codes
        state.task_id = "task-new"
        state.batch_failure_streak = 2
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/partial.zip",
        }
        client = FakeClient(
            binary_result=make_dimension_zip(
                ["甲公司"],
                [(codes[0], True), (codes[1], False)],
            )
        )
        exporter = CicpaExporter(client, artifact_dir=artifact_dir)

        completed = exporter.download_current_group(state, task, state_path)

        self.assertEqual(completed, codes[:2])
        self.assertEqual(state.dimension_results[codes[0]], "completed")
        self.assertEqual(state.dimension_results[codes[1]], "no_data")
        self.assertNotIn(codes[2], state.dimension_results)
        self.assertEqual(state.pending_groups[0], [codes[2]])
        self.assertEqual(state.batch_failure_streak, 0)
        self.assertTrue((artifact_dir / "基础工商信息.xlsx").exists())
        self.assertTrue((artifact_dir / "股东信息.xlsx").exists())
        self.assertTrue((artifact_dir / "dimension-batches" / "batch-001.zip").exists())

    def test_同名不同哈希文件暂停且不覆盖旧证据(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_conflict_"))
        state_path = artifact_dir / "state.json"
        code = "S0000002"
        existing = artifact_dir / "基础工商信息.xlsx"
        existing.write_bytes(b"old evidence")
        state = waiting_state(artifact_dir)
        state.required_dimensions = [code]
        state.active_group = [code]
        state.task_id = "task-new"
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "url": "/conflict.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=make_dimension_zip(["甲公司"], [(code, True)])),
            artifact_dir=artifact_dir,
        )

        with self.assertRaises(ExportValidationError):
            exporter.download_current_group(state, task, state_path)

        self.assertEqual(existing.read_bytes(), b"old evidence")
        self.assertEqual(state.status, "paused")
        self.assertTrue(
            (artifact_dir / "dimension-batches" / "batch-001_files" / existing.name).exists()
        )

    def test_下载文件中的公司不一致时拒绝完成(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_mismatch_"))
        state_path = artifact_dir / "state.json"
        state = waiting_state(artifact_dir)
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "date": "2026-07-26 15:01",
            "url": "/wrong.zip",
        }
        client = FakeClient(binary_result=make_export_zip(["乙公司"], artifact_dir))
        exporter = CicpaExporter(client, staging_root=artifact_dir)

        with self.assertRaises(ExportValidationError):
            exporter.download_and_validate(state, task, state_path)

        self.assertNotEqual(state.status, "completed")

    def test_zip_路径穿越在写文件前被拒绝(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_zip_slip_"))
        output = BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("../outside.txt", "unsafe")
        state = waiting_state(artifact_dir)
        task = {
            "task_id": "task-new",
            "batch_id": "batch-new",
            "type": "批量查询",
            "status": 1,
            "date": "2026-07-26 15:01",
            "url": "/unsafe.zip",
        }
        exporter = CicpaExporter(
            FakeClient(binary_result=output.getvalue()),
            staging_root=artifact_dir,
        )

        with self.assertRaises(ExportValidationError):
            exporter.download_and_validate(state, task, artifact_dir / "state.json")

        self.assertFalse((artifact_dir.parent / "outside.txt").exists())

    def test_状态文件不带_bom_且不接受认证秘密字段(self):
        artifact_dir = Path(tempfile.mkdtemp(prefix="rpi_export_state_"))
        state_path = artifact_dir / "state.json"
        state = waiting_state(artifact_dir)

        save_export_state(state_path, state)
        raw = state_path.read_bytes()
        payload = json.loads(raw.decode("utf-8"))

        self.assertFalse(raw.startswith(b"\xef\xbb\xbf"))
        self.assertNotIn("cookie", json.dumps(payload, ensure_ascii=False).lower())
        self.assertNotIn("password", json.dumps(payload, ensure_ascii=False).lower())
        self.assertEqual(load_export_state(state_path), state)


if __name__ == "__main__":
    unittest.main()
