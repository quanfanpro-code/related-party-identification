# 关联方报告的业务验收：用真实 Excel 输入和真实核查引擎检查输出。
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path
import tempfile
import unittest
import json
from unittest.mock import patch

import openpyxl

from scripts.discovery import Candidate, DiscoveryResult
from scripts.related_party_check import run_check
from scripts.related_party_workflow import TaskState, run_workflow, extract_seed_export_candidates
from scripts.cicpa.exporter import ExportState, ExportError, save_export_state
from scripts.tests.test_related_party_check import blank_basic_row


def write_basic(folder, names, *, blank_line=False, shared_phone=True):
    folder.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工商原表"
    ws.append([f"字段{i}" for i in range(1, 24)])
    for index, name in enumerate(names):
        if blank_line and index == 1:
            ws.append([None] * 23)
        row = list(blank_basic_row(name))
        row[3] = ["张三", "李四", "王五"][index % 3]
        row[9] = "028-87654321" if shared_phone else f"1380000000{index}"
        row[21] = ["成都一环路111号", "北京长安街222号", "杭州余杭路333号"][index % 3]
        row[22] = "软件开发以及信息技术咨询服务"
        ws.append(row)
    wb.save(folder / "基础工商信息.xlsx")
    wb.close()


class LocalExporter:
    """仅替代联网导出，真实写入收到的公司名单，其余走生产流程。"""
    def __init__(self, root, fail_candidates=False):
        self.root = root
        self.names = []
        self.fail_candidates = fail_candidates

    def start_export(self, company_names, state_path):
        self.names.append(list(company_names))
        folder = self.root / ("被审计单位" if company_names == ["甲公司"] else "候选公司")
        state = ExportState(
            batch_no=str(len(self.names)), company_names=list(company_names),
            created_at="2026-10-07 20:00", triggered_at="2026-10-07 20:00",
            staging_dir=str(folder), required_dimensions=["S0000002"],
        )
        save_export_state(state_path, state)
        return state

    def run_to_completion(self, state, state_path):
        if self.fail_candidates and state.company_names != ["甲公司"]:
            raise ExportError("模拟候选下载中断")
        write_basic(Path(state.staging_dir), state.company_names)
        state.status = "completed"
        state.extract_dir = state.staging_dir
        state.dimension_results = {"S0000002": "completed"}
        save_export_state(state_path, state)
        return Path(state.extract_dir)


def candidate_result():
    return DiscoveryResult(status="completed", seed_name="甲公司", candidates={
        "乙公司": Candidate(name="乙公司", parent_name="甲公司", depth=1,
            relation_type="公开供应商关系", reasons=["公开供应商关系"],
            paths=["甲公司 --公开供应商关系--> 乙公司"]),
    })


def write_dimension(folder, filename, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "原始记录"
    ws.append([f"字段{i}" for i in range(1, max(map(len, rows)) + 1)])
    for row in rows:
        ws.append(row)
    wb.save(folder / filename)
    wb.close()


class ReportAcceptanceTests(unittest.TestCase):
    def test_公众号命中追溯名称列而非编号列(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        write_dimension(self.root, "微信公众号.xlsx", [[1, "甲公司", "账号甲", "同名公众号"], [2, "乙公司", "账号乙", "同名公众号"]])
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        hit = next(item for item in result.hits if item["field"] == "wechat")
        self.assertEqual({source["cell"] for source in hit["sources"]}, {"D2", "D3"})
        self.assertTrue(all(source["value"] == "同名公众号" for source in hit["sources"]))

    def test_损坏的范围说明与候选记录进入缺口(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        (self.root / "取数说明.json").write_text('{"dimension_results": []}', encoding="utf-8")
        (self.root / "候选发现记录.json").write_text('[]', encoding="utf-8")
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        self.assertEqual({item["category"] for item in result.errors}, {"取数说明读取失败", "候选发现记录读取失败"})

    def setUp(self):
        # 保留测试产物供复核，不删除文件。
        self.root = Path(tempfile.mkdtemp(prefix="rpi_report_acceptance_"))

    def run_silent(self, **kwargs):
        with redirect_stdout(StringIO()):
            return run_check(as_of_date=date(2026, 10, 7), **kwargs)

    def test_发现候选后确实取数比对且只交付一个报告(self):
        exporter = LocalExporter(self.root / "原始导出")
        state = TaskState(task_id="discovery", mode="discovery", audited_entity="甲公司",
                          output_dir=str(self.root))
        with redirect_stdout(StringIO()):
            result = run_workflow(state, self.root / "task.json", client_factory=lambda: object(),
                exporter_factory=lambda _: exporter, discoverer=lambda *a, **k: candidate_result())
        self.assertEqual(exporter.names, [["甲公司"], ["乙公司"]])
        reports = list(self.root.glob("*.xlsx"))
        self.assertEqual(len(reports), 1)
        wb = openpyxl.load_workbook(reports[0])
        try:
            rows = list(wb["关系核查汇总"].values)
            self.assertTrue(any("甲公司" in row and "乙公司" in row for row in rows))
            self.assertTrue(any("公开供应商关系" in str(row) for row in wb["核查对象与来源"].values))
        finally:
            wb.close()
        self.assertEqual(result.status, "completed")

    def test_候选批次恢复不重复目标导出或丢失来源(self):
        exporter = LocalExporter(self.root / "原始导出", fail_candidates=True)
        state = TaskState(task_id="resume", mode="discovery", audited_entity="甲公司", output_dir=str(self.root))
        with redirect_stdout(StringIO()):
            result = run_workflow(state, self.root / "task.json", client_factory=lambda: object(),
                exporter_factory=lambda _: exporter, discoverer=lambda *a, **k: candidate_result())
        self.assertNotEqual(result.status, "completed")
        exporter.fail_candidates = False
        with redirect_stdout(StringIO()):
            result = run_workflow(state, self.root / "task.json", client_factory=lambda: object(),
                exporter_factory=lambda _: exporter,
                discoverer=lambda *a, **k: self.fail("已有候选来源必须从状态恢复"))
        self.assertEqual(result.status, "completed")
        self.assertEqual(exporter.names, [["甲公司"], ["乙公司"]])

    def test_命中精确定位真实行号且原文可回读(self):
        write_basic(self.root, ["甲公司", "乙公司"], blank_line=True)
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        phone = next(hit for hit in result.hits if hit["field"] == "phone")
        sources = phone.get("sources", [])
        self.assertEqual({(s["company"], s["cell"]) for s in sources}, {("甲公司", "J2"), ("乙公司", "J4")})
        for source in sources:
            wb = openpyxl.load_workbook(source["file"], data_only=True)
            try:
                self.assertEqual(wb[source["sheet"]][source["cell"]].value, source["value"])
            finally:
                wb.close()

    def test_七表导航统计和缺口不混淆(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        wb = openpyxl.load_workbook(result.output_path)
        try:
            self.assertEqual(wb.sheetnames, ["目录", "核查概览", "核查对象与来源", "关系核查汇总", "证据明细", "数据覆盖与缺口", "任务说明"])
            for ws in wb.worksheets[1:]:
                self.assertIsNotNone(ws["A1"].hyperlink)
            overview = {row[0]: row[1] for row in wb["核查概览"].iter_rows(min_row=4, values_only=True) if row[0]}
            self.assertEqual(overview["实际尝试比对公司对数"], 1)
            self.assertEqual(overview["命中公司对数"], 0)
            self.assertGreater(overview["存在数据缺口公司数"], 0)
            for name in ("核查对象与来源", "关系核查汇总", "证据明细", "数据覆盖与缺口"):
                self.assertTrue(wb[name].auto_filter.ref)
                self.assertTrue(wb[name].freeze_panes)
        finally:
            wb.close()

    def test_七类命中均有真实来源且同一公司对只计一次(self):
        write_basic(self.root, ["甲公司", "乙公司"])
        write_dimension(self.root, "主要人员（高管）.xlsx", [[1, "甲公司", "王六", "董事"], [2, "乙公司", "王六", "经理"]])
        write_dimension(self.root, "对外投资（新）.xlsx", [[1, "甲公司", "乙公司", "", "60%"]])
        write_dimension(self.root, "法定代表人变更.xlsx", [["甲公司", "李四", "张三"]])
        write_dimension(self.root, "股权质押.xlsx", [[1, "甲公司", "", "甲公司", "", "乙公司"]])
        write_dimension(self.root, "商标.xlsx", [[1, "甲公司", "测试商标"], [2, "乙公司", "测试商标"]])
        wb = openpyxl.load_workbook(self.root / "基础工商信息.xlsx")
        wb.active["P3"] = 0
        wb.save(self.root / "基础工商信息.xlsx")
        wb.close()
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        self.assertEqual({hit["dimension"] for hit in result.hits}, {"工商指纹重合", "关键人员重合", "客商异常画像", "股权控制穿透", "历史关联痕迹", "担保资金链", "无形资产共用"})
        self.assertEqual(len(result.summary), 1)
        for hit in result.hits:
            self.assertTrue(hit["sources"], hit["field"])
            for source in hit["sources"]:
                wb = openpyxl.load_workbook(source["file"], data_only=True)
                self.assertEqual(wb[source["sheet"]][source["cell"]].value, source["value"])
                wb.close()
        again = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "再生成.xlsx")
        self.assertEqual({h["evidence_id"] for h in result.hits}, {h["evidence_id"] for h in again.hits})

    def test_缺少经营范围仅作为数据缺口(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        wb = openpyxl.load_workbook(self.root / "基础工商信息.xlsx")
        wb.active["W3"] = None
        wb.save(self.root / "基础工商信息.xlsx")
        wb.close()
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        self.assertEqual(result.hits, [])

    def test_明确无数据折叠显示与静默缺失逐家显示(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        (self.root / "取数说明.json").write_text(json.dumps({"company_names": ["甲公司", "乙公司"],
            "dimension_results": {"S0000041": "no_data"}}, ensure_ascii=False), encoding="utf-8")
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        wb = openpyxl.load_workbook(result.output_path)
        rows = list(wb["数据覆盖与缺口"].values)
        # 全体公司都是"明确无数据"的维度折叠成一条（统计表确认无数据），与静默缺失区分开；
        # 静默缺失（未取得）是缺口，仍逐家列出。
        folded = [row for row in rows if row[:3] == ("全部核查对象", "商标", "明确无数据")]
        self.assertEqual(len(folded), 1)
        self.assertIn("2 家", folded[0][4])
        self.assertFalse(any(row[1] == "商标" and row[0] == "甲公司" for row in rows))
        self.assertTrue(any(row[:3] == ("甲公司", "软件著作权", "未取得") for row in rows))
        wb.close()

    def test_规则失败在概览与对象记录中可见(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        with patch("scripts.related_party_check.rule1_fingerprint", side_effect=ValueError("模拟规则异常")):
            result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        self.assertTrue(any("模拟规则异常" in item["message"] for item in result.errors))
        wb = openpyxl.load_workbook(result.output_path)
        self.assertTrue(any("规则执行失败" in str(row) for row in wb["核查对象与来源"].values))
        wb.close()

    def test_非审计对象之间命中不推定与审计对象关联(self):
        write_basic(self.root, ["甲公司", "乙公司", "丙公司"], shared_phone=False)
        write_dimension(self.root, "商标.xlsx", [[1, "乙公司", "共用品牌"], [2, "丙公司", "共用品牌"]])
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        self.assertEqual(len(result.summary), 1)
        self.assertIn("不推定", result.summary[0]["relation_type"])

    def test_失去基础资料的候选仍出现在对象名单(self):
        write_basic(self.root, ["甲公司"])
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx",
            object_records=[{"name": "乙公司", "relation_type": "公开客户关系", "reasons": ["公开客户关系"]}])
        wb = openpyxl.load_workbook(result.output_path)
        rows = [row for row in wb["核查对象与来源"].values if row[0] == "乙公司"]
        self.assertEqual(len(rows), 1)
        self.assertIn("未纳入比对", rows[0][7])
        self.assertEqual(result.hits, [])
        wb.close()

    def test_源值以等号开头仍按原始文本保存(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        write_dimension(self.root, "商标.xlsx", [[1, "甲公司", "=1+1"], [2, "乙公司", "=1+1"]])
        source = self.root / "商标.xlsx"
        wb = openpyxl.load_workbook(source)
        for address in ("C2", "C3"):
            wb.active[address].data_type = "s"
        wb.save(source)
        wb.close()
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        wb = openpyxl.load_workbook(result.output_path)
        cells = [c for row in wb["证据明细"] for c in row if c.value == "=1+1"]
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == "s" for c in cells))
        wb.close()

    def test_汇总证据往返用编号定位而不是固定行号(self):
        write_basic(self.root, ["甲公司", "乙公司"])
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        wb = openpyxl.load_workbook(result.output_path)
        try:
            location = wb["关系核查汇总"]["I4"].hyperlink.location
            self.assertIn(location, wb.defined_names)
            self.assertIn("MATCH(", wb.defined_names[location].attr_text)
            back = wb["证据明细"]["K4"].hyperlink.location
            self.assertIn(back, wb.defined_names)
            self.assertIn("MATCH(", wb.defined_names[back].attr_text)
        finally:
            wb.close()

    def test_质押证据不混入同一出质人的其他记录(self):
        write_basic(self.root, ["甲公司", "乙公司"], shared_phone=False)
        write_dimension(self.root, "股权质押.xlsx", [
            [1, "甲公司", "", "甲公司", "", "乙公司"],
            [2, "甲公司", "", "甲公司", "", "无关公司"],
        ])
        result = self.run_silent(data_dir=self.root, target_names=["甲公司"], output_path=self.root / "报告.xlsx")
        hit = next(h for h in result.hits if h["field"] == "pledge")
        self.assertEqual({s["cell"] for s in hit["sources"] if Path(s["file"]).name == "股权质押.xlsx"}, {"D2", "F2"})

    def test_客商候选保留两种来源及实际单元格(self):
        # 表头必须带"关联方名称"才会提取（不再按第 8 列猜位置）；本用例验证来源单元格定位准确。
        for name in ("客户.xlsx", "供应商.xlsx"):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "原始记录"
            ws.append(["序号", "公司名称", "公告时间", "金额", "占比（%）", "与本公司关系", "货币代码", "关联方名称"])
            ws.append([1, "甲公司", "2026-01-01", 20, 1, "", "CNY", "乙公司"])
            wb.save(self.root / name)
            wb.close()
        candidates = extract_seed_export_candidates(self.root)
        self.assertEqual({item.relation_type for item in candidates}, {"公开客户关系", "公开供应商关系"})
        self.assertTrue(all(getattr(item, "sources", None) for item in candidates))
        self.assertTrue(all(item.sources[0]["cell"] == "H2" for item in candidates))


if __name__ == "__main__":
    unittest.main()
