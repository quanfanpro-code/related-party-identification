# 报告可读性（复核发现修复批次四第 21-26 项）的行为测试。
# 用真实引擎 run_check 生成八表报告到临时目录，再用 openpyxl 回读断言；
# 涉及"取数通道降级"提示的用例直接调用 write_report（该类别由保底通道批次产生）。
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
import json
import tempfile
import unittest

import openpyxl

from scripts.related_party_check import run_check, write_report
from scripts.cicpa.exporter import DIMENSION_NAMES


BASIC_HEADER = ["企业名称", "公司ID", "登记状态", "法定代表人", "注册资本", "成立日期",
                "所在省份", "所在城市", "所在区县", "电话", "网址", "邮箱",
                "统一社会信用代码", "注册号", "组织机构代码", "参保人数", "企业类型",
                "行业门类", "行业大类", "行业中类", "曾用名", "企业地址", "经营范围"]

_BASIC_COL = {"legal_person": 3, "capital": 4, "phone": 9, "email": 11, "insured": 15,
              "address": 21, "scope": 22}


def basic_row(name, idx, **kw):
    """默认各字段互不相同、无命中；用关键字参数覆盖需要撞车的字段。"""
    row = [name, f"id{idx}", "存续", f"法人{idx}", "1000万", "2010-01-01",
           "省", "市", "区", f"139000{idx:05d}", "", f"mail{idx}@corp{idx}.com",
           f"credit{idx}", "", "", "50", "有限责任公司",
           "门类", "大类", "中类", "", f"某市高新区科技园{idx}号",
           "一般项目：信息技术咨询服务；软件开发；企业管理咨询。"]
    for key, value in kw.items():
        row[_BASIC_COL[key]] = value
    return row


def write_xlsx(folder, filename, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(Path(folder) / filename)
    wb.close()


def run(folder, target="甲测试有限公司", **kwargs):
    with redirect_stdout(StringIO()):
        return run_check(data_dir=folder, target_names=target,
                         output_path=Path(folder) / "报告.xlsx", as_of_date="2026-10-07", **kwargs)


def overview_rows(path):
    wb = openpyxl.load_workbook(path)
    try:
        return [row for row in wb["核查结论"].iter_rows(min_row=4, values_only=True) if row[0]]
    finally:
        wb.close()


def overview_names(path):
    return [row[0] for row in overview_rows(path)]


def gap_rows(path):
    wb = openpyxl.load_workbook(path)
    try:
        return [row for row in wb["各家资料取得情况"].iter_rows(min_row=4, values_only=True)
                if any(value is not None for value in row)]
    finally:
        wb.close()


def write_phone_hit_pair(folder):
    """甲乙共用同一手机号 → 一条高风险命中（两个来源行）。"""
    write_xlsx(folder, "基础工商信息.xlsx", BASIC_HEADER,
               [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2, phone="13900000001")])


class Test概览结论先行(unittest.TestCase):
    def test_概览首行直接给出发现了什么(self):
        with tempfile.TemporaryDirectory() as d:
            write_phone_hit_pair(d)
            result = run(d)
            names = overview_names(result.output_path)
            self.assertEqual(names[0], "本次发现了什么")
            self.assertLess(names.index("本次发现了什么"), names.index("疑似关联方名单"))
            self.assertNotIn("纳入核查的公司数（不含被审计单位）", names, "范围与数量说明已移到独立工作表")

    def test_首行结论写明家数与风险分档(self):
        with tempfile.TemporaryDirectory() as d:
            write_phone_hit_pair(d)
            result = run(d)
            first = overview_rows(result.output_path)[0]
            self.assertIn("发现疑似关联方 1 家", first[3])
            self.assertIn("高风险 1 家", first[3])
            self.assertIn("乙测试有限公司", str(overview_rows(result.output_path)))

    def test_范围与数量说明在独立工作表(self):
        with tempfile.TemporaryDirectory() as d:
            write_phone_hit_pair(d)
            result = run(d)
            wb = openpyxl.load_workbook(result.output_path)
            try:
                self.assertIn("核查范围与数量说明", wb.sheetnames)
                names = [row[0] for row in wb["核查范围与数量说明"].iter_rows(min_row=4, values_only=True) if row[0]]
            finally:
                wb.close()
            for expected in ("被审计单位", "纳入核查的公司数（不含被审计单位）", "发现疑似关联方", "存在资料缺口的公司数"):
                self.assertIn(expected, names)

    def test_取数通道降级提示紧跟必读说明(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "报告.xlsx"
            write_report(out, [], [], {}, "甲测试有限公司",
                         limitations=[{"category": "取数通道降级", "source": "批量导出",
                                       "message": "批量通道连续失败，已改用逐家取数补采"}])
            names = overview_names(out)
            self.assertEqual(names[0], "本次发现了什么")
            self.assertEqual(names[1], "这份报告能说明什么（必读）")
            self.assertEqual(names[2], "取数通道降级提示")
            self.assertIn("逐家取数补采", overview_rows(out)[2][3])

    def test_零命中时声明不提供完整性保证(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            result = run(d)
            text = "\n".join(str(value) for row in overview_rows(result.output_path) for value in row if value)
            self.assertIn("本结果不提供关联方完整性保证", text)

    def test_全部低风险命中时同样声明不提供完整性保证(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "微信公众号.xlsx", ["序号", "公司名称", "账号", "公众号名"],
                       [[1, "甲测试有限公司", "账号甲", "同名公众号"], [2, "乙测试有限公司", "账号乙", "同名公众号"]])
            result = run(d)
            rows = overview_rows(result.output_path)
            self.assertIn("低风险 1 家", rows[0][3])
            text = "\n".join(str(value) for row in rows for value in row if value)
            self.assertIn("本结果不提供关联方完整性保证", text)

    def test_存在高风险命中时不出现完整性保证声明(self):
        with tempfile.TemporaryDirectory() as d:
            write_phone_hit_pair(d)
            result = run(d)
            text = "\n".join(str(value) for row in overview_rows(result.output_path) for value in row if value)
            self.assertNotIn("本结果不提供关联方完整性保证", text)


class Test名单疑似关系列(unittest.TestCase):
    def test_名单直接列出公司_风险_疑似关系_发现(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, phone="13900000001"),
                        basic_row("丙测试有限公司", 3)])
            write_xlsx(d, "商标.xlsx", ["序号", "公司名称", "商标名"],
                       [[1, "乙测试有限公司", "共用品牌"], [2, "丙测试有限公司", "共用品牌"]])
            result = run(d)
            wb = openpyxl.load_workbook(result.output_path)
            try:
                ws = wb["核查结论"]
                header_row = next(row[0].row for row in ws.iter_rows(min_row=4)
                                  if row[0].value == "疑似关联方名单")
                headers = [ws.cell(header_row, col).value for col in range(1, 5)]
                clues = [row for row in ws.iter_rows(min_row=header_row + 1, values_only=True)
                         if row[0] and str(row[1] or "") in {"高", "中", "低"}]
            finally:
                wb.close()
            self.assertEqual(headers, ["疑似关联方名单", "风险等级", "与被审计单位的疑似关系", "发现了什么（完整证据见证据明细）"])
            # 乙丙互相撞商标但与被审计单位无关，不再出现；只剩与被审计单位撞电话的乙。
            self.assertEqual([row[0] for row in clues], ["乙测试有限公司"])
            self.assertEqual(clues[0][1], "高")
            self.assertEqual(clues[0][2], "注册地址、电话或邮箱与被审计单位相同")


class Test证据明细行块分组(unittest.TestCase):
    def test_同一证据编号多行共享块底色_相邻块交替_块首行加粗上边框(self):
        with tempfile.TemporaryDirectory() as d:
            write_phone_hit_pair(d)
            write_xlsx(d, "商标.xlsx", ["序号", "公司名称", "商标名"],
                       [[1, "甲测试有限公司", "同名牌"], [2, "乙测试有限公司", "同名牌"]])
            result = run(d)
            wb = openpyxl.load_workbook(result.output_path)
            try:
                rows = list(wb["证据明细"].iter_rows(min_row=4, max_row=wb["证据明细"].max_row))
            finally:
                wb.close()
            ids = [row[0].value for row in rows]
            self.assertEqual(len(rows), 4, "两条证据各两个来源行")
            self.assertEqual(ids[0], ids[1])
            self.assertEqual(ids[2], ids[3])
            self.assertNotEqual(ids[0], ids[2])
            fills = [row[0].fill.fgColor.rgb for row in rows]
            self.assertEqual(fills[0], fills[1], "同一证据块内底色应一致")
            self.assertEqual(fills[2], fills[3], "同一证据块内底色应一致")
            self.assertNotEqual(fills[0], fills[2], "相邻证据块底色应交替")
            self.assertEqual(rows[0][0].border.top.style, "medium", "块首行应有加粗上边框")
            self.assertEqual(rows[2][0].border.top.style, "medium", "块首行应有加粗上边框")
            non_first_top = lambda row: row[0].border.top.style if row[0].border.top else None
            self.assertNotEqual(non_first_top(rows[1]), "medium")
            self.assertNotEqual(non_first_top(rows[3]), "medium")


class Test对人可读序号(unittest.TestCase):
    def test_汇总序号连续且保留GX锚点(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, phone="13900000001"),
                        basic_row("丙测试有限公司", 3)])
            write_xlsx(d, "商标.xlsx", ["序号", "公司名称", "商标名"],
                       [[1, "甲测试有限公司", "共用品牌"], [2, "丙测试有限公司", "共用品牌"]])
            result = run(d)
            wb = openpyxl.load_workbook(result.output_path)
            try:
                rows = list(wb["疑似关联方复核底稿"].iter_rows(min_row=4, values_only=True))
            finally:
                wb.close()
            self.assertEqual(len(rows), 2)
            self.assertEqual([row[1] for row in rows], ["乙测试有限公司", "丙测试有限公司"],
                             "底稿按风险从高到低排，高风险公司在前")
            self.assertEqual([row[-1] for row in rows], ["01", "02"])
            self.assertTrue(all(str(row[0]).startswith("GX-") for row in rows))

    def test_明细序号按证据编号连续且同一证据同号(self):
        with tempfile.TemporaryDirectory() as d:
            write_phone_hit_pair(d)
            write_xlsx(d, "商标.xlsx", ["序号", "公司名称", "商标名"],
                       [[1, "甲测试有限公司", "同名牌"], [2, "乙测试有限公司", "同名牌"]])
            result = run(d)
            wb = openpyxl.load_workbook(result.output_path)
            try:
                rows = list(wb["证据明细"].iter_rows(min_row=4, values_only=True))
            finally:
                wb.close()
            self.assertEqual([row[-1] for row in rows], ["证-01", "证-01", "证-02", "证-02"])
            self.assertTrue(all(str(row[0]).startswith("ZJ-") for row in rows))


class Test缺口页汇总与排序(unittest.TestCase):
    def test_保底模式下全体未取得维度只保留一条顶部汇总(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "报告.xlsx"
            write_report(out, [], [], {}, "甲公司 / 乙公司",
                         limitations=[{"category": "取数通道降级", "source": "批量导出",
                                       "message": "批量通道连续失败，已改用逐家取数补采"}])
            rows = gap_rows(out)
            self.assertEqual(rows[0][1], "维度级汇总", "首行应为维度级汇总")
            self.assertIn("本次未取得维度", rows[0][2])
            self.assertIn("影响全部 2 家", rows[0][2])
            self.assertFalse(any(row[0] in {"甲公司", "乙公司"} for row in rows),
                             "保底模式下全体未取得维度不再逐家重复列出")

    def test_非保底模式保留逐家明细且顶部仍有维度汇总(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "报告.xlsx"
            write_report(out, [], [], {}, "甲公司 / 乙公司")
            rows = gap_rows(out)
            self.assertEqual(rows[0][1], "维度级汇总")
            self.assertTrue(any(row[:3] == ("甲公司", "商标", "未取得") for row in rows),
                            "非保底模式应保留逐公司未取得明细行")

    def test_涉及命中的缺口行排在前面并标记(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, phone="13900000001", address=None),
                        basic_row("丙测试有限公司", 3, address=None)])
            result = run(d)
            rows = gap_rows(result.output_path)
            marked = [i for i, row in enumerate(rows) if row[4] and "涉及命中" in str(row[4])]
            self.assertTrue(marked, "涉及命中公司的缺口行应带标记")
            third_gap = [i for i, row in enumerate(rows)
                         if row[0] == "丙测试有限公司" and row[2] != "已取得记录"]
            self.assertTrue(third_gap)
            self.assertLess(max(marked), min(third_gap), "涉及命中的缺口行应排在未涉及命中的前面")
            yi_basic = next(row for row in rows if row[0] == "乙测试有限公司" and row[1] == "基础工商信息")
            self.assertIn("涉及命中", yi_basic[4])

    def test_全体未见记录的维度折叠为一条(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "动产抵押.xlsx", ["序号", "公司名称", "c", "d", "抵押人", "抵押权人"], [])
            (Path(d) / "取数说明.json").write_text(json.dumps({
                "company_names": ["甲测试有限公司", "乙测试有限公司"],
                "dimension_results": {code: "completed" for code in DIMENSION_NAMES},
            }, ensure_ascii=False), encoding="utf-8")
            result = run(d)
            rows = gap_rows(result.output_path)
            per_company = [row for row in rows if row[1] == "动产抵押" and row[0] not in {"全部核查对象", "其余核查对象"}]
            self.assertEqual(per_company, [], "全体未见记录的维度不该再逐家列：{}".format(per_company))
            folded = [row for row in rows if row[0] == "全部核查对象" and row[1] == "动产抵押"]
            self.assertEqual(len(folded), 1, "折叠行应只有一条")
            self.assertEqual(folded[0][2], "已核对范围，未见记录")
            self.assertIn("2 家", folded[0][4])

    def test_个别未见记录不折叠(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1), basic_row("乙测试有限公司", 2)])
            write_xlsx(d, "商标.xlsx", ["序号", "公司名称", "商标名"], [[1, "甲测试有限公司", "某商标"]])
            (Path(d) / "取数说明.json").write_text(json.dumps({
                "company_names": ["甲测试有限公司", "乙测试有限公司"],
                "dimension_results": {code: "completed" for code in DIMENSION_NAMES},
            }, ensure_ascii=False), encoding="utf-8")
            result = run(d)
            rows = gap_rows(result.output_path)
            self.assertTrue(any(row[0] == "乙测试有限公司" and row[1] == "商标" for row in rows),
                            "有公司取得记录的维度应按公司保留明细")


class Test缺口计数分两档(unittest.TestCase):
    def test_重度与轻度缺口公司分开统计(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, address=None),
                        basic_row("丙测试有限公司", 3)])
            (Path(d) / "取数说明.json").write_text(json.dumps({
                "company_names": ["甲测试有限公司", "乙测试有限公司", "丙测试有限公司"],
                "dimension_results": {code: "no_data" for code in DIMENSION_NAMES},
            }, ensure_ascii=False), encoding="utf-8")
            result = run(d, object_records=[{"name": "丁测试有限公司", "relation_type": "公开客户关系",
                                             "reasons": ["公开客户关系"]}])
            wb = openpyxl.load_workbook(result.output_path)
            try:
                quality = {row[0]: row[1] for row in wb["核查范围与数量说明"].iter_rows(min_row=4, values_only=True) if row[0]}
            finally:
                wb.close()
            self.assertEqual(quality["存在资料缺口的公司数"], 2)
            self.assertEqual(quality["其中重度缺口公司数"], 1, "丁公司未取得基础工商记录应为重度")
            self.assertEqual(quality["其中轻度缺口公司数"], 1, "乙公司仅缺个别字段应为轻度")


class Test对手方画像措辞(unittest.TestCase):
    def test_客商异常画像说明为对手方自身特征(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, capital="10万", insured="0")])
            result = run(d)
            hits = [h for h in result.hits if h["field"] == "profile"]
            self.assertTrue(hits, "客商异常画像未命中")
            self.assertTrue(hits[0]["evidence"].startswith("对手方"),
                            "画像线索应点明讲的是对手方自身特征：" + hits[0]["evidence"])
            self.assertIn("乙测试有限公司", hits[0]["evidence"])
            self.assertIn("真实交易", hits[0]["evidence"])

    def test_中关村等含村的正常地名不判居民楼(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, address="北京市海淀区中关村大街27号")])
            result = run(d)
            hits = [h for h in result.hits if h["field"] == "profile"]
            self.assertEqual(hits, [], "中关村这类正常地名不应判居民楼")

    def test_村组门牌仍判居民楼(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, address="四川省某县某镇幸福村3组12号")])
            result = run(d)
            hits = [h for h in result.hits if h["field"] == "profile"]
            self.assertTrue(hits, "村组门牌应仍判居民楼")
            self.assertIn("居民楼", hits[0]["evidence"])


class Test客商画像只对真实客商(unittest.TestCase):
    CUST_HEADER = ["序号", "公司名称", "c", "d", "e", "f", "g", "关联方名称"]
    SHELL = {"capital": "10万", "insured": "0"}

    def test_有客户表时未出现在客户表的公司不做画像(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, **self.SHELL),
                        basic_row("丙测试有限公司", 3, **self.SHELL)])
            write_xlsx(d, "客户.xlsx", self.CUST_HEADER, [[1, "甲测试有限公司", "", "", "", "", "", "丙测试有限公司"]])
            result = run(d)
            pairs = {frozenset({h["company_a"], h["company_b"]}) for h in result.hits if h["field"] == "profile"}
            self.assertNotIn(frozenset({"甲测试有限公司", "乙测试有限公司"}), pairs,
                             "乙公司不在客户表里，不该被当成客商做画像")
            self.assertIn(frozenset({"甲测试有限公司", "丙测试有限公司"}), pairs,
                          "丙公司是真实客商，应保留画像")

    def test_导出里没有任何客商数据时保持画像(self):
        with tempfile.TemporaryDirectory() as d:
            write_xlsx(d, "基础工商信息.xlsx", BASIC_HEADER,
                       [basic_row("甲测试有限公司", 1),
                        basic_row("乙测试有限公司", 2, **self.SHELL)])
            result = run(d)
            hits = [h for h in result.hits if h["field"] == "profile"]
            self.assertTrue(hits, "没有客商资料时仍应画像，避免静默丢掉线索")


class Test摘要按整条证据截断(unittest.TestCase):
    LONG_ADDRESS = "北京市朝阳区建国路88号甲座18层1808室" + "（园区北门内，紧邻地铁1号线大望路站C出口东侧100米）" * 3

    def _write_three_hits(self, folder):
        write_xlsx(folder, "基础工商信息.xlsx", BASIC_HEADER,
                   [basic_row("甲测试有限公司", 1, phone="13800000001",
                              address=self.LONG_ADDRESS, email="gongyong@corp1.com"),
                    basic_row("乙测试有限公司", 2, phone="13800000001",
                              address=self.LONG_ADDRESS, email="gongyong@corp1.com")])

    def test_摘要不把一条证据从中间截断并注明剩余条数(self):
        with tempfile.TemporaryDirectory() as d:
            self._write_three_hits(d)
            result = run(d)
            wb = openpyxl.load_workbook(result.output_path)
            try:
                summary_cell = next(row[4].value for row in wb["疑似关联方复核底稿"].iter_rows(min_row=4)
                                    if row[4].value)
            finally:
                wb.close()
            all_evidence = {h["evidence"] for h in result.hits}
            self.assertGreaterEqual(len(all_evidence), 3, "本用例需要至少三条证据")
            self.assertNotIn("…（完整见明细）", summary_cell, "摘要不应按字符硬截断")
            self.assertLess(len(summary_cell), sum(len(text) for text in all_evidence),
                            "摘要未做长度收敛，仍是全文")
            shown = [part for part in summary_cell.split("…（另")[0].split(" | ") if part]
            self.assertTrue(shown, "摘要不应为空")
            for part in shown:
                self.assertIn(part, all_evidence, "摘要把一条证据从中间截断了：" + part)
            self.assertLess(len(shown), len(all_evidence), "本用例应触发按条省略")
            self.assertRegex(summary_cell, r"…（另 {:d} 条见明细）$".format(len(all_evidence) - len(shown)))

    def test_概览线索摘要与汇总口径一致(self):
        with tempfile.TemporaryDirectory() as d:
            self._write_three_hits(d)
            result = run(d)
            rows = overview_rows(result.output_path)
            clue_row = next(row for row in rows if row[0] == "乙测试有限公司")
            self.assertIn("…（另", clue_row[3])
            self.assertNotIn("…（完整见明细）", clue_row[3])


if __name__ == "__main__":
    unittest.main()
