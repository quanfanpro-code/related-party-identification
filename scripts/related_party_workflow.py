# 关联方识别与核查 — 原创编排层
# Copyright (C) 2026 CPA-Q (quanfanpro-code)
#
# 本文件是 related-party-identification 的原创编排层,采用 GNU Affero General
# Public License v3.0 (AGPL-3.0) 发布。
# 完整协议见项目根目录 LICENSE 文件(AGPL-3.0)。
#
# 本项目的 scripts/cicpa/ 目录包含改编自 nigo/nigo-skills(MIT) 和
# jackwener/OpenCLI(Apache-2.0) 的代码,分别保留原始许可证。
# 详见 NOTICE 和 references/SOURCES.json。
#"""关联方识别 AI 对话式工作流入口。"""

import argparse
import csv
from dataclasses import asdict, dataclass, field, replace
from datetime import date, datetime
import importlib.util
import json
import os
from pathlib import Path
import re
import sys
from typing import Callable, Dict, List, Optional, Tuple

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts.cicpa.auth import (
    auth_status as _bridge_auth_status,
    default_browser_profile_path,
    default_status_path,
    detect_browser,
    read_login_status,
    start_guided_login,
)
from scripts.cicpa.browser_transport import OpenCliTransport
from scripts.cicpa.client import CicpaClient
from scripts.cicpa.exporter import (
    CicpaExporter,
    ExportError,
    ExportState,
    load_export_state,
)
from scripts.cicpa.opencli_setup import (
    download_opencli_extension,
    open_install_guide,
    opencli_extension_installed,
)
from scripts.discovery import (
    DiscoveryPolicy,
    SeedExportCandidate,
    discover,
)
from scripts.state_io import atomic_write_json
from scripts.逐家取数 import SingleCompanyError, collect_by_company


DEPENDENCIES = ("requests", "openpyxl")


@dataclass
class TaskState:
    task_id: str
    mode: str
    audited_entity: str
    stage: str = "created"
    status: str = "waiting_user"
    initial_depth: int = 2
    maximum_depth: int = 5
    approved_depth: int = 2
    candidate_cap: int = 100
    input_files: List[str] = field(default_factory=list)
    output_dir: str = ""
    export_state_path: str = ""
    message_zh: str = ""
    selected_column: str = ""
    candidate_columns: List[str] = field(default_factory=list)
    rejected_inputs: List[str] = field(default_factory=list)
    disclosed_parties: List[str] = field(default_factory=list)
    report_paths: List[str] = field(default_factory=list)
    candidate_count: int = 0
    candidate_export_state_path: str = ""
    candidate_records: List[dict] = field(default_factory=list)
    discovery_completed: bool = False
    discovery_warnings: List[str] = field(default_factory=list)
    raw_export_dir: str = ""
    degraded_export: bool = False
    degraded_reason: str = ""
    as_of_date: str = ""
    channel_warnings: List[str] = field(default_factory=list)


@dataclass
class NameListResult:
    status: str
    names: List[str] = field(default_factory=list)
    rejected: List[str] = field(default_factory=list)
    selected_column: str = ""
    candidate_columns: List[str] = field(default_factory=list)
    message_zh: str = ""


@dataclass
class WorkflowResult:
    status: str
    message_zh: str
    report_paths: List[Path] = field(default_factory=list)
    candidate_count: int = 0
    rejected_inputs: List[str] = field(default_factory=list)


def save_task_state(path: Path, state: TaskState) -> None:
    """以不带 BOM 的 UTF-8 原子保存非敏感任务状态。"""
    atomic_write_json(path, asdict(state))


def load_task_state(path: Path) -> TaskState:
    with open(path, "r", encoding="utf-8") as stream:
        payload = json.load(stream)
    return TaskState(**payload)


def auth_status(*, bridge_factory=None) -> Dict[str, str]:
    """实时验证 Edge 内注协会话;不读取、不返回任何凭据。"""
    if bridge_factory is None:
        return _bridge_auth_status()
    return _bridge_auth_status(bridge_factory=bridge_factory)


def ensure_login(
    *,
    auth_checker: Callable[[], Dict[str, str]] = auth_status,
    login_starter: Callable[[], Dict[str, object]] = start_guided_login,
) -> Dict[str, object]:
    """每次联网任务先验证登录，失效时直接打开系统浏览器。"""
    current = auth_checker()
    if current.get("status") == "authenticated":
        return current
    return login_starter()


def preflight(
    *,
    platform_name: str = os.name,
    version_info: Tuple[int, int, int] = tuple(sys.version_info[:3]),
    import_finder: Callable[[str], object] = importlib.util.find_spec,
    browser_detector: Callable[[], object] = detect_browser,
    opencli_checker: Callable[[Path], bool] = opencli_extension_installed,
    auth_checker: Callable[[], Dict[str, str]] = auth_status,
    local_app_data: Optional[str] = None,
    command_runner=None,
) -> Dict[str, object]:
    """只检查环境并给出下一步，不执行安装。"""
    _ = command_runner
    app_data = local_app_data or os.environ.get("LOCALAPPDATA", "")
    recommended_venv = str(
        Path(app_data) / "related-party-identification" / "venv"
    ) if app_data else ""
    result: Dict[str, object] = {
        "status": "ready",
        "message_zh": "环境检查通过",
        "python_version": "{}.{}.{}".format(*version_info),
        "python_executable": sys.executable,
        "recommended_venv": recommended_venv,
        "missing_dependencies": [],
        "browser": None,
        "login": {"status": "idle", "message_zh": "尚未检查"},
    }
    if platform_name != "nt":
        result["status"] = "unsupported_os"
        result["message_zh"] = "此 skill 仅支持 Windows"
        return result
    if version_info < (3, 9, 0):
        result["status"] = "unsupported_python"
        result["message_zh"] = "需要 Python 3.9 或更高版本"
        return result

    missing = [name for name in DEPENDENCIES if import_finder(name) is None]
    result["missing_dependencies"] = missing
    browser = browser_detector()
    result["browser"] = browser
    login = auth_checker()
    result["login"] = login

    if missing:
        result["status"] = "needs_user_approval"
        result["message_zh"] = (
            "缺少 {}。建议安装到独立环境 {}；说明用途并取得用户同意后才能安装。".format(
                "、".join(missing),
                recommended_venv or "当前用户本地应用数据目录",
            )
        )
    elif not browser:
        result["status"] = "needs_browser"
        result["message_zh"] = (
            "未检测到 Microsoft Edge 或 Google Chrome。"
            "本技能统一走 OpenCLI 驱动 Edge 的路线,请使用 Microsoft Edge。"
        )
    elif browser.get("name") in {"Microsoft Edge", "Google Chrome"} and not opencli_checker(
        default_browser_profile_path(browser_name=str(browser["name"]))
    ):
        result["status"] = "needs_opencli_extension"
        result["message_zh"] = (
            "需要一次性准备 OpenCLI:Edge 扩展(可从 Gitee 国内镜像自动准备)"
            "和 opencli 命令行工具。取得同意后技能会打开安装引导页面。"
        )
    elif login.get("status") != "authenticated":
        result["status"] = "ready_to_login"
        result["message_zh"] = "环境可用，下一步需要在注协官方页面登录"
    return result


def setup_opencli(
    *,
    browser: Optional[Dict[str, str]] = None,
    browser_detector: Callable[[], object] = detect_browser,
    downloader: Callable[[], Path] = download_opencli_extension,
    guide_opener: Callable[[Dict[str, str], Path], None] = open_install_guide,
) -> Dict[str, object]:
    """为默认 Edge 或 Chrome 准备国内镜像扩展并打开图形化安装入口。"""
    selected = browser or browser_detector()
    if not selected or selected.get("name") not in {"Microsoft Edge", "Google Chrome"}:
        return {
            "status": "not_applicable",
            "message_zh": "OpenCLI 安装只用于 Edge 或 Chrome;Firefox 不在本技能路线内。",
        }
    extension_dir = downloader()
    guide_opener(selected, extension_dir)
    return {
        "status": "waiting_user",
        "message_zh": (
            "OpenCLI 扩展文件已从 Gitee 国内镜像准备好，扩展管理页和文件夹已经打开。"
            "请先打开扩展管理页右侧的“开发人员模式”，完成后回复“打开了”。"
        ),
        "browser": selected.get("name"),
        "extension_path": str(extension_dir),
    }


def pick_path(
    kind: str,
    *,
    root_factory=None,
    filedialog_module=None,
) -> Dict[str, str]:
    """通过 Windows 选择窗口取得输入文件或输出目录。"""
    if kind not in {"input-file", "output-folder"}:
        raise ValueError("选择类型只能是 input-file 或 output-folder")
    if root_factory is None or filedialog_module is None:
        import tkinter
        from tkinter import filedialog

        root_factory = root_factory or tkinter.Tk
        filedialog_module = filedialog_module or filedialog

    root = root_factory()
    try:
        root.withdraw()
        root.attributes("-topmost", True)
        if kind == "input-file":
            path = filedialog_module.askopenfilename(
                title="选择关联方核查输入文件",
                filetypes=[
                    ("支持的名单文件", "*.xlsx *.xlsm *.csv *.txt"),
                    ("Excel 文件", "*.xlsx *.xlsm"),
                    ("CSV 文件", "*.csv"),
                    ("文本文件", "*.txt"),
                    ("所有文件", "*.*"),
                ],
            )
        else:
            path = filedialog_module.askdirectory(title="选择关联方核查输出文件夹")
    finally:
        root.destroy()

    if not path:
        return {
            "status": "cancelled",
            "path": "",
            "message_zh": "用户取消了选择，可以稍后重新打开窗口",
        }
    return {
        "status": "selected",
        "path": str(path),
        "message_zh": "已取得用户通过窗口选择的路径",
    }


KNOWN_NAME_COLUMNS = (
    "企业名称",
    "公司名称",
    "客户名称",
    "供应商名称",
    "单位名称",
    "名称",
)
LEDGER_TERMS = ("总账", "明细账", "序时账", "销售台账", "采购台账")


def _clean_names(values) -> Tuple[List[str], List[str]]:
    names = []
    rejected = []
    seen = set()
    for value in values:
        if value is None:
            continue
        name = str(value).strip()
        if not name or name in KNOWN_NAME_COLUMNS:
            continue
        if name.isdigit():
            rejected.append(name)
            continue
        if name not in seen:
            seen.add(name)
            names.append(name)
    return names, rejected


def _select_name_column(headers, requested_column=None):
    normalized = [str(value).strip() if value is not None else "" for value in headers]
    candidates = [value for value in normalized if value]
    if requested_column:
        if requested_column not in normalized:
            return None, candidates
        return normalized.index(requested_column), candidates
    for known in KNOWN_NAME_COLUMNS:
        if known in normalized:
            return normalized.index(known), candidates
    if len(candidates) == 1:
        return normalized.index(candidates[0]), candidates
    return None, candidates


def read_name_list(path, column_name=None) -> NameListResult:
    """只读取明确名单文件，不解析总账、销售或采购明细账。"""
    source = Path(path)
    if any(term in source.name for term in LEDGER_TERMS):
        return NameListResult(
            status="unsupported_input",
            message_zh="该文件看起来是账簿或台账；本模式只接受已经整理好的公司名单",
        )
    if not source.is_file():
        return NameListResult(status="invalid_input", message_zh="名单文件不存在")
    suffix = source.suffix.lower()
    if suffix == ".txt":
        try:
            text = source.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = source.read_text(encoding="gb18030")
        names, rejected = _clean_names(text.splitlines())
        return NameListResult(
            status="ready" if names else "empty",
            names=names,
            rejected=rejected,
            message_zh="名单读取完成" if names else "名单中没有可用公司名称",
        )

    if suffix == ".csv":
        try:
            stream = open(source, "r", encoding="utf-8-sig", newline="")
            rows = list(csv.reader(stream))
            stream.close()
        except UnicodeDecodeError:
            with open(source, "r", encoding="gb18030", newline="") as stream:
                rows = list(csv.reader(stream))
    elif suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError:
            return NameListResult(
                status="missing_dependency",
                message_zh="缺少 openpyxl，需要先取得用户同意后安装",
            )
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        return NameListResult(
            status="unsupported_input",
            message_zh="只支持 XLSX、XLSM、CSV 或 TXT 名单文件",
        )

    if not rows:
        return NameListResult(status="empty", message_zh="名单文件没有内容")
    column_index, candidates = _select_name_column(rows[0], column_name)
    if column_index is None:
        return NameListResult(
            status="needs_column_confirmation",
            candidate_columns=candidates,
            message_zh="无法唯一判断公司名称列，需要用户确认",
        )
    values = [
        row[column_index] if column_index < len(row) else None
        for row in rows[1:]
    ]
    names, rejected = _clean_names(values)
    selected = str(rows[0][column_index]).strip()
    return NameListResult(
        status="ready" if names else "empty",
        names=names,
        rejected=rejected,
        selected_column=selected,
        candidate_columns=candidates,
        message_zh="名单读取完成" if names else "所选列没有可用公司名称",
    )


# 客户/供应商表的对手方列只认这几个表头名，命中不了就跳过并留痕
COUNTERPARTY_HEADER_ALIASES = ("关联方名称", "客户名称", "供应商名称", "对手方名称", "交易对手方")


def extract_seed_export_candidates(data_dir, warnings=None) -> List[SeedExportCandidate]:
    """从被审计单位完整导出的客户、供应商表提取一层候选。

    对手方列只认明确的表头名（关联方名称/客户名称/供应商名称等小别名集），
    表头未命中时跳过该表并留 warning，不再按第 8 列猜位置。
    """
    try:
        import openpyxl
    except ImportError:
        return []
    candidates = []
    seen = {}
    for filename, relation_type in (
        ("客户.xlsx", "公开客户关系"),
        ("供应商.xlsx", "公开供应商关系"),
    ):
        path = Path(data_dir) / filename
        if not path.is_file():
            continue
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            sheet_name = workbook.active.title
        finally:
            workbook.close()
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        counterparty_index = next(
            (
                index
                for index, header in enumerate(headers)
                if header in COUNTERPARTY_HEADER_ALIASES
            ),
            None,
        )
        if counterparty_index is None:
            if warnings is not None:
                warnings.append(
                    "{}表头未命中对手方列（{}），已跳过该表候选提取".format(
                        filename, "、".join(COUNTERPARTY_HEADER_ALIASES)
                    )
                )
            continue
        for row_number, row in enumerate(rows[1:], 2):
            candidate_name = ""
            if (
                counterparty_index is not None
                and counterparty_index < len(row)
                and row[counterparty_index]
            ):
                candidate_name = str(row[counterparty_index]).strip()
            if not candidate_name:
                continue
            source = {"file": str(path.resolve()), "sheet": sheet_name,
                      "cell": f"{openpyxl.utils.get_column_letter(counterparty_index + 1)}{row_number}", "value": candidate_name}
            key = (candidate_name, relation_type)
            if key in seen:
                index = seen[key]
                candidates[index] = replace(candidates[index], sources=candidates[index].sources + (source,))
                continue
            seen[key] = len(candidates)
            candidates.append(
                SeedExportCandidate(
                    name=candidate_name,
                    relation_type=relation_type,
                    sources=(source,),
                )
            )
    return candidates


def _default_client_factory():
    status = _bridge_auth_status()
    if status.get("status") != "authenticated":
        raise RuntimeError("Edge 内注协会话无效，需要先完成登录")
    return CicpaClient(session=OpenCliTransport())


def _default_checker(**kwargs):
    from scripts.related_party_check import run_check

    return run_check(**kwargs)


_WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *{"COM{}".format(index) for index in range(1, 10)},
    *{"LPT{}".format(index) for index in range(1, 10)},
}


def _safe_company_stem(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name).strip())
    cleaned = cleaned.rstrip(" .")[:80].rstrip(" .") or "未命名公司"
    if cleaned.upper() in _WINDOWS_RESERVED_NAMES:
        cleaned = "_" + cleaned
    return cleaned


def _safe_output_path(folder: Path, filename: str) -> Path:
    path = folder / filename
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidate = folder / "{}_{}{}".format(stem, timestamp, suffix)
    counter = 2
    while candidate.exists():
        candidate = folder / "{}_{}_{}{}".format(
            stem,
            timestamp,
            counter,
            suffix,
        )
        counter += 1
    return candidate


def _ensure_export(state, state_path, exporter, company_names, *, state_field="export_state_path") -> Path:
    saved_path = getattr(state, state_field)
    if saved_path:
        export_state_path = Path(saved_path)
        export_state = load_export_state(export_state_path)
        if set(export_state.company_names) != set(company_names):
            raise ExportError("已保存导出任务的企业范围与本次不一致，未复用或覆盖旧证据")
        if export_state.status == "completed" and Path(export_state.extract_dir).is_dir():
            return _record_export_scope(export_state)
        if not export_state.required_dimensions:
            export_state_path = Path(state_path).with_name(
                "{}-adaptive-export-state.json".format(state.task_id)
            )
            setattr(state, state_field, str(export_state_path))
            save_task_state(state_path, state)
            export_state = exporter.start_export(company_names, export_state_path)
    else:
        export_state_path = Path(state_path).with_name(
            "{}-{}-state.json".format(state.task_id, "candidate-export" if state_field == "candidate_export_state_path" else "export")
        )
        setattr(state, state_field, str(export_state_path))
        save_task_state(state_path, state)
        export_state = exporter.start_export(company_names, export_state_path)
    exporter.run_to_completion(export_state, export_state_path)
    return _record_export_scope(load_export_state(export_state_path))


def _record_export_scope(export_state):
    """把范围和明确无数据说明随原始文件交付，离线重读仍能解释缺口。"""
    folder = Path(export_state.extract_dir)
    metadata = {
        "company_names": export_state.company_names,
        "created_at": export_state.created_at,
        "required_dimensions": export_state.required_dimensions,
        "dimension_results": export_state.dimension_results,
        "batch_history": export_state.batch_history,
    }
    path = folder / "取数说明.json"
    if path.exists():
        if json.loads(path.read_text(encoding="utf-8-sig")) != metadata:
            raise ExportError("取数说明与现有任务不一致，未覆盖原始资料")
    else:
        atomic_write_json(path, metadata)
    return folder


def _pause_export(state: TaskState, state_path: Path, error: ExportError) -> WorkflowResult:
    state.stage = "paused_export"
    state.status = "paused_export"
    state.message_zh = str(error)
    save_task_state(state_path, state)
    return WorkflowResult(state.status, state.message_zh)


def _fallback_to_single_company(
    state: TaskState,
    state_path: Path,
    client,
    data_dir,
    company_names: List[str],
    reason: str,
) -> Path:
    """批量导出通道不可用时的保底做法：改成逐个企业取数。

    逐家取数只能覆盖基础工商信息、最新公示股东、实际控制人、最终受益人和对外投资，
    其余维度（客户、供应商、发票、变更记录、主要人员等）本次取不到，报告会如实标注为缺口。

    保底文件固定写入 data_dir 下的"逐家取数补采"子目录，永不覆盖主目录里已有的批量导出证据。
    """
    state.degraded_export = True
    state.degraded_reason = str(reason)
    warning = (
        "批量导出通道不可用（{}），已改用逐家取数：本次只覆盖基础工商信息、最新公示股东、"
        "实际控制人、最终受益人、对外投资；客户、供应商、发票信息、变更记录、主要人员等"
        "维度本次无法取得，报告中如实标注为数据缺口。".format(reason)
    )
    if warning not in state.channel_warnings:
        state.channel_warnings.append(warning)
    save_task_state(state_path, state)
    fallback_dir = Path(data_dir) / "逐家取数补采"
    try:
        collect_by_company(client, company_names, fallback_dir, fallback_reason=str(reason))
    except SingleCompanyError as exc:
        save_task_state(state_path, state)
        raise ExportError(
            "{}；逐家取数保底也未取得资料：{}".format(reason, exc)
        ) from exc
    return Path(data_dir)




def run_workflow(
    state: TaskState,
    state_path,
    *,
    client_factory: Callable = _default_client_factory,
    exporter_factory: Optional[Callable] = None,
    discoverer: Callable = discover,
    checker: Callable = _default_checker,
) -> WorkflowResult:
    """运行三种模式，并在每个外部动作前后保存可恢复状态。"""
    state_file = Path(state_path)
    output_dir = Path(state.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    company_stem = _safe_company_stem(state.audited_entity)
    state.stage = "running"
    state.status = "running"
    save_task_state(state_file, state)

    additional_data_dirs = []
    object_records = []
    if state.mode == "existing_export":
        if not state.input_files:
            state.status = "waiting_user"
            state.message_zh = "需要通过窗口选择已有完整导出目录"
            save_task_state(state_file, state)
            return WorkflowResult(state.status, state.message_zh)
        data_dir = Path(state.input_files[0])
        if not data_dir.is_dir():
            state.status = "waiting_user"
            state.message_zh = "选择的已有数据目录不存在"
            save_task_state(state_file, state)
            return WorkflowResult(state.status, state.message_zh)
        candidate_count = 0
    elif state.mode in {"list_check", "discovery"}:
        client = client_factory()
        raw_export_dir = Path(state.raw_export_dir) if state.raw_export_dir else _safe_output_path(
            output_dir, "{}_注协原始导出".format(company_stem))
        state.raw_export_dir = str(raw_export_dir)
        save_task_state(state_file, state)
        if exporter_factory is None:
            exporter = CicpaExporter(client, artifact_dir=raw_export_dir)
        else:
            exporter = exporter_factory(client)
        if state.mode == "list_check":
            names = []
            rejected = []
            for input_file in state.input_files:
                parsed = read_name_list(
                    input_file,
                    column_name=state.selected_column or None,
                )
                if parsed.status != "ready":
                    state.status = parsed.status
                    state.message_zh = parsed.message_zh
                    state.candidate_columns = parsed.candidate_columns
                    save_task_state(state_file, state)
                    return WorkflowResult(state.status, state.message_zh)
                names.extend(parsed.names)
                rejected.extend(parsed.rejected)
                if parsed.selected_column:
                    state.selected_column = parsed.selected_column
            names = list(dict.fromkeys([state.audited_entity] + names))
            state.rejected_inputs = list(dict.fromkeys(rejected))
            if len(names) - 1 > state.candidate_cap:
                state.status = "needs_user_confirmation"
                state.candidate_count = len(names) - 1
                state.message_zh = (
                    "名单包含 {} 家候选，需用户确认后继续；"
                    "确认后请以 --candidate-cap {} 重新运行 run 命令".format(
                        state.candidate_count, state.candidate_count
                    )
                )
                save_task_state(state_file, state)
                return WorkflowResult(
                    state.status,
                    state.message_zh,
                    candidate_count=state.candidate_count,
                    rejected_inputs=state.rejected_inputs,
                )
            try:
                data_dir = _ensure_export(state, state_file, exporter, names)
            except ExportError as exc:
                try:
                    data_dir = _fallback_to_single_company(
                        state, state_file, client, raw_export_dir, names, str(exc)
                    )
                except ExportError as fallback_error:
                    return _pause_export(state, state_file, fallback_error)
            state.report_paths = list(
                dict.fromkeys(state.report_paths + [str(data_dir)])
            )
            candidate_count = len(names) - 1
            object_records = [{"name": name, "relation_type": "用户提供名单", "reasons": ["用户提供名单"],
                               "paths": [], "notes": [], "source_files": list(state.input_files)} for name in names if name != state.audited_entity]
        else:
            try:
                data_dir = _ensure_export(
                    state,
                    state_file,
                    exporter,
                    [state.audited_entity],
                )
            except ExportError as exc:
                try:
                    data_dir = _fallback_to_single_company(
                        state,
                        state_file,
                        client,
                        raw_export_dir,
                        [state.audited_entity],
                        str(exc),
                    )
                except ExportError as fallback_error:
                    return _pause_export(state, state_file, fallback_error)
            state.report_paths = list(
                dict.fromkeys(state.report_paths + [str(data_dir)])
            )
            if not state.discovery_completed:
                seed_warnings = []
                seed_candidates = (
                    [] if state.degraded_export else extract_seed_export_candidates(data_dir, warnings=seed_warnings)
                )
                discovery_result = discoverer(
                    state.audited_entity, client=client,
                    policy=DiscoveryPolicy(initial_depth=state.initial_depth,
                        maximum_depth=state.maximum_depth, candidate_cap=state.candidate_cap),
                    approved_depth=state.approved_depth,
                    seed_export_candidates=seed_candidates,
                )
                state.candidate_count = len(discovery_result.candidates)
                state.candidate_records = [asdict(candidate) for candidate in discovery_result.candidates.values()]
                state.discovery_warnings = list(discovery_result.warnings) + seed_warnings
                state.discovery_completed = discovery_result.status == "completed"
                state.status = discovery_result.status
                state.message_zh = discovery_result.message_zh
                save_task_state(state_file, state)
                if not state.discovery_completed:
                    return WorkflowResult(state.status, state.message_zh, candidate_count=state.candidate_count)
                snapshot = data_dir / "候选发现记录.json"
                snapshot_data = {"candidates": state.candidate_records, "warnings": state.discovery_warnings}
                if snapshot.exists():
                    if json.loads(snapshot.read_text(encoding="utf-8-sig")) != snapshot_data:
                        raise ExportError("候选发现记录与现有资料不一致，未覆盖原文件")
                else:
                    atomic_write_json(snapshot, snapshot_data)
            object_records = state.candidate_records
            candidate_count = state.candidate_count
            if object_records and state.degraded_export:
                # 保底路线：候选企业也用逐家取数补齐，维度文件在补采子目录按全部企业重写一次
                try:
                    collect_by_company(
                        client,
                        [state.audited_entity] + [record["name"] for record in object_records],
                        Path(data_dir) / "逐家取数补采",
                        fallback_reason=state.degraded_reason,
                    )
                except SingleCompanyError as exc:
                    state.channel_warnings.append(
                        "逐家取数补充候选企业时未完整取得：{}".format(exc)
                    )
                    save_task_state(state_file, state)
            elif object_records:
                candidate_exporter = (CicpaExporter(client, artifact_dir=data_dir / "候选公司原始导出")
                                      if exporter_factory is None else exporter_factory(client))
                try:
                    candidate_dir = _ensure_export(state, state_file, candidate_exporter,
                        [record["name"] for record in object_records], state_field="candidate_export_state_path")
                except ExportError as exc:
                    return _pause_export(state, state_file, exc)
                if candidate_dir != data_dir:
                    additional_data_dirs.append(candidate_dir)
    else:
        raise ValueError("未知工作流模式：{}".format(state.mode))

    report_path = _safe_output_path(
        output_dir,
        "{}_关联方核查报告.xlsx".format(company_stem),
    )
    checker(
        data_dir=data_dir,
        target_names=[state.audited_entity],
        output_path=report_path,
        as_of_date=date.fromisoformat(state.as_of_date) if state.as_of_date else date.today(),
        disclosed_parties=set(state.disclosed_parties) if state.disclosed_parties else None,
        task_mode=state.mode,
        additional_data_dirs=additional_data_dirs,
        object_records=object_records,
        discovery_warnings=state.discovery_warnings,
        channel_warnings=state.channel_warnings,
        scope_metadata={
            "初始穿透层数": state.initial_depth,
            "批准穿透层数": state.approved_depth,
            "最多穿透层数": state.maximum_depth,
            "候选上限": state.candidate_cap,
            "候选数量": candidate_count,
        },
    )
    state.report_paths = list(dict.fromkeys(state.report_paths + [str(report_path)]))
    state.candidate_count = candidate_count
    state.stage = "completed"
    state.status = "completed"
    state.message_zh = "关联方核查已完成"
    save_task_state(state_file, state)
    return WorkflowResult(
        status=state.status,
        message_zh=state.message_zh,
        report_paths=[Path(path) for path in state.report_paths],
        candidate_count=candidate_count,
        rejected_inputs=state.rejected_inputs,
    )


def _emit(payload: Dict[str, object]) -> None:
    print(json.dumps(payload, ensure_ascii=False))


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="关联方识别 AI 对话式工作流")
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("preflight", help="检查运行环境")
    picker = subparsers.add_parser("pick", help="打开 Windows 路径选择窗口")
    picker.add_argument("--kind", required=True, choices=["input-file", "output-folder"])
    subparsers.add_parser("auth-status", help="验证已有登录状态")
    opencli_parser = subparsers.add_parser(
        "opencli-setup",
        help="为 Edge 或 Chrome 准备 OpenCLI 扩展",
    )
    opencli_parser.add_argument(
        "--browser",
        choices=["edge", "chrome"],
        default="edge",
    )
    login_parser = subparsers.add_parser(
        "login-start", help="在 Edge 中打开注协登录页并等待登录"
    )
    subparsers.add_parser("login-status", help="读取后台登录进度")
    run_parser = subparsers.add_parser("run", help="运行关联方识别任务")
    run_parser.add_argument(
        "--mode",
        required=True,
        choices=["discovery", "list_check", "existing_export"],
    )
    run_parser.add_argument("--audited-entity", required=True)
    run_parser.add_argument("--state-path", required=True, type=Path)
    run_parser.add_argument("--input", action="append", default=[])
    run_parser.add_argument("--output-dir", required=True)
    run_parser.add_argument("--approved-depth", type=int, default=2)
    run_parser.add_argument("--as-of-date", default="", help="核查基准日，格式 YYYY-MM-DD；缺省为运行当天")
    run_parser.add_argument("--candidate-cap", type=int, default=100, help="候选企业数量上限；确认超出上限后用更大的值重新运行即可继续")
    resume_parser = subparsers.add_parser("resume", help="恢复关联方识别任务")
    resume_parser.add_argument("--state-path", required=True, type=Path)
    subparsers.add_parser("validate", help="验证 skill 分发包")
    args = parser.parse_args(argv)

    if args.command == "preflight":
        payload = preflight()
    elif args.command == "pick":
        payload = pick_path(args.kind)
    elif args.command == "auth-status":
        payload = auth_status()
    elif args.command == "opencli-setup":
        browser_name = {
            "edge": "Microsoft Edge",
            "chrome": "Google Chrome",
        }[args.browser]
        payload = setup_opencli(
            browser=detect_browser(preferred_name=browser_name)
        )
    elif args.command == "login-start":
        payload = start_guided_login()
    elif args.command == "login-status":
        payload = read_login_status(default_status_path())
    elif args.command == "run":
        state = TaskState(
            task_id=args.state_path.stem,
            mode=args.mode,
            audited_entity=args.audited_entity,
            input_files=args.input,
            output_dir=args.output_dir,
            approved_depth=args.approved_depth,
            candidate_cap=args.candidate_cap,
            as_of_date=args.as_of_date,
        )
        login = (
            {"status": "authenticated"}
            if state.mode == "existing_export"
            else ensure_login()
        )
        if login.get("status") != "authenticated":
            state.stage = "login"
            state.status = "waiting_user"
            state.message_zh = str(login.get("message_zh", "请在系统浏览器中完成登录"))
            save_task_state(args.state_path, state)
            payload = {
                "status": state.status,
                "message_zh": state.message_zh,
                "report_paths": [],
                "candidate_count": 0,
                "rejected_inputs": [],
            }
        else:
            result = run_workflow(state, args.state_path)
            payload = {
                "status": result.status,
                "message_zh": result.message_zh,
                "report_paths": [str(path) for path in result.report_paths],
                "candidate_count": result.candidate_count,
                "rejected_inputs": result.rejected_inputs,
            }
    elif args.command == "resume":
        state = load_task_state(args.state_path)
        login = (
            {"status": "authenticated"}
            if state.mode == "existing_export"
            else ensure_login()
        )
        if login.get("status") != "authenticated":
            state.stage = "login"
            state.status = "waiting_user"
            state.message_zh = str(login.get("message_zh", "请在系统浏览器中完成登录"))
            save_task_state(args.state_path, state)
            payload = {
                "status": state.status,
                "message_zh": state.message_zh,
                "report_paths": [],
                "candidate_count": state.candidate_count,
                "rejected_inputs": state.rejected_inputs,
            }
        else:
            result = run_workflow(state, args.state_path)
            payload = {
                "status": result.status,
                "message_zh": result.message_zh,
                "report_paths": [str(path) for path in result.report_paths],
                "candidate_count": result.candidate_count,
                "rejected_inputs": result.rejected_inputs,
            }
    else:
        payload = {
            "status": "not_ready",
            "message_zh": "该命令将在后续工作流任务中启用",
        }
    _emit(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
