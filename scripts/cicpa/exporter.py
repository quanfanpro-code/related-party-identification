# 关联方识别与核查 — 注协完整维度导出器
# 本文件改编自 nigo/nigo-skills/cicpa-company-query(MIT)
# 上游作者: nigo(涂佳兵) | 原始仓库: https://github.com/nigo81/nigo-skills
# 上游协议: MIT
#
# 本地修改: CPA-Q(quanfanpro-code)
# 本文件的修改部分同样以 MIT 协议发布,与上游保持一致。
#"""中注协完整维度导出的安全批次、任务关联与文件验证。"""

from dataclasses import asdict, dataclass, field
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import random
import shutil
import tempfile
import time
from typing import Any, Callable, Dict, List, Optional
import unicodedata
import zipfile

from .client import CicpaClient, CicpaError
from scripts.state_io import atomic_write_json


class ExportError(CicpaError):
    """完整维度导出基础错误。"""


class ExportTaskNotCorrelated(ExportError):
    """下载中心任务无法与本次批次唯一对应。"""


class ExportTaskAmbiguous(ExportTaskNotCorrelated):
    """下载中心同时出现多个可能属于本批次的新任务。"""


class ExportValidationError(ExportError):
    """下载文件未通过结构或企业名称验证。"""


class ExportWorkbookIncomplete(ExportValidationError):
    """维度工作簿残缺（无法打开或连表头都没有），按缺文件处理。"""


REQUIRED_DIMENSIONS = (
    ("S0000002", "基础工商信息"),
    ("S0000006", "股东信息"),
    ("S0000103", "最新公示股东"),
    ("S0000032", "实际控制人"),
    ("S0000020", "最终受益人"),
    ("S0000037", "主要人员（高管）"),
    ("S0000123", "核心团队"),
    ("S0000104", "对外投资（新）"),
    ("S0000105", "参控股企业"),
    ("S0000107", "发票信息"),
    ("S0000119", "客户"),
    ("S0000118", "供应商"),
    ("S0000016", "变更记录"),
    ("S0000019", "法定代表人变更"),
    ("S0000018", "经营异常"),
    ("S0000013", "股权质押"),
    ("S0000012", "动产抵押"),
    ("S0000041", "商标"),
    ("S0000036", "软件著作权"),
    ("S0000101", "微信公众号"),
)
DIMENSION_NAMES = dict(REQUIRED_DIMENSIONS)
# 终态集合：completed/no_data 是正常结果；missing 是"没拿到"（重试上限后放弃），
# 与 no_data（服务端确认没有）严格区分，报告侧对非 no_data 的缺失文件走"缺失维度"通道
TERMINAL_DIMENSION_RESULTS = {"completed", "no_data", "missing"}
# 同一维度在下载批次中缺文件达到此次数后，记 missing 终态，不再无限重排
MISSING_FILE_RETRY_LIMIT = 2


def initial_dimension_groups() -> List[List[str]]:
    codes = [code for code, _name in REQUIRED_DIMENSIONS]
    return [codes[:10], codes[10:]]


def next_dimension_group(state: "ExportState") -> List[str]:
    while state.pending_groups:
        group = state.pending_groups.pop(0)
        remaining = [
            code
            for code in group
            if state.dimension_results.get(code) not in TERMINAL_DIMENSION_RESULTS
        ]
        if remaining:
            return remaining
    return []


@dataclass
class ExportState:
    batch_no: str
    company_names: List[str]
    created_at: str
    triggered_at: str
    preexisting_task_ids: List[str] = field(default_factory=list)
    staging_dir: str = ""
    task_id: str = ""
    task_name: str = ""
    status: str = "waiting"
    archive_path: str = ""
    extract_dir: str = ""
    message_zh: str = ""
    direct_delivery: bool = False
    required_dimensions: List[str] = field(default_factory=list)
    dimension_results: Dict[str, str] = field(default_factory=dict)
    # 各维度"下载批次里没有它的文件"的累计次数，用于缺文件重试上限
    missing_counts: Dict[str, int] = field(default_factory=dict)
    pending_groups: List[List[str]] = field(default_factory=list)
    active_group: List[str] = field(default_factory=list)
    batch_history: List[Dict[str, Any]] = field(default_factory=list)
    poll_windows: int = 0
    batch_failure_streak: int = 0
    single_dimension_mode: bool = False


def save_export_state(path: Path, state: ExportState) -> None:
    """以不带 BOM 的 UTF-8 原子保存非敏感任务状态。"""
    atomic_write_json(path, asdict(state))


def load_export_state(path: Path) -> ExportState:
    with open(path, "r", encoding="utf-8") as stream:
        payload = json.load(stream)
    return ExportState(**payload)


class CicpaExporter:
    ZSK_BASE = "https://zsk-cmis.cicpa.org.cn"

    def __init__(
        self,
        client: CicpaClient,
        *,
        staging_root: Optional[Path] = None,
        artifact_dir: Optional[Path] = None,
        now: Callable[[], str] = lambda: datetime.now().astimezone().strftime(
            "%Y-%m-%d %H:%M"
        ),
        sleeper: Callable[[float], None] = time.sleep,
        randint: Callable[[int, int], int] = random.randint,
    ):
        self.client = client
        self.staging_root = Path(staging_root) if staging_root is not None else None
        self.artifact_dir = Path(artifact_dir) if artifact_dir is not None else None
        self._now = now
        self._sleep = sleeper
        self._randint = randint

    @staticmethod
    def _business_data(payload: Any, alias: str) -> Any:
        if not isinstance(payload, dict):
            raise ExportError("{}返回结构无效".format(alias))
        if payload.get("status_code") not in (None, 0):
            raise ExportError(
                "{}失败：{}".format(alias, payload.get("status_msg") or "未知业务错误")
            )
        return payload.get("data", {})

    def _new_staging_dir(self) -> Path:
        if self.artifact_dir is not None:
            self.artifact_dir.mkdir(parents=True, exist_ok=True)
            return self.artifact_dir
        if self.staging_root is not None:
            self.staging_root.mkdir(parents=True, exist_ok=True)
        return Path(
            tempfile.mkdtemp(
                prefix="rpi_export_",
                dir=str(self.staging_root) if self.staging_root is not None else None,
            )
        )

    @staticmethod
    def _normalize_names(company_names: List[str]) -> List[str]:
        result = []
        seen = set()
        for value in company_names:
            name = str(value).strip()
            if name and name not in seen:
                seen.add(name)
                result.append(name)
        if not result:
            raise ValueError("企业名单不能为空")
        return result

    def _upload_companies(self, company_names: List[str], staging_dir: Path) -> str:
        try:
            import openpyxl
        except ImportError as exc:
            raise ExportError("缺少 openpyxl，需先完成依赖预检") from exc

        workbook_path = staging_dir / "upload-companies.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["企业名称"])
        for name in company_names:
            sheet.append([name])
        workbook.save(workbook_path)
        content = workbook_path.read_bytes()
        payload = self.client.request_json(
            "POST",
            self.ZSK_BASE + "/open/industry_chain_api/v1/batch_search/upload_local",
            files={
                "file": (
                    "companies.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            timeout=60,
        )
        data = self._business_data(payload, "上传企业名单")
        batch_no = str(data.get("batch_id") or "")
        if not batch_no:
            raise ExportError("上传成功响应中缺少批次号")
        return batch_no

    def _trigger_search(self, batch_no: str) -> None:
        payload = self.client.request_json(
            "POST",
            self.ZSK_BASE + "/open/industry_chain_api/v1/batch_search/enter_search",
            json={
                "batch_id": batch_no,
                "page": 1,
                "page_size": 10,
                "type": "credit",
            },
        )
        self._business_data(payload, "触发批量查询")

    def _get_dimension_catalog(self) -> Dict[str, str]:
        payload = self.client.request_json(
            "GET",
            self.ZSK_BASE
            + "/open/industry_chain_api/v1/search/export/get_dimension_class",
            params={"internal": "false"},
            cache_key=("export", "dimensions"),
        )
        data = self._business_data(payload, "获取导出维度")
        dimensions = {}
        for group in data if isinstance(data, list) else []:
            for child in group.get("children", []) if isinstance(group, dict) else []:
                code = child.get("second_dimension_code") if isinstance(child, dict) else None
                if code:
                    name = (
                        child.get("second_dimension_name")
                        or child.get("dimension_name")
                        or child.get("name")
                        or ""
                    )
                    dimensions[str(code)] = str(name)
        if not dimensions:
            raise ExportError("未取得可用导出维度")
        return dimensions

    def _get_dimensions(self) -> List[str]:
        return list(self._get_dimension_catalog())

    def _list_tasks(self, kind: str = "normal") -> List[Dict[str, Any]]:
        payload = self.client.request_json(
            "GET",
            self.ZSK_BASE + "/open/industry_chain_api/v1/download/task_list",
            kind=kind,
            params={"page": 1, "pageSize": 20},
        )
        data = self._business_data(payload, "读取下载中心")
        if not isinstance(data, dict):
            return []
        return [item for item in data.get("list", []) if isinstance(item, dict)]

    def _trigger_export(self, batch_no: str, dimensions: List[str]) -> str:
        payload = self.client.request_json(
            "POST",
            self.ZSK_BASE
            + "/open/industry_chain_api/v1/batch_search/enter_search_out_type",
            json={"batch_id": batch_no, "dimensions": dimensions},
        )
        data = self._business_data(payload, "触发完整维度导出")
        if isinstance(data, dict):
            return str(data.get("task_id") or "")
        return ""

    def start_export(self, company_names: List[str], state_path: Path) -> ExportState:
        """上传一次名单，初始化 10+10 并触发第一组。"""
        names = self._normalize_names(company_names)
        staging_dir = self._new_staging_dir()
        batch_no = self._upload_companies(names, staging_dir)
        self._trigger_search(batch_no)
        catalog = self._get_dimension_catalog()
        required = [code for code, _name in REQUIRED_DIMENSIONS]
        missing = [code for code in required if code not in catalog]
        if missing:
            raise ExportError("行业库缺少关联方核查维度：{}".format("、".join(missing)))
        created_at = self._now()
        state = ExportState(
            batch_no=batch_no,
            company_names=names,
            created_at=created_at,
            triggered_at="",
            staging_dir=str(staging_dir),
            status="waiting",
            message_zh="已初始化 20 个关联方核查维度",
            direct_delivery=self.artifact_dir is not None,
            required_dimensions=required,
            pending_groups=initial_dimension_groups(),
        )
        save_export_state(state_path, state)
        self.trigger_next_group(state, state_path)
        return state

    def trigger_next_group(self, state: ExportState, state_path: Path) -> bool:
        """在没有活动任务时触发下一组未完成维度。"""
        if state.active_group or state.task_id:
            return False
        while True:
            group = next_dimension_group(state)
            if not group:
                return False
            state.active_group = group
            state.triggered_at = self._now()
            try:
                before = self._list_tasks()
                state.preexisting_task_ids = [
                    str(task.get("task_id"))
                    for task in before
                    if task.get("task_id") is not None
                ]
                state.task_id = self._trigger_export(state.batch_no, group)
            except CicpaError as exc:
                if len(group) > 1 and not state.single_dimension_mode:
                    self._record_batch_failure(
                        state,
                        state_path,
                        result="trigger_error",
                        detail=str(exc),
                    )
                    continue
                state.status = "paused"
                state.message_zh = "单维度 {}（{}）下载请求失败，已暂停：{}".format(
                    group[0],
                    DIMENSION_NAMES.get(group[0], group[0]),
                    exc,
                )
                save_export_state(state_path, state)
                raise
            state.task_name = ""
            state.poll_windows = 0
            state.status = "waiting"
            state.message_zh = "已触发 {} 个维度，等待下载任务".format(len(group))
            save_export_state(state_path, state)
            return True

    @staticmethod
    def _unfinished_dimensions(state: ExportState) -> List[str]:
        return [
            code
            for code in state.required_dimensions
            if state.dimension_results.get(code) not in TERMINAL_DIMENSION_RESULTS
        ]

    def _record_batch_failure(
        self,
        state: ExportState,
        state_path: Path,
        *,
        result: str,
        detail: str = "",
    ) -> None:
        """记录连续批量失败；达到三次后只逐项处理尚未完成维度。"""
        failed_group = [
            code
            for code in state.active_group
            if state.dimension_results.get(code) not in TERMINAL_DIMENSION_RESULTS
        ]
        state.batch_failure_streak += 1
        state.batch_history.append(
            {
                "dimensions": list(failed_group),
                "task_id": state.task_id,
                "result": result,
                "detail": detail,
                "consecutive_batch_failures": state.batch_failure_streak,
            }
        )
        state.active_group = []
        state.task_id = ""
        state.task_name = ""
        state.poll_windows = 0
        state.status = "waiting"
        if state.batch_failure_streak >= 3:
            state.single_dimension_mode = True
            state.pending_groups = [
                [code] for code in self._unfinished_dimensions(state)
            ]
            state.message_zh = (
                "批量接口连续失败 3 次，已改为逐个下载尚未完成的维度"
            )
        else:
            existing = {code for code in failed_group}
            remaining_groups = []
            for group in state.pending_groups:
                remaining = [
                    code
                    for code in group
                    if code not in existing
                    and state.dimension_results.get(code)
                    not in TERMINAL_DIMENSION_RESULTS
                ]
                if remaining:
                    remaining_groups.append(remaining)
            state.pending_groups = ([failed_group] if failed_group else []) + remaining_groups
            state.message_zh = "批量接口连续失败 {}/3，继续批量重试未完成维度".format(
                state.batch_failure_streak
            )
        save_export_state(state_path, state)

    def handle_poll_window_failure(
        self,
        state: ExportState,
        state_path: Path,
    ) -> None:
        """批量轮询失败计入连续次数；单维度保留两个轮询窗口，仍失败则记 missing 继续。"""
        state.poll_windows += 1
        if len(state.active_group) > 1 and not state.single_dimension_mode:
            self._record_batch_failure(
                state,
                state_path,
                result="poll_timeout",
            )
            return
        if state.poll_windows < 2:
            state.message_zh = "第一个轮询窗口结束，保留当前维度组继续等待"
            save_export_state(state_path, state)
            return
        # 单维度两个轮询窗口仍失败：记为未取得并继续后续维度，不再暂停等人处理
        self._mark_active_group_missing(
            state,
            state_path,
            detail="经过两个轮询窗口仍未取得文件",
        )

    def _mark_active_group_missing(
        self,
        state: ExportState,
        state_path: Path,
        *,
        detail: str,
    ) -> None:
        """把活动组里尚未完成的维度记为 missing（未取得），清空活动组后继续。"""
        unfinished = [
            code
            for code in state.active_group
            if state.dimension_results.get(code) not in TERMINAL_DIMENSION_RESULTS
        ]
        for code in unfinished:
            state.dimension_results[code] = "missing"
        names = "、".join(
            "{}（{}）".format(code, DIMENSION_NAMES.get(code, code))
            for code in unfinished
        )
        state.batch_history.append(
            {
                "dimensions": list(unfinished),
                "task_id": state.task_id,
                "result": "marked_missing",
                "detail": detail,
            }
        )
        state.active_group = []
        state.task_id = ""
        state.task_name = ""
        state.poll_windows = 0
        state.status = "waiting"
        state.message_zh = "维度 {}{}，记为未取得并继续后续维度".format(names, detail)
        save_export_state(state_path, state)

    @staticmethod
    def correlate_task(
        state: ExportState,
        tasks: List[Dict[str, Any]],
        *,
        allow_missing: bool = False,
    ):
        """只接受能与本次批次唯一对应的下载任务。"""
        batch_tasks = [task for task in tasks if task.get("type") == "批量查询"]
        if state.task_id:
            exact = [
                task
                for task in batch_tasks
                if str(task.get("task_id") or "") == state.task_id
            ]
            if len(exact) == 1:
                return exact[0]
            if allow_missing:
                return None
            raise ExportTaskNotCorrelated("下载中心没有找到已记录的任务号")

        candidates = []
        for task in batch_tasks:
            task_id = str(task.get("task_id") or "")
            if not task_id or task_id in state.preexisting_task_ids:
                continue
            task_batch = str(task.get("batch_id") or "")
            if task_batch and task_batch != state.batch_no:
                continue
            task_date = str(task.get("date") or "")
            if not task_date or task_date < state.triggered_at:
                continue
            candidates.append(task)

        matching_batch = [
            task for task in candidates if str(task.get("batch_id") or "") == state.batch_no
        ]
        if len(matching_batch) == 1:
            return matching_batch[0]
        if len(candidates) == 1:
            return candidates[0]
        if len(candidates) > 1:
            raise ExportTaskAmbiguous("同时出现多个新下载任务，无法唯一确认本次任务")
        if allow_missing:
            return None
        raise ExportTaskNotCorrelated("尚未发现能与本次批次对应的下载任务")

    def wait_for_task(
        self,
        state: ExportState,
        state_path: Path,
        *,
        max_polls: int = 18,
    ) -> Dict[str, Any]:
        """按轮询随机等待策略寻找并跟踪唯一任务。"""
        for _attempt in range(max_polls):
            task = self.correlate_task(
                state,
                self._list_tasks(kind="poll"),
                allow_missing=True,
            )
            if task is not None:
                state.task_id = str(task.get("task_id") or "")
                state.task_name = str(task.get("name") or "")
                state.message_zh = "已关联下载任务，正在等待文件可下载"
                save_export_state(state_path, state)
                if task.get("status") == 1:
                    return task
        state.message_zh = "轮询结束，未发现可唯一对应且已完成的下载任务"
        save_export_state(state_path, state)
        raise ExportTaskNotCorrelated(state.message_zh)

    @staticmethod
    def _safe_extract(archive_path: Path, extract_dir: Path) -> None:
        extract_root = extract_dir.resolve()
        with zipfile.ZipFile(archive_path, "r") as archive:
            members = archive.infolist()
            for member in members:
                destination = (extract_dir / member.filename).resolve()
                try:
                    outside = os.path.commonpath(
                        [str(extract_root), str(destination)]
                    ) != str(extract_root)
                except ValueError:
                    # 跨盘符等无法比较的路径按校验失败处理，不裸抛 ValueError
                    raise ExportValidationError("ZIP 中存在无法校验安全性的路径")
                if outside:
                    raise ExportValidationError("ZIP 中存在越出暂存目录的路径")
            for member in members:
                destination = extract_dir / member.filename
                if member.is_dir():
                    destination.mkdir(parents=True, exist_ok=True)
                    continue
                destination.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(member) as source, open(destination, "wb") as target:
                    target.write(source.read())

    @staticmethod
    def _normalize_company_name(value: str) -> str:
        normalized = unicodedata.normalize("NFKC", str(value))
        return "".join(normalized.split())

    @classmethod
    def _validate_companies(cls, extract_dir: Path, expected_names: List[str]) -> None:
        try:
            import openpyxl
        except ImportError as exc:
            raise ExportValidationError("缺少 openpyxl，无法验证导出文件") from exc

        workbooks = list(extract_dir.rglob("*.xlsx"))
        base_file = next(
            (path for path in workbooks if "基础工商信息" in path.stem),
            None,
        )
        if base_file is None:
            base_file = next(
                (
                    path
                    for path in workbooks
                    if "基础" in path.stem and "工商" in path.stem
                ),
                None,
            )
        if base_file is None:
            raise ExportValidationError("完整导出中缺少基础工商信息工作簿")

        workbook = openpyxl.load_workbook(base_file, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            actual_names = [
                str(row[0]).strip()
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if row and row[0]
            ]
        finally:
            workbook.close()

        expected = {cls._normalize_company_name(name) for name in expected_names}
        actual = {cls._normalize_company_name(name) for name in actual_names}
        if expected != actual:
            raise ExportValidationError(
                "下载文件企业集合与本次上传名单不一致：预期 {} 家，实际 {} 家".format(
                    len(expected), len(actual)
                )
            )

    @staticmethod
    def _file_sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with open(path, "rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _normalize_dimension_name(value: str) -> str:
        normalized = unicodedata.normalize("NFKC", str(value))
        return "".join(normalized.split()).replace("(", "").replace(")", "")

    @classmethod
    def _find_dimension_workbook(
        cls,
        extract_dir: Path,
        dimension_name: str,
    ) -> Optional[Path]:
        expected = cls._normalize_dimension_name(dimension_name)
        matches = [
            path
            for path in extract_dir.rglob("*.xlsx")
            if expected in cls._normalize_dimension_name(path.stem)
        ]
        if len(matches) > 1:
            raise ExportValidationError("维度 {} 出现多个工作簿".format(dimension_name))
        return matches[0] if matches else None

    @classmethod
    def _validate_dimension_workbook(
        cls,
        workbook_path: Path,
        expected_names: List[str],
    ) -> str:
        try:
            import openpyxl
        except ImportError as exc:
            raise ExportValidationError("缺少 openpyxl，无法验证导出文件") from exc

        try:
            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=True,
                data_only=True,
            )
            try:
                sheet = workbook.active
                sheet.reset_dimensions()
                rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        except (OSError, zipfile.BadZipFile) as exc:
            # 损坏到无法打开的工作簿属于残缺文件，按缺文件处理
            raise ExportWorkbookIncomplete(
                "{} 无法作为工作簿打开，按缺文件处理".format(workbook_path.name)
            ) from exc
        if not rows or not any(value not in (None, "") for value in rows[0]):
            # 连表头都没有的残缺文件不能当作"无数据"结论，按缺文件处理
            raise ExportWorkbookIncomplete(
                "{} 连表头都没有，按缺文件处理".format(workbook_path.name)
            )
        data_rows = [row for row in rows[1:] if any(value not in (None, "") for value in row)]
        if not data_rows:
            return "no_data"
        headers = [str(value or "").strip() for value in rows[0]]
        name_indexes = [
            index
            for index, header in enumerate(headers)
            if header in {"企业名称", "公司名称", "主体名称", "查询企业名称", "被查询企业名称"}
            or "被查询企业" in header
        ]
        if not name_indexes:
            raise ExportValidationError(
                "{} 缺少可核对的被查企业名称列".format(workbook_path.name)
            )
        name_index = name_indexes[0]
        actual = {
            cls._normalize_company_name(row[name_index])
            for row in data_rows
            if len(row) > name_index and row[name_index] not in (None, "")
        }
        expected = {cls._normalize_company_name(name) for name in expected_names}
        if actual != expected:
            raise ExportValidationError(
                "{} 的企业集合与本次上传名单不一致".format(workbook_path.name)
            )
        return "completed"

    @classmethod
    def _statistics_no_data_dimensions(
        cls,
        extract_dir: Path,
        expected_names: List[str],
        dimension_codes: List[str],
    ) -> set:
        statistics_path = next(
            (path for path in extract_dir.rglob("*.xlsx") if path.stem == "统计表"),
            None,
        )
        if statistics_path is None:
            return set()
        try:
            import openpyxl
        except ImportError as exc:
            raise ExportValidationError("缺少 openpyxl，无法验证统计表") from exc
        workbook = openpyxl.load_workbook(
            statistics_path,
            read_only=True,
            data_only=True,
        )
        try:
            sheet = workbook.active
            sheet.reset_dimensions()
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        header_index = next(
            (
                index
                for index, row in enumerate(rows)
                if row and str(row[0] or "").strip() == "公司名称"
            ),
            None,
        )
        if header_index is None:
            return set()
        headers = [str(value or "").strip() for value in rows[header_index]]
        data_rows = [
            row
            for row in rows[header_index + 1 :]
            if row and row[0] not in (None, "")
        ]
        expected = {cls._normalize_company_name(name) for name in expected_names}
        actual = {cls._normalize_company_name(row[0]) for row in data_rows}
        if actual != expected:
            raise ExportValidationError("统计表企业集合与本次上传名单不一致")
        no_data = set()
        for code in dimension_codes:
            name = DIMENSION_NAMES.get(code, code)
            try:
                column = headers.index(name)
            except ValueError:
                continue
            values = [row[column] for row in data_rows if len(row) > column]
            if values and all(
                value not in (None, "") and float(value) == 0 for value in values
            ):
                no_data.add(code)
        return no_data

    def download_current_group(
        self,
        state: ExportState,
        task: Dict[str, Any],
        state_path: Path,
    ) -> List[str]:
        """逐维保留本批有效结果，只重排缺失维度；缺文件达到重试上限后记 missing。"""
        url = str(task.get("url") or "")
        if not url:
            raise ExportValidationError("对应下载任务没有下载地址")
        if not url.startswith("http"):
            url = self.ZSK_BASE + "/" + url.lstrip("/")
        staging_dir = Path(state.staging_dir)
        batch_root = staging_dir / "dimension-batches"
        batch_root.mkdir(parents=True, exist_ok=True)
        batch_number = len(state.batch_history) + 1
        archive_path = batch_root / "batch-{:03d}.zip".format(batch_number)
        extract_dir = batch_root / "batch-{:03d}_files".format(batch_number)
        while archive_path.exists() or extract_dir.exists():
            batch_number += 1
            archive_path = batch_root / "batch-{:03d}.zip".format(batch_number)
            extract_dir = batch_root / "batch-{:03d}_files".format(batch_number)
        archive_path.write_bytes(self.client.request_bytes("GET", url, timeout=120))
        if not zipfile.is_zipfile(archive_path):
            state.status = "paused"
            state.message_zh = "本批下载内容不是有效 ZIP"
            save_export_state(state_path, state)
            raise ExportValidationError(state.message_zh)
        extract_dir.mkdir(parents=True, exist_ok=True)
        try:
            self._safe_extract(archive_path, extract_dir)
        except ExportValidationError as exc:
            # 整批无法安全解包（含跨盘符路径）：视同本批各维度均未取得文件，留痕后继续
            requeue = []
            marked = []
            for code in state.active_group:
                if state.dimension_results.get(code) in TERMINAL_DIMENSION_RESULTS:
                    continue
                if self._count_missing_file(state, code, requeue):
                    marked.append(code)
            state.batch_history.append(
                {
                    "dimensions": list(state.active_group),
                    "task_id": str(task.get("task_id") or state.task_id),
                    "archive_path": str(archive_path),
                    "completed_dimensions": [],
                    "missing_dimensions": list(requeue),
                    "marked_missing_dimensions": list(marked),
                    "result": "archive_invalid",
                    "detail": str(exc),
                }
            )
            if requeue:
                state.pending_groups.insert(0, requeue)
            finished = self._close_batch(state, archive_path, staging_dir)
            state.message_zh = (
                self._completion_message(state)
                if finished
                else "本批 ZIP 未通过安全校验（{}），未取得维度将重新触发".format(exc)
            )
            save_export_state(state_path, state)
            return []

        completed = []
        missing = []
        marked_missing = []
        try:
            explicit_no_data = self._statistics_no_data_dimensions(
                extract_dir,
                state.company_names,
                state.active_group,
            )
        except (ExportValidationError, TypeError, ValueError) as exc:
            state.status = "paused"
            state.message_zh = str(exc)
            save_export_state(state_path, state)
            raise ExportValidationError(state.message_zh) from exc
        for code in state.active_group:
            if state.dimension_results.get(code) in TERMINAL_DIMENSION_RESULTS:
                continue
            name = DIMENSION_NAMES.get(code, code)
            try:
                workbook_path = self._find_dimension_workbook(extract_dir, name)
            except ExportValidationError as exc:
                state.status = "paused"
                state.message_zh = str(exc)
                save_export_state(state_path, state)
                raise
            if workbook_path is None:
                if code in explicit_no_data:
                    state.dimension_results[code] = "no_data"
                    completed.append(code)
                elif self._count_missing_file(state, code, missing):
                    marked_missing.append(code)
                continue
            try:
                result = self._validate_dimension_workbook(
                    workbook_path,
                    state.company_names,
                )
            except ExportWorkbookIncomplete:
                # 残缺文件（无法打开或连表头都没有）按缺文件处理，不写成"无数据"结论
                if code in explicit_no_data:
                    state.dimension_results[code] = "no_data"
                    completed.append(code)
                elif self._count_missing_file(state, code, missing):
                    marked_missing.append(code)
                continue
            except ExportValidationError as exc:
                state.status = "paused"
                state.message_zh = str(exc)
                save_export_state(state_path, state)
                raise
            target = staging_dir / workbook_path.name
            if target.exists():
                if self._file_sha256(target) != self._file_sha256(workbook_path):
                    # 同名不同内容：新文件改名保留双份、沿用旧文件、留痕后继续
                    renamed = self._conflict_copy_path(staging_dir, workbook_path)
                    shutil.copy2(workbook_path, renamed)
                    state.batch_history.append(
                        {
                            "dimensions": [code],
                            "task_id": str(task.get("task_id") or state.task_id),
                            "archive_path": str(archive_path),
                            "result": "file_conflict",
                            "detail": "{} 两次取数内容不一致：沿用旧文件，新文件已改名 {}".format(
                                workbook_path.name,
                                renamed.name,
                            ),
                        }
                    )
            else:
                shutil.copy2(workbook_path, target)
            state.dimension_results[code] = result
            completed.append(code)

        state.batch_history.append(
            {
                "dimensions": list(state.active_group),
                "task_id": str(task.get("task_id") or state.task_id),
                "archive_path": str(archive_path),
                "completed_dimensions": list(completed),
                "missing_dimensions": list(missing),
                "marked_missing_dimensions": list(marked_missing),
                "result": "partial" if (missing or marked_missing) else "completed",
            }
        )
        if len(state.active_group) > 1 and not state.single_dimension_mode:
            state.batch_failure_streak = 0
        if missing:
            state.pending_groups.insert(0, missing)
        finished = self._close_batch(state, archive_path, staging_dir)
        if finished:
            state.message_zh = self._completion_message(state)
        elif marked_missing:
            state.message_zh = (
                "本批已保留 {} 个维度，{} 个维度记为未取得，仅继续未完成维度".format(
                    len(completed),
                    len(marked_missing),
                )
            )
        else:
            state.message_zh = "本批已保留 {} 个维度，仅继续未完成维度".format(len(completed))
        save_export_state(state_path, state)
        return completed

    @staticmethod
    def _count_missing_file(state: ExportState, code: str, requeue: List[str]) -> bool:
        """累计"下载批次里没有该维度文件"的次数；达到上限记 missing 终态并返回 True。"""
        state.missing_counts[code] = state.missing_counts.get(code, 0) + 1
        if state.missing_counts[code] >= MISSING_FILE_RETRY_LIMIT:
            state.dimension_results[code] = "missing"
            return True
        requeue.append(code)
        return False

    @staticmethod
    def _conflict_copy_path(staging_dir: Path, workbook_path: Path) -> Path:
        """同名冲突时给新文件改名：原名.本批.xlsx；再冲突追加序号，永不覆盖。"""
        candidate = staging_dir / "{}.本批{}".format(
            workbook_path.stem, workbook_path.suffix
        )
        serial = 1
        while candidate.exists():
            serial += 1
            candidate = staging_dir / "{}.本批{}{}".format(
                workbook_path.stem, serial, workbook_path.suffix
            )
        return candidate

    @staticmethod
    def _completion_message(state: ExportState) -> str:
        """完成消息如实统计：有未取得维度时不得笼统说"全部完成"。"""
        total = len(state.required_dimensions)
        missing_total = sum(
            1
            for code in state.required_dimensions
            if state.dimension_results.get(code) == "missing"
        )
        if missing_total:
            return "完成 {} 个维度，{} 个维度未取得".format(
                total - missing_total, missing_total
            )
        return "{} 个关联方核查维度已全部完成".format(total)

    def _close_batch(
        self,
        state: ExportState,
        archive_path: Path,
        staging_dir: Path,
    ) -> bool:
        """清空活动组并刷新完成状态；返回是否全部维度进入终态。"""
        state.active_group = []
        state.task_id = ""
        state.task_name = ""
        state.poll_windows = 0
        state.archive_path = str(archive_path)
        state.extract_dir = str(staging_dir)
        finished = all(
            state.dimension_results.get(code) in TERMINAL_DIMENSION_RESULTS
            for code in state.required_dimensions
        )
        state.status = "completed" if finished else "waiting"
        return finished

    def run_to_completion(
        self,
        state: ExportState,
        state_path: Path,
        *,
        max_polls: int = 18,
    ) -> Path:
        """串行驱动全部未完成维度，全部进入终态（含 missing）后才返回。"""
        while True:
            if state.status == "paused":
                raise ExportError(state.message_zh or "维度导出已暂停")
            finished = bool(state.required_dimensions) and all(
                state.dimension_results.get(code) in TERMINAL_DIMENSION_RESULTS
                for code in state.required_dimensions
            )
            if finished:
                state.status = "completed"
                state.extract_dir = state.extract_dir or state.staging_dir
                state.message_zh = self._completion_message(state)
                save_export_state(state_path, state)
                return Path(state.extract_dir)
            if not state.active_group and not self.trigger_next_group(state, state_path):
                raise ExportError("仍有未完成维度，但没有可触发的维度组")
            try:
                task = self.wait_for_task(state, state_path, max_polls=max_polls)
            except ExportTaskAmbiguous as exc:
                state.status = "paused"
                state.message_zh = str(exc)
                save_export_state(state_path, state)
                raise
            except ExportTaskNotCorrelated:
                self.handle_poll_window_failure(state, state_path)
                continue
            self.download_current_group(state, task, state_path)
            if state.status != "completed":
                self._sleep(self._randint(15, 30))

    def download_and_validate(
        self,
        state: ExportState,
        task: Dict[str, Any],
        state_path: Path,
    ) -> Path:
        """下载到暂存目录并核对 ZIP、路径和公司集合。"""
        url = str(task.get("url") or "")
        if not url:
            raise ExportValidationError("对应下载任务没有下载地址")
        if not url.startswith("http"):
            url = self.ZSK_BASE + "/" + url.lstrip("/")
        staging_dir = Path(state.staging_dir)
        staging_dir.mkdir(parents=True, exist_ok=True)
        archive_path = staging_dir / "complete-dimensions.zip"
        archive_path.write_bytes(self.client.request_bytes("GET", url, timeout=120))
        if not zipfile.is_zipfile(archive_path):
            state.status = "failed"
            state.message_zh = "下载内容不是有效 ZIP"
            save_export_state(state_path, state)
            raise ExportValidationError(state.message_zh)

        extract_dir = (
            staging_dir
            if state.direct_delivery
            else staging_dir / "complete-dimensions_files"
        )
        extract_dir.mkdir(parents=True, exist_ok=True)
        try:
            self._safe_extract(archive_path, extract_dir)
            self._validate_companies(extract_dir, state.company_names)
        except ExportValidationError as exc:
            state.status = "failed"
            state.message_zh = str(exc)
            save_export_state(state_path, state)
            raise

        state.task_id = str(task.get("task_id") or state.task_id)
        state.task_name = str(task.get("name") or state.task_name)
        state.status = "completed"
        state.archive_path = str(archive_path)
        state.extract_dir = str(extract_dir)
        state.message_zh = "完整维度导出已下载并通过企业名称验证"
        save_export_state(state_path, state)
        return extract_dir
