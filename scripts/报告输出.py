# 八张工作表共享同一核查结果，使用 Excel 原生导航、筛选和相对文件链接。
from collections import Counter
from datetime import datetime
import hashlib
import math
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.workbook.defined_name import DefinedName


SHEETS = {
    "目录": "先看核查结论，再按公司查看证据；所有原始资料保持原样。",
    "核查结论": "本次发现了哪些疑似关联方：叫什么名字、风险多高、与被审计单位疑似什么关系、依据是什么。",
    "疑似关联方复核底稿": "每家疑似关联方一行；自动初判与人工复核意见分开填写。",
    "证据明细": "每条证据的判断依据及实际参与匹配的原始字段，可按证据编号筛选。",
    "纳入核查的公司名单": "每家公司为何纳入、取得什么资料、实际查到哪一步。",
    "各家资料取得情况": "每家公司、每个维度的资料取得情况及其对核查的影响。",
    "核查范围与数量说明": "本次核查覆盖的范围，以及报告里各项数字是怎么算出来的。",
    "核查任务说明": "任务入口、范围、日期、原始目录和判断口径。",
}
NAVY = "1F4E79"
PALE = "D6E4F0"
LIGHT = "F3F7FB"
RISK_COLORS = {"高": "FCE4D6", "中": "FFF2CC", "低": "E2F0D9"}
# 概览与汇总的"主要线索摘要"长度上限：按整条证据取舍，不切断单条证据。
SUMMARY_LIMIT = 240
# "取数范围外"状态：导出本就不覆盖该公司该维度，属正常情况，不计入数据缺口。
OUT_OF_SCOPE_STATUS = "本次取数范围外，导出未覆盖该企业"
DIMENSIONS = {
    "basic": ("基础工商信息", "工商指纹、人员、客商画像、历史痕迹"),
    "shareholder": ("股东信息", "人员、共同股东"),
    "shareholder_new": ("最新公示股东", "人员、共同股东"),
    "actual_controller": ("实际控制人", "共同实际控制人"),
    "ultimate_beneficiary": ("最终受益人", "共同最终受益人"),
    "main_persons": ("主要人员（高管）", "关键人员"),
    "core_team": ("核心团队", "关键人员"),
    "invest_new": ("对外投资（新）", "直接投资"),
    "holding": ("参控股企业", "参控股关系"),
    "invoice": ("发票信息", "地址补充"),
    "customer": ("客户", "仅用于候选来源"),
    "supplier": ("供应商", "仅用于候选来源"),
    "change": ("变更记录", "历史地址"),
    "legal_change": ("法定代表人变更", "历史法人"),
    "abnormal": ("经营异常", "背景资料，本次规则不直接使用"),
    "pledge": ("股权质押", "质押关系"),
    "mortgage": ("动产抵押", "抵押关系"),
    "trademark": ("商标", "同名商标"),
    "software": ("软件著作权", "同名软件著作权"),
    "wechat": ("微信公众号", "同名公众号"),
}


def risk_text(value):
    text = str(value)
    if "硬" in text or text.startswith("高"):
        return "高"
    if "可疑" in text or text.startswith("中"):
        return "中"
    return "低"


def pair_key(a, b):
    return tuple(sorted((a, b)))


def pair_id(a, b):
    return "GX-" + hashlib.sha256("\n".join(pair_key(a, b)).encode("utf-8")).hexdigest()[:8].upper()


def put(ws, row, column, value):
    cell = ws.cell(row, column)
    cell.value = value
    if isinstance(value, str):
        # 原始企业名称和字段是数据，不能被 Excel 当成公式执行。
        cell.data_type = "s"
    return cell


def append_row(ws, values):
    """超过 Excel 单元格上限时续行保留全文，编号和短字段在续行重复。"""
    start = max(ws.max_row + 1, 4)
    parts = max([1] + [math.ceil(len(v) / 32767) for v in values if isinstance(v, str)])
    for offset in range(parts):
        for column, value in enumerate(values, 1):
            part = value[offset * 32767:(offset + 1) * 32767] if isinstance(value, str) and len(value) > 32767 else value
            put(ws, start + offset, column, part)
    return start


def section_header(ws, values):
    """同一张工作表里第二张表的列头：与第 3 行表头同款深蓝样式，肉眼可辨两张表的分界。
    finish_sheet 会把正文行统一改回浅色，调用方需在 finish_sheet 之后用 restyle_section_headers 还原。"""
    return append_row(ws, values)


def restyle_section_headers(ws, rows):
    for row in rows:
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="微软雅黑", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 30


def link(cell, sheet, row):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet.replace(chr(39), chr(39)*2)}'!A{row}", display=str(cell.value))
    cell.font = Font(name="微软雅黑", size=10, color=NAVY, underline="single")


def source_link(cell, source, output):
    relative = os.path.relpath(source["file"], Path(output).parent).replace("\\", "/")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=relative,
        location=f"'{source['sheet'].replace(chr(39), chr(39)*2)}'!{source['cell']}", display=str(cell.value))
    cell.font = Font(name="微软雅黑", size=10, color=NAVY, underline="single")


def keyed_link(cell, sheet, key):
    """按唯一编号定位，明细或汇总排序以后仍指向同一条记录。"""
    wb = cell.parent.parent
    name = "_位置_" + key.replace("-", "_")
    quoted = "'" + sheet.replace("'", "''") + "'"
    if name not in wb.defined_names:
        wb.defined_names.add(DefinedName(name, attr_text=f'INDEX({quoted}!$A:$A,MATCH("{key}",{quoted}!$A:$A,0))'))
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=name, display=str(cell.value))
    cell.font = Font(name="微软雅黑", size=10, color=NAVY, underline="single")


def create_sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = NAVY
    put(ws, 1, 1, "返回目录" if title != "目录" else "关联方识别与核查")
    if title != "目录":
        link(ws["A1"], "目录", 1)
    else:
        ws["A1"].font = Font(name="微软雅黑", size=14, bold=True, color=NAVY)
    for column in range(1, len(headers) + 1):
        ws.cell(1, column).fill = PatternFill("solid", fgColor=PALE)
    if len(headers) > 1:
        put(ws, 1, 2, title)
        ws["B1"].font = Font(name="微软雅黑", size=16, bold=True, color=NAVY)
    put(ws, 2, 1, SHEETS[title])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"].font = Font(name="微软雅黑", size=10, color="526579")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 30
    for column, header in enumerate(headers, 1):
        cell = put(ws, 3, column, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="微软雅黑", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(column)].width = widths[column - 1]
    ws.freeze_panes = "C4" if len(headers) > 3 else "A4"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if len(headers) > 8 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    return ws


def finish_sheet(ws, risk_column=None):
    border = Border(bottom=Side(style="hair", color="D9E2F3"))
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if not cell.hyperlink:
                cell.font = Font(name="微软雅黑", size=10, color="243746")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=LIGHT if cell.row % 2 == 0 else "FFFFFF")
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00" if isinstance(cell.value, float) else "#,##0"
        if risk_column:
            cell = row[risk_column - 1]
            if cell.value in RISK_COLORS:
                cell.fill = PatternFill("solid", fgColor=RISK_COLORS[cell.value])
                cell.font = Font(name="微软雅黑", size=10, bold=True, color="243746")
        lines = max([1] + [sum(max(1, math.ceil(len(part) * 1.6 / max(10, ws.column_dimensions[cell.column_letter].width)))
                              for part in str(cell.value or "").split("\n")) for cell in row])
        ws.row_dimensions[row[0].row].height = min(210, max(30, lines * 15))
    ws.auto_filter.ref = f"A3:{get_column_letter(ws.max_column)}{max(3, ws.max_row)}"
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{max(4, ws.max_row)}"


def group_evidence_blocks(ws, risk_column=5):
    """同一证据编号的连续多行视为一块：相邻块交替浅底色、块首行加粗上边框，肉眼可分辨块界。"""
    if ws.max_row < 4:
        return
    previous = None
    block = -1
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        if row[0].value != previous:
            previous = row[0].value
            block += 1
            for cell in row:
                cell.border = Border(top=Side(style="medium", color=NAVY), bottom=Side(style="hair", color="D9E2F3"))
        fill = PALE if block % 2 else "FFFFFF"
        for cell in row:
            if cell.column != risk_column:  # 风险列保留红黄绿底色双通道
                cell.fill = PatternFill("solid", fgColor=fill)


def clue_summary(item, limit=SUMMARY_LIMIT):
    """概览与汇总共用的一份摘要：按整条证据取舍，宁少列一条也不切断一条。"""
    items = list(item.get("evidence_items") or []) or [item.get("evidence", "")]
    kept, total = [], 0
    for text in items:
        if kept and total + len(text) + 3 > limit:
            break
        kept.append(text)
        total += len(text) + 3
    omitted = len(items) - len(kept)
    joined = " | ".join(kept)
    if omitted:
        joined += f"…（另 {omitted} 条见明细）"
    prefix = "多条独立线索相互印证；" if item.get("corroborated") else ""
    if item.get("converged"):
        prefix += "【收敛：全部命中来自高共用度指纹，疑似集中注册或代理记账】"
    return prefix + joined


def coverage_rows(names, companies, file_info, errors, target_set=None):
    target_set = set(target_set or ())
    rows = []
    incomplete = set()
    gap_dimensions = Counter()
    for name in names:
        for key, (label, impact) in DIMENSIONS.items():
            infos = [item for item in file_info if item["key"] == key]
            own = [item for item in infos if name in item["companies"] or name in item["scope"]]
            record = next((item for item in own if name in item["companies"]), None)
            if any(item.get("malformed") for item in own or infos):
                status = "资料字段不足"
            elif record:
                status = "已取得记录"
            elif any(item["exists"] and not item["readable"] for item in own or infos):
                status = "读取失败"
            elif any(item["terminal"] == "no_data" and name in item["scope"] for item in own):
                status = "明确无数据"
            elif any(item["readable"] and item["terminal"] == "completed" and name in item["scope"] for item in own):
                status = "已核对范围，未见记录"
            elif any(item["readable"] for item in infos):
                # 维度文件存在但该公司没有记录：按取数范围区分两类——
                # 范围内（取数说明列名或被审计单位）应取得而未取得，属真缺口；
                # 范围外（导出本就不覆盖该公司该维度，如名单里的客户供应商）属正常，不计缺口。
                if name in target_set or any(name in item["scope"] for item in infos):
                    status = "未见该企业记录，范围待核实"
                else:
                    status = OUT_OF_SCOPE_STATUS
            else:
                status = "未取得"
            detail = ""
            if key == "basic" and name in companies:
                company = companies[name]
                fields = {"法定代表人": company.legal_person, "联系电话": company.phones,
                    "邮箱": company.emails, "地址": company.addresses, "注册资本": company.capital,
                    "成立日期": company.found_date, "参保人数": company.insured, "经营范围": company.business_scope}
                blanks = [label for label, value in fields.items() if value is None or value == "" or value == [] or value == set()]
                if blanks:
                    status = "已取得记录，部分字段缺失或不可用"
                    detail = "缺失或不能按本次口径解析的字段：" + "、".join(blanks)
            if status not in {"已取得记录", "明确无数据", "已核对范围，未见记录", OUT_OF_SCOPE_STATUS} and key not in {"customer", "supplier", "abnormal"}:
                incomplete.add(name)
                gap_dimensions[name] += 1
            chosen = record or next(iter(own or infos), {})
            rows.append([name, label, status, impact, detail, chosen.get("file", ""), chosen.get("created_at", "未记录") or "未记录"])
        if name not in companies:
            incomplete.add(name)
        if any(name in str(item.get("message", "")) or name == item.get("source") for item in errors):
            incomplete.add(name)
    # 重度缺口：未取得基础工商记录，或缺失维度达 10 个及以上；其余缺口公司为轻度。
    severe = {name for name in names if name not in companies or gap_dimensions.get(name, 0) >= 10}
    return rows, incomplete, severe


def write_report(out_path, summary, all_hits, companies, target_display,
                 *, errors=None, limitations=None, task_mode="", scope_metadata=None,
                 file_info=None, object_records=None, target_set=None):
    errors, limitations = list(errors or []), list(limitations or [])
    scope = dict(scope_metadata or {})
    target_set = set(target_set or target_display.split(" / "))
    file_info = list(file_info or [])
    objects = {}
    for record in object_records or []:
        name = record["name"]
        if name not in objects:
            objects[name] = dict(record)
        else:
            for key in ("reasons", "paths", "notes", "source_files"):
                objects[name][key] = list(dict.fromkeys(objects[name].get(key, []) + record.get(key, [])))
            objects[name]["sources"] = objects[name].get("sources", []) + record.get("sources", [])
    names = sorted(set(companies) | target_set | set(objects) | {name for info in file_info for name in info["scope"]}
                   | {name for info in file_info if info["key"] == "basic" for name in info["companies"]})
    coverage, incomplete, severe = coverage_rows(names, companies, file_info, errors, target_set)
    counts = Counter(risk_text(item["max_level"]) for item in summary)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = create_sheet(wb, "目录", ["工作表", "阅读用途", "数据行数"], [26, 84, 16])
    for title, description in list(SHEETS.items())[1:]:
        row = append_row(ws, [title, description, None])
        link(ws.cell(row, 1), title, 1)
    overview = create_sheet(wb, "核查结论", ["项目", "数量", "", "内容 / 说明"], [34, 10, 36, 82])
    # 第一行直接给结论：发现了几家、各是什么风险。读者打开报告十秒内要知道答案。
    checked_count = len(set(companies) - target_set)
    if summary:
        risk_parts = [f"{risk}风险 {counts[risk]} 家" for risk in ("高", "中", "低") if counts[risk]]
        conclusion = f"在被审计单位“{target_display}”之外，本次共核查 {checked_count} 家公司，发现疑似关联方 {len(summary)} 家（{'、'.join(risk_parts)}），名单见下表。"
    else:
        conclusion = f"在被审计单位“{target_display}”之外，本次共核查 {checked_count} 家公司，未发现疑似关联方。"
    append_row(overview, ["本次发现了什么", "", "", conclusion])
    append_row(overview, ["这份报告能说明什么（必读）", "", "",
        scope.get("本次核查结论（必读）", "本报告只是自动初判线索，不是审计结论：列出的疑似关联方需经审计程序核实；未发现不代表不存在关联关系，关联方是否完整须结合其他审计程序确认")])
    for item in limitations:
        if item.get("category") == "取数通道降级":
            append_row(overview, ["取数通道降级提示", "", "",
                f"{item.get('message', '')}（来源：{item.get('source', '')}；受影响范围详见各家资料取得情况）"])
    if not counts["高"] and not counts["中"]:
        append_row(overview, ["重要提示", "", "",
            "本结果不提供关联方完整性保证：未发现高、中风险线索；未发现不代表不存在关联关系，关联方是否完整须结合其他审计程序确认"])
    if summary:
        # 名单是这张表的正文：公司叫什么、风险多高、疑似什么关系、发现了什么，一眼看完。
        clue_header_row = section_header(overview,
            ["疑似关联方名单", "风险等级", "与被审计单位的疑似关系", "发现了什么（完整证据见证据明细）"])
    else:
        clue_header_row = None
    overview_links = []
    converged_items = [item for item in summary if item.get("converged")]
    for item in summary:
        if item.get("converged"):
            continue  # 只有集中注册类线索的公司不列入名单，折叠为一行说明
        row = append_row(overview, [item["company"], risk_text(item["max_level"]),
                                    item["relation_to_target"],
                                    clue_summary(item)])
        overview_links.append((row, pair_key(item["company_a"], item["company_b"])))
    if converged_items:
        append_row(overview, [f"另有 {len(converged_items)} 家公司未列入名单", "", "",
            "这些公司与被审计单位撞上的地址或电话同时被 5 家及以上公司共用，疑似代理记账或集中注册，审计价值低；逐家明细见复核底稿后段"])
    quality = create_sheet(wb, "核查范围与数量说明", ["项目", "数量", "这个数字怎么算出来的"], [36, 46, 92])
    metrics = [
        ["被审计单位", target_display, "本次核查以被审计单位为基准，每家公司都与它比对"],
        ["纳入核查的公司数（不含被审计单位）", len([n for n in names if n not in target_set]), "候选及资料未取得的也算在内，按名称去重"],
        ["其中由程序主动发现的公司数", len(set(objects) - target_set) if task_mode == "discovery" else 0, "只有主动发现模式才有；被发现不等于关联方"],
        ["取得基础资料的公司数（不含被审计单位）", len(set(companies) - target_set), "基础工商表可读取、实际参与了比对的公司"],
        ["与被审计单位实际比对的次数", scope.get("被审计单位与各公司的实际比对数", checked_count), "每家取得基础资料的公司与被审计单位比对一次"],
        ["发现疑似关联方", len(summary), "同一家公司命中多条证据只计一家"],
        ["证据条数", len({hit.get("evidence_id", str(i)) for i, hit in enumerate(all_hits)}), "按证据编号去重；同一条证据的多个来源行不重复计数"],
    ]
    if converged_items:
        metrics.append(["其中只有集中注册类线索的公司数", len(converged_items),
            "撞上的地址或电话被 5 家及以上公司共用，疑似代理记账或集中注册，已列于复核底稿后段"])
    metrics += [
        ["存在资料缺口的公司数", len(incomplete), "缺资料、空字段或规则执行失败；具体影响见各家资料取得情况"],
        ["其中重度缺口公司数", len(severe), "未取得基础工商记录，或缺失维度达 10 个及以上；缺口实质削弱核查结论"],
        ["其中轻度缺口公司数", len(incomplete - severe), "其余存在缺口的公司；个别维度或字段缺失，影响相对有限"],
        ["读取、字段或规则错误条数", len(errors), "有错误不能当作未发现"],
        ["发现过程提示条数", sum(item.get("category") == "候选发现不完整" for item in limitations), "发现过程中取数失败会影响候选范围是否完整"],
    ]
    for record in metrics:
        append_row(quality, record)

    roster = create_sheet(wb, "纳入核查的公司名单",
        ["公司名称", "怎么进名单的", "发现理由 / 名单来源", "层级", "直接来源", "关系路径", "资料情况", "核查结果", "法定代表人", "成立日期", "注册资本（人民币元）", "参保人数", "来源记录 / 说明"],
        [30, 21, 42, 9, 28, 54, 23, 28, 17, 15, 20, 12, 52])
    involved = {value for item in summary for value in (item["company_a"], item["company_b"])}
    for name in names:
        record = objects.get(name, {})
        company = companies.get(name)
        origin = "被审计单位" if name in target_set else record.get("relation_type", "已有数据中的企业")
        reason = "；".join(record.get("reasons", []))
        if record.get("source_files"):
            reason += "\n名单文件：" + "；".join(Path(p).name for p in record["source_files"])
        if name in target_set:
            outcome = "被审计单位（比对基准）"
        else:
            outcome = "未纳入比对，缺少基础资料" if not company else ("比对发现疑似关联线索" if name in involved else "已比对，未发现关联线索")
        if any(item.get("category") == "规则异常" and name in item.get("message", "") for item in errors):
            outcome = "部分规则执行失败；" + ("已有命中线索" if name in involved else "其余已执行规则未发现")
        source_notes = []
        for source in record.get("sources", []):
            if source.get("file"):
                source_notes.append(f"{Path(source['file']).name} / {source['sheet']}!{source['cell']}：{source['value']}")
            else:
                ratio_text = "比例未取得" if source.get("ratio") is None else str(source["ratio"]) + "%"
                source_notes.append(f"股权查询：{source.get('company', '')}（{source.get('company_id', '')}） / {source.get('relation_type', '')} / {source.get('counterparty', '')} / {ratio_text}")
        append_row(roster, [name, origin, reason, record.get("depth", ""), record.get("parent_name", ""),
            "\n".join(record.get("paths", [])), "存在资料缺口" if name in incomplete else "已取得本次核查资料", outcome,
            company.legal_person if company else "", company.found_date if company else "", company.capital if company else None,
            company.insured if company else None, "\n".join(source_notes + record.get("notes", []))])

    summary_ws = create_sheet(wb, "疑似关联方复核底稿",
        ["线索编号", "公司名称", "风险等级", "与被审计单位的疑似关系", "发现了什么", "证据条数", "查看全部证据", "建议审计程序", "资料限制", "人工复核意见", "复核人 / 日期", "序号"],
        [16, 30, 12, 34, 48, 12, 20, 45, 28, 40, 24, 10])
    summary_rows = {}
    for index, item in enumerate(summary, 1):
        key = pair_key(item["company_a"], item["company_b"])
        # GX- 编号保留作链接锚点；末尾序号供底稿引用和口头沟通。
        row = append_row(summary_ws, [pair_id(*key), item["company"], risk_text(item["max_level"]),
            item["relation_to_target"], clue_summary(item),
            item["hit_count"], "查看全部证据", item["suggestion"],
            "存在资料缺口，详见各家资料取得情况" if set(key) & incomplete else "", "", "", f"{index:02d}"])
        summary_rows[key] = row
    for row, key in overview_links:
        keyed_link(overview.cell(row, 1), "疑似关联方复核底稿", pair_id(*key))

    evidence_ws = create_sheet(wb, "证据明细",
        ["证据编号", "被审计单位", "公司", "核查类别", "风险初判", "判断依据", "来源公司", "原始字段", "原始值", "原始文件 / 工作表 / 单元格", "返回底稿", "案例参考（非本次证据）", "序号"],
        [24, 27, 27, 22, 12, 55, 27, 23, 44, 48, 18, 33, 10])
    first_evidence = {}
    evidence_seq = {}
    for number, hit in enumerate(all_hits, 1):
        key = pair_key(hit["company_a"], hit["company_b"])
        evidence_id = hit.get("evidence_id", f"ZJ-{number:05d}")
        if evidence_id not in evidence_seq:
            # 同一证据编号的多行来源共享同一个"证-XX"序号，按首次出现顺序编号。
            evidence_seq[evidence_id] = f"证-{len(evidence_seq) + 1:02d}"
        # 每对必有一方是被审计单位，固定放"被审计单位"列，另一方放"公司"列。
        target_side = hit["company_a"] if hit["company_a"] in target_set else hit["company_b"]
        other_side = hit["company_b"] if target_side == hit["company_a"] else hit["company_a"]
        sources = hit.get("sources") or [None]
        for source in sources:
            source_text = (os.path.relpath(source["file"], Path(out_path).parent) + "\n" + source["sheet"] + "!" + source["cell"]) if source else "来源位置未记录，须人工核对"
            row = append_row(evidence_ws, [evidence_id, target_side, other_side, hit["dimension"],
                risk_text(hit["level"]), hit["evidence"], source["company"] if source else "", source["field"] if source else "",
                source["value"] if source else "", source_text, "返回对应底稿", hit["case_ref"], evidence_seq[evidence_id]])
            first_evidence.setdefault(key, evidence_id)
            if source:
                source_link(evidence_ws.cell(row, 10), source, out_path)
            if key in summary_rows:
                keyed_link(evidence_ws.cell(row, 11), "疑似关联方复核底稿", pair_id(*key))
    for key, row in summary_rows.items():
        if key in first_evidence:
            keyed_link(summary_ws.cell(row, 7), "证据明细", first_evidence[key])
    gaps_ws = create_sheet(wb, "各家资料取得情况",
        ["公司 / 项目", "资料维度 / 类型", "取得或处理情况", "影响的核查", "说明", "来源文件", "取数任务创建时间"],
        [30, 25, 36, 38, 60, 48, 25])
    # 维度级汇总置顶：全部公司都未取得的维度先总起一笔，说清影响面。
    label_key = {label: key for key, (label, _impact) in DIMENSIONS.items()}
    all_missing = []
    for key, (label, _impact) in DIMENSIONS.items():
        dim_rows = [record for record in coverage if record[1] == label]
        if dim_rows and all(record[2] == "未取得" for record in dim_rows):
            all_missing.append(label)
    degraded_mode = any(item.get("category") == "取数通道降级" for item in limitations)
    if all_missing:
        impacts = "、".join(dict.fromkeys(DIMENSIONS[label_key[label]][1] for label in all_missing))
        append_row(gaps_ws, ["全部核查对象", "维度级汇总",
            f"本次未取得维度：{'、'.join(all_missing)}，影响全部 {len(names)} 家", impacts,
            "相关核查规则整体受限" + ("；逐公司重复行已折叠，不再逐家列出" if degraded_mode else "；逐公司明细见下文"), "", ""])
    # 走完流程的公司若在某维度上"已核对、未见记录"或"明确无数据"，那不是缺口；
    # 两条以上就折叠成一条，缺口（未取得、读取失败、范围待核实）永远逐家列出。
    quiet_statuses = {"已核对范围，未见记录", "明确无数据"}
    checked_names = set(companies)
    by_label = {}
    for record in coverage:
        by_label.setdefault(record[1], []).append(record)
    quiet = {}
    for label, records in by_label.items():
        quiet_rows = [record for record in records
                      if record[0] in checked_names and record[2] in quiet_statuses]
        statuses = {record[2] for record in quiet_rows}
        if len(quiet_rows) >= 2 and len(statuses) == 1:
            quiet[label] = (len(quiet_rows), next(iter(statuses)), len(records) - len(quiet_rows))
    # "取数范围外"不是缺口：同一维度两家及以上时折叠为一行说明，个别出现的保留单行。
    out_scope = {label: len([record for record in records if record[2] == OUT_OF_SCOPE_STATUS])
                 for label, records in by_label.items()}
    out_scope = {label: count for label, count in out_scope.items() if count >= 2}
    # 涉及命中公司的缺口行排在前面并标记，读者先看到需要结合结论阅读的缺口。
    ok_status = {"已取得记录", "明确无数据", "已核对范围，未见记录", OUT_OF_SCOPE_STATUS}
    candidate_keys = {"customer", "supplier", "abnormal"}
    front, rest = [], []
    for record in coverage:
        if degraded_mode and record[1] in all_missing:
            continue
        if record[1] in quiet and record[0] in checked_names and record[2] in quiet_statuses:
            continue
        if record[1] in out_scope and record[2] == OUT_OF_SCOPE_STATUS:
            continue
        values = list(record)
        is_gap = record[2] not in ok_status and label_key.get(record[1]) not in candidate_keys
        if record[0] in involved and is_gap:
            values[4] = ("涉及命中；" + values[4]) if values[4] else \
                "涉及命中；本行缺口指该公司在本维度的自身记录未取得，已形成的命中来自对方登记或其他维度，两者不必然互相削弱"
            front.append(values)
        else:
            rest.append(values)
    for label, count in out_scope.items():
        rest.append(["取数范围外公司", label, OUT_OF_SCOPE_STATUS, DIMENSIONS[label_key[label]][1],
            f"该维度 {count} 家不在本次取数范围内（导出本就不覆盖），属正常情况，不计入数据缺口", "", ""])
    for label, (count, status, others) in quiet.items():
        note = "；其余公司单独列于本表前面各行" if others else ""
        rest.append(["全部核查对象" if not others else "其余核查对象", label,
            status, DIMENSIONS[label_key[label]][1],
            (f"该维度 {count} 家已核对，未见记录" if status == "已核对范围，未见记录"
             else f"该维度 {count} 家经统计表确认无数据") + note, "", ""])
    for values in front + rest:
        if values[5]:
            values[5] = os.path.relpath(values[5], Path(out_path).parent)
        append_row(gaps_ws, values)
    for item in errors + limitations:
        append_row(gaps_ws, [item.get("source", ""), item.get("category", ""), "需复核", "按说明确认受影响范围", item.get("message", ""), "", ""])
    explanation = create_sheet(wb, "核查任务说明", ["项目", "内容"], [32, 115])
    mode_names = {"discovery": "主动发现", "list_check": "名单核查", "existing_export": "已有数据离线核查"}
    for values in [["任务模式", mode_names.get(task_mode, task_mode or "未说明")], ["被审计单位", target_display],
                   ["报告生成时间", datetime.now().astimezone().isoformat(timespec="seconds")]] + [[key, value] for key, value in scope.items()]:
        append_row(explanation, values)
    append_row(explanation, ["阅读方法", "核查结论 → 复核底稿 → 证据明细 → 原始单元格。证据明细同一编号的多行是同一条证据的不同来源。"])
    append_row(explanation, ["筛选与链接", "底稿与证据之间按编号定位，排序后仍可往返。查看全部证据时，按公司列筛选；来源文件同时标明工作表和单元格。"])
    append_row(explanation, ["原始资料", "报告与原始资料目录应一起移动。链接使用相对路径；若客户端只打开文件，按同时显示的工作表和单元格定位。"])
    append_row(explanation, ["日期口径", "核查基准日用于计算成立时间等规则；原始记录日期见对应原表字段。取数任务创建时间不等同于记录生效日期。"])
    # 标签页顺序与目录一致：核查结论 → 复核底稿 → 证据明细 → 其余参考表。
    wb._sheets = [wb[title] for title in SHEETS]
    for ws in wb.worksheets:
        finish_sheet(ws, {"疑似关联方复核底稿": 3, "证据明细": 5}.get(ws.title))
    if clue_header_row:
        restyle_section_headers(overview, [clue_header_row])
    group_evidence_blocks(evidence_ws)
    for row, title in enumerate(list(SHEETS)[1:], 4):
        put(wb["目录"], row, 3, max(0, wb[title].max_row - 3))
    wb.active = 0
    wb.save(out_path)
    wb.close()
